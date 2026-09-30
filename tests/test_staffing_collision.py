"""Name/kind collision review and enriched summary."""

from __future__ import annotations

from dataclasses import asdict

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import PERM_STAFFING_READ, PERM_STAFFING_WRITE
from app.identity.principal import Principal
from app.models.identity import Project
from app.models.staffing import StaffingAttendanceFact, StaffingImportBatch
from app.staffing import service as staffing_service
from app.staffing.parse import KIND_CONTRACT, KIND_FORMAL, parse_daily_report
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401
import app.models.staffing  # noqa: F401
from tests.test_staffing_parse import _build_sample_xlsx


def _principal(project_ids: tuple[str, ...] = ("proj2515",)) -> Principal:
    return Principal(
        user_id="u1",
        tenant_id="autley",
        username="liu",
        display_name="liu",
        site="taiyuan",
        clearance="general",
        permissions=(PERM_STAFFING_READ, PERM_STAFFING_WRITE),
        project_ids=project_ids,
        staffing_org_ok=True,
    )


def _ack_all(warnings: list, collision: dict) -> dict:
    resolutions = {
        w["code"]: "acknowledged"
        for w in warnings
        if w.get("code") and w["code"] != "name_kind_collision"
    }
    resolutions["name_kind_collision"] = collision
    return resolutions


def test_parse_emits_name_kind_collision_for_wang() -> None:
    result = parse_daily_report(_build_sample_xlsx(), filename="2515项目日报.xlsx")
    collide = [
        w
        for w in result.warnings
        if w.code == "name_kind_collision" and w.detail.get("person_name") == "王亮"
    ]
    assert collide
    kinds = {r.person_kind for r in result.rows if r.person_name == "王亮"}
    assert kinds == {KIND_FORMAL, KIND_CONTRACT}


@pytest.mark.asyncio
async def test_confirm_requires_collision_resolution_and_split(tmp_path) -> None:
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'c.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()
    project_id = "proj2515"

    async with factory() as session:
        session.add(
            Project(
                id=project_id,
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        result = parse_daily_report(_build_sample_xlsx(), filename="2515项目日报.xlsx")
        warnings = [asdict(w) for w in result.warnings]
        rows = [asdict(r) for r in result.rows]
        session.add(
            StaffingImportBatch(
                id="b1",
                tenant_id="autley",
                project_id=project_id,
                status="needs_review",
                filename="2515项目日报.xlsx",
                storage_key="x",
                uploaded_by=principal.user_id,
                warnings=warnings,
                parse_result={"rows": rows, "warnings": warnings},
                resolved_warnings={},
                match_info={"status": "matched", "project_id": project_id},
            )
        )
        await session.commit()

        with pytest.raises(AppError):
            await staffing_service.confirm_batch(
                session, principal=principal, batch_id="b1"
            )

        await staffing_service.review_batch(
            session,
            principal=principal,
            batch_id="b1",
            resolutions=_ack_all(warnings, {"王亮": {"action": "split"}}),
        )
        confirmed = await staffing_service.confirm_batch(
            session, principal=principal, batch_id="b1"
        )
        assert confirmed.status == "confirmed"

        facts = (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.batch_id == "b1",
                    StaffingAttendanceFact.person_name == "王亮",
                    StaffingAttendanceFact.status == "active",
                )
            )
        ).scalars().all()
        kinds = {f.person_kind for f in facts}
        assert KIND_FORMAL in kinds and KIND_CONTRACT in kinds

        summary = await staffing_service.project_summary(
            session, principal=principal, project_id=project_id
        )
        wang_rows = [p for p in summary["people"] if p["person_name"] == "王亮"]
        assert len(wang_rows) == 2
        assert summary["project_code"] == "2515"
        assert summary["person_day_total"] >= summary["person_count"]

        xlsx = await staffing_service.export_project_xlsx(
            session, principal=principal, project_id=project_id
        )
        assert xlsx[:2] == b"PK"

    await engine.dispose()


@pytest.mark.asyncio
async def test_project_summary_stints_and_export(tmp_path) -> None:
    """Gapped on-site dates yield multiple stints on summary + xlsx 汇总."""
    from datetime import date
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.entity import new_id

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'stints.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()
    project_id = "proj2515"

    async with factory() as session:
        session.add(
            Project(
                id=project_id,
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        batch_id = "bstint"
        session.add(
            StaffingImportBatch(
                id=batch_id,
                tenant_id="autley",
                project_id=project_id,
                status="confirmed",
                filename="stints.xlsx",
                storage_key="x",
                uploaded_by=principal.user_id,
                warnings=[],
                parse_result={"rows": []},
                resolved_warnings={},
                match_info={"status": "matched"},
            )
        )
        for d in [
            date(2026, 1, 11),
            date(2026, 1, 12),
            date(2026, 1, 13),
            date(2026, 3, 9),
            date(2026, 3, 10),
            date(2026, 3, 11),
        ]:
            session.add(
                StaffingAttendanceFact(
                    id=new_id(),
                    tenant_id="autley",
                    project_id=project_id,
                    batch_id=batch_id,
                    work_date=d,
                    person_name="张三",
                    person_kind=KIND_FORMAL,
                    status="active",
                    stage="电气安装" if d.month == 3 else "机械安装",
                )
            )
        await session.commit()

        summary = await staffing_service.project_summary(
            session, principal=principal, project_id=project_id
        )
        person = next(p for p in summary["people"] if p["person_name"] == "张三")
        assert person["days_on_site"] == 6
        assert person["stint_count"] == 2
        assert person["stints"][0]["entry_date"] == "2026-01-11"
        assert person["stints"][0]["exit_date"] == "2026-01-13"
        assert person["stints"][1]["entry_date"] == "2026-03-09"
        assert person["stints"][1]["exit_date"] == "2026-03-11"
        assert person["stages"] == [
            {"stage": "机械安装", "days": 3},
            {"stage": "电气安装", "days": 3},
        ]
        assert person["stints"][0]["stage_label"] == "机械安装"
        assert person["stints"][1]["stage_label"] == "电气安装"

        from app.chat.tools.staffing_orchestrator import _summary_card

        card = _summary_card(summary)
        card_person = next(p for p in card["people"] if p["person_name"] == "张三")
        assert card_person["stint_count"] == 2
        assert len(card_person["stints"]) == 2
        assert card_person["stages"][0]["stage"] == "机械安装"
        assert card_person["stints"][1]["stage_label"] == "电气安装"

        xlsx = await staffing_service.export_project_xlsx(
            session, principal=principal, project_id=project_id
        )
        wb = load_workbook(BytesIO(xlsx))
        assert wb.sheetnames == ["人员投入汇总表"]
        ws = wb["人员投入汇总表"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert headers == [
            "序号",
            "日期",
            "姓名",
            "项目号",
            "项目阶段",
            "部门",
            "职务",
            "项目名称",
        ]
        detail_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(detail_rows) == 6
        assert detail_rows[0][0] == 1
        assert detail_rows[0][2] == "张三"
        assert detail_rows[0][3] == 2515
        assert detail_rows[0][4] == "机械安装"
        assert detail_rows[0][7] == "2515"
        march = [r for r in detail_rows if r[4] == "电气安装"]
        assert len(march) == 3
        assert all(r[2] == "张三" for r in march)
        # Dates are contiguous in two blocks — reconstructible as stints from person-days.
        dates = [r[1].date() if hasattr(r[1], "date") else r[1] for r in detail_rows]
        assert dates[0].isoformat() == "2026-01-11"
        assert dates[-1].isoformat() == "2026-03-11"

    await engine.dispose()


@pytest.mark.asyncio
async def test_export_roster_fills_department_title(tmp_path, monkeypatch) -> None:
    """Roster 姓名 match fills 部门/职务; contract fallback is 机电服务处."""
    from datetime import date
    from io import BytesIO
    from pathlib import Path

    from openpyxl import Workbook, load_workbook

    from app.core.config import get_settings
    from app.models.entity import new_id
    from app.staffing import roster as roster_mod

    roster_path = Path(tmp_path) / "roster.xlsx"
    rwb = Workbook()
    rws = rwb.active
    rws.append(["序列", "姓名", "部门", "工种", "备注", "待进场"])
    rws.append([1, "王亮", "项目管理部", "项目经理", "", ""])
    rwb.save(roster_path)
    monkeypatch.setenv("KB_STAFFING_ROSTER_PATH", str(roster_path))
    get_settings.cache_clear()
    roster_mod._cached_roster.cache_clear()

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'roster_export.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal(project_ids=("proj_roster",))
    project_id = "proj_roster"

    async with factory() as session:
        session.add(
            Project(
                id=project_id,
                tenant_id="autley",
                code="2532",
                name="岩心库二期项目",
                status="active",
            )
        )
        batch_id = "broster"
        session.add(
            StaffingImportBatch(
                id=batch_id,
                tenant_id="autley",
                project_id=project_id,
                status="confirmed",
                filename="r.xlsx",
                storage_key="x",
                uploaded_by=principal.user_id,
                warnings=[],
                parse_result={"rows": []},
                resolved_warnings={},
                match_info={"status": "matched"},
            )
        )
        session.add(
            StaffingAttendanceFact(
                id=new_id(),
                tenant_id="autley",
                project_id=project_id,
                batch_id=batch_id,
                work_date=date(2026, 1, 1),
                person_name="王亮",
                person_kind=KIND_FORMAL,
                status="active",
                stage="软件调试",
            )
        )
        session.add(
            StaffingAttendanceFact(
                id=new_id(),
                tenant_id="autley",
                project_id=project_id,
                batch_id=batch_id,
                work_date=date(2026, 1, 1),
                person_name="史承志",
                person_kind=KIND_CONTRACT,
                status="active",
                stage="机械安装",
            )
        )
        await session.commit()

        xlsx = await staffing_service.export_project_xlsx(
            session, principal=principal, project_id=project_id
        )
        ws = load_workbook(BytesIO(xlsx)).active
        by_name = {r[2]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        assert by_name["王亮"][5] == "项目管理部"
        assert by_name["王亮"][6] == "项目经理"
        assert by_name["史承志"][5] == "机电服务处"
        assert not by_name["史承志"][6]

    get_settings.cache_clear()
    roster_mod._cached_roster.cache_clear()
    await engine.dispose()


@pytest.mark.asyncio
async def test_confirm_merge_keeps_one_kind(tmp_path) -> None:
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'm.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()
    project_id = "proj2515"

    async with factory() as session:
        session.add(
            Project(
                id=project_id,
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        result = parse_daily_report(_build_sample_xlsx(), filename="2515项目日报.xlsx")
        warnings = [asdict(w) for w in result.warnings]
        rows = [asdict(r) for r in result.rows]
        session.add(
            StaffingImportBatch(
                id="b2",
                tenant_id="autley",
                project_id=project_id,
                status="needs_review",
                filename="2515项目日报.xlsx",
                storage_key="x",
                uploaded_by=principal.user_id,
                warnings=warnings,
                parse_result={"rows": rows},
                resolved_warnings={},
                match_info={"status": "matched"},
            )
        )
        await session.commit()
        await staffing_service.review_batch(
            session,
            principal=principal,
            batch_id="b2",
            resolutions=_ack_all(
                warnings, {"王亮": {"action": "merge", "keep_kind": KIND_FORMAL}}
            ),
        )
        await staffing_service.confirm_batch(
            session, principal=principal, batch_id="b2"
        )
        facts = (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.batch_id == "b2",
                    StaffingAttendanceFact.person_name == "王亮",
                    StaffingAttendanceFact.status == "active",
                )
            )
        ).scalars().all()
        assert facts
        assert all(f.person_kind == KIND_FORMAL for f in facts)

    await engine.dispose()
