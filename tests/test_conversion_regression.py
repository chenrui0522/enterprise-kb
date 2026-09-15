"""Per-format conversion regression: every supported input still yields text.

These tests drive the ingestion pipeline's conversion step (routing + report)
so a change in triage, fallback order or a converter cannot silently drop
content for one of the corpus formats.
"""

from types import SimpleNamespace

import pytest

from app.core.config import Settings
from app.ingestion.converters import ConversionResult
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.storage import FileDocumentStorage


def _settings(directory, **overrides) -> Settings:
    values = {
        "mineru_url": "",
        "docling_url": "",
        "document_storage_dir": str(directory),
    }
    values.update(overrides)
    return Settings(**values)


def _pipeline(tmp_path, settings) -> IngestionPipeline:
    return IngestionPipeline(
        store=object(),
        embedder=object(),
        storage=FileDocumentStorage(str(tmp_path / "docs")),
        settings=settings,
    )


def _run(pipeline, path, filename: str):
    document = SimpleNamespace(filename=filename)
    version = SimpleNamespace(parse_mode="auto")
    return pipeline, document, version, path


def _write_text_pdf(path, text: str, pages: int = 1) -> None:
    import fitz

    document = fitz.open()
    for index in range(pages):
        page = document.new_page()
        page.insert_text((72, 72), f"{text} (page {index + 1})", fontsize=12)
    document.save(path)
    document.close()


def _write_blank_pdf(path, pages: int = 1) -> None:
    import fitz

    document = fitz.open()
    for _ in range(pages):
        document.new_page()
    document.save(path)
    document.close()


def _write_table_pdf(path) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Quarterly Metrics", fontsize=18)
    left, top, col_w, row_h, cols, rows = 72, 140, 90, 20, 2, 3
    for index in range(rows + 1):
        page.draw_line(fitz.Point(left, top + index * row_h), fitz.Point(left + cols * col_w, top + index * row_h))
    for index in range(cols + 1):
        page.draw_line(fitz.Point(left + index * col_w, top), fitz.Point(left + index * col_w, top + rows * row_h))
    cells = [["Indicator", "Weight"], ["Completion", "40%"], ["Quality", "60%"]]
    for row in range(rows):
        for column in range(cols):
            page.insert_text(
                fitz.Point(left + 6 + column * col_w, top + 14 + row * row_h),
                cells[row][column],
                fontsize=11,
            )
    document.save(path)
    document.close()


def _write_docx(path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("绩效管理办法", level=1)
    doc.add_paragraph("本办法适用于全体员工。")
    doc.save(path)


def _write_xlsx(path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "花名册"
    sheet.append(["姓名", "部门"])
    sheet.append(["张三", "研发部"])
    workbook.save(path)


def _write_pptx(path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "季度复盘"
    slide.placeholders[1].text = "达成率 112%"
    presentation.save(path)


@pytest.mark.asyncio
async def test_text_pdf_conversion(tmp_path) -> None:
    pdf = tmp_path / "manual.pdf"
    _write_text_pdf(pdf, "Warranty covers twelve months for ZB-100.")
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="manual.pdf"), SimpleNamespace(parse_mode="auto"), pdf
    )
    assert "Warranty covers twelve months" in result.markdown
    assert "<!-- PAGE:1 -->" in result.markdown
    assert pipeline._last_conversion_meta["label"] == "text"


@pytest.mark.asyncio
async def test_multi_page_text_pdf_keeps_every_page(tmp_path) -> None:
    pdf = tmp_path / "manual.pdf"
    _write_text_pdf(pdf, "Chapter content", pages=3)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="manual.pdf"), SimpleNamespace(parse_mode="auto"), pdf
    )
    for page in (1, 2, 3):
        assert f"<!-- PAGE:{page} -->" in result.markdown
    assert result.page_count == 3


@pytest.mark.asyncio
async def test_scanned_pdf_goes_through_ocr(tmp_path, monkeypatch) -> None:
    from app.ingestion.mineru import MinerUConverter

    pdf = tmp_path / "scan.pdf"
    _write_blank_pdf(pdf, pages=2)

    async def fake_ocr(self, file_path):  # noqa: ARG001
        return ConversionResult(markdown="# 扫描件\n\nOCR 识别出的正文。", page_count=None)

    monkeypatch.setattr(MinerUConverter, "convert_async", fake_ocr)
    settings = _settings(tmp_path / "docs", mineru_url="http://mineru.test")
    pipeline = _pipeline(tmp_path, settings)
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="scan.pdf"), SimpleNamespace(parse_mode="auto"), pdf
    )
    assert "OCR 识别出的正文" in result.markdown
    assert pipeline._last_conversion_meta["label"] == "ocr"
    assert pipeline._last_triage["kind"] == "pdf_scanned"


@pytest.mark.asyncio
async def test_table_pdf_conversion_keeps_table(tmp_path) -> None:
    pdf = tmp_path / "table.pdf"
    _write_table_pdf(pdf)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="table.pdf"), SimpleNamespace(parse_mode="auto"), pdf
    )
    assert "Completion" in result.markdown
    assert "|" in result.markdown
    assert pipeline._last_triage["has_tables"] is True


@pytest.mark.asyncio
async def test_docx_conversion(tmp_path) -> None:
    target = tmp_path / "policy.docx"
    _write_docx(target)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="policy.docx"), SimpleNamespace(parse_mode="auto"), target
    )
    assert "绩效管理办法" in result.markdown
    assert pipeline._last_conversion_meta["label"] == "text"


@pytest.mark.asyncio
async def test_xlsx_conversion_is_semantic(tmp_path) -> None:
    target = tmp_path / "roster.xlsx"
    _write_xlsx(target)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="roster.xlsx"), SimpleNamespace(parse_mode="auto"), target
    )
    assert "姓名=张三；部门=研发部" in result.markdown
    assert pipeline._last_conversion_meta["label"] == "xlsx"
    assert len(result.tables) == 1


@pytest.mark.asyncio
async def test_pptx_conversion(tmp_path) -> None:
    target = tmp_path / "deck.pptx"
    _write_pptx(target)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    result = await pipeline._convert_to_markdown(
        SimpleNamespace(filename="deck.pptx"), SimpleNamespace(parse_mode="auto"), target
    )
    assert "季度复盘" in result.markdown
    assert pipeline._last_conversion_meta["label"] == "markitdown"


@pytest.mark.asyncio
async def test_forced_ocr_requires_service(tmp_path) -> None:
    from app.core.errors import AppError

    pdf = tmp_path / "scan.pdf"
    _write_blank_pdf(pdf)
    pipeline = _pipeline(tmp_path, _settings(tmp_path / "docs"))
    with pytest.raises(AppError, match="KB_MINERU_URL"):
        await pipeline._convert_to_markdown(
            SimpleNamespace(filename="scan.pdf"), SimpleNamespace(parse_mode="ocr"), pdf
        )