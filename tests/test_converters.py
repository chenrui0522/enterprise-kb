from pathlib import Path

import pytest

from app.core.errors import AppError
from app.ingestion.converters import (
    DocxMarkdownConverter,
    MarkItDownConverter,
    PDFMarkdownConverter,
    TextMarkdownConverter,
    converter_for,
)
from app.ingestion.chunker import ChunkContext, chunk_markdown


def _write_pdf(path: Path, text: str) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text, fontsize=12)
    document.save(path)
    document.close()


def test_pdf_converter_emits_page_marker(tmp_path) -> None:
    target = tmp_path / "manual.pdf"
    _write_pdf(target, "Warranty covers twelve months for ZB-100.")
    result = PDFMarkdownConverter(extract_tables=False).convert(str(target))
    assert result.page_count == 1
    assert "<!-- PAGE:1 -->" in result.markdown
    assert "ZB-100" in result.markdown


def test_docx_converter_keeps_headings_and_table(tmp_path) -> None:
    from docx import Document

    target = tmp_path / "policy.docx"
    doc = Document()
    doc.add_heading("绩效考核办法", level=1)
    doc.add_paragraph("本制度适用于全体员工。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "权重"
    table.cell(1, 0).text = "完成率"
    table.cell(1, 1).text = "40%"
    doc.save(target)

    result = DocxMarkdownConverter().convert(str(target))
    assert "# 绩效考核办法" in result.markdown
    assert "完成率" in result.markdown
    assert result.page_count is None


def test_text_converter_passthrough(tmp_path) -> None:
    target = tmp_path / "note.md"
    target.write_text("## 备份策略\n\n每日备份一次。", encoding="utf-8")
    result = TextMarkdownConverter().convert(str(target))
    assert "## 备份策略" in result.markdown


def test_unknown_format_rejected() -> None:
    with pytest.raises(AppError):
        converter_for("image.png")


def test_markitdown_pptx_converter(tmp_path) -> None:
    from pptx import Presentation

    target = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "季度销售复盘"
    slide.placeholders[1].text = "本季度达成率 112%"
    prs.save(target)

    result = MarkItDownConverter().convert(str(target))
    assert "季度销售复盘" in result.markdown
    assert "112%" in result.markdown


def test_markitdown_xlsx_converter(tmp_path) -> None:
    from openpyxl import Workbook

    target = tmp_path / "spreadsheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "部门", "状态"])
    ws.append(["张三", "研发部", "在职"])
    wb.save(target)

    result = MarkItDownConverter().convert(str(target))
    assert "张三" in result.markdown
    assert "研发部" in result.markdown


def test_markitdown_html_converter(tmp_path) -> None:
    target = tmp_path / "page.html"
    target.write_text(
        "<html><body><h1>产品介绍</h1><p>这是一段说明文字。</p></body></html>",
        encoding="utf-8",
    )
    result = MarkItDownConverter().convert(str(target))
    assert "产品介绍" in result.markdown


def test_markitdown_csv_converter(tmp_path) -> None:
    target = tmp_path / "roster.csv"
    target.write_text("姓名,部门\n李四,财务部\n王五,人事部\n", encoding="utf-8")
    result = MarkItDownConverter().convert(str(target))
    assert "李四" in result.markdown
    assert "财务部" in result.markdown


def test_markitdown_json_converter(tmp_path) -> None:
    target = tmp_path / "config.json"
    target.write_text(
        '{"title": "知识库配置", "items": [{"name": "产品手册", "pages": 100}]}',
        encoding="utf-8",
    )
    result = MarkItDownConverter().convert(str(target))
    assert "知识库配置" in result.markdown


def test_chunk_markdown_respects_page_and_section(tmp_path) -> None:
    markdown = "<!-- PAGE:1 -->\n\n# 保修政策\n\n打印机保修十二个月。\n\n<!-- PAGE:2 -->\n\n# 退换货\n\n七天无理由。"
    context = ChunkContext(doc_id="doc1", version_id="v1", tenant_id="t1", title="手册")
    drafts = chunk_markdown(markdown, context, chunk_size=500, overlap=20)
    pages = {draft.page for draft in drafts}
    assert pages == {1, 2}
    assert any(draft.section == "保修政策" for draft in drafts)


def _write_table_pdf(path: Path) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Warranty Manual", fontsize=20)
    page.insert_text(
        (72, 104),
        "This section explains the warranty policy in detail for all supported products.",
        fontsize=11,
    )
    left, top, col_w, row_h, cols, rows = 72, 140, 100, 20, 2, 3
    for index in range(rows + 1):
        y = top + index * row_h
        page.draw_line(fitz.Point(left, y), fitz.Point(left + cols * col_w, y))
    for index in range(cols + 1):
        x = left + index * col_w
        page.draw_line(fitz.Point(x, top), fitz.Point(x, top + rows * row_h))
    cells = [["Indicator", "Weight"], ["Completion", "40%"], ["Quality", "60%"]]
    for row in range(rows):
        for column in range(cols):
            page.insert_text(
                fitz.Point(left + 6 + column * col_w, top + 14 + row * row_h),
                cells[row][column],
                fontsize=11,
            )
    document.save(path)
    document.close()


def test_pdf_converter_detects_heading_and_table_blocks(tmp_path) -> None:
    target = tmp_path / "table.pdf"
    _write_table_pdf(target)
    result = PDFMarkdownConverter().convert(str(target))
    assert result.page_count == 1
    structure = result.structure
    assert structure is not None
    headings = [block for block in structure.blocks if block.type == "heading"]
    assert headings and headings[0].text == "Warranty Manual"
    tables = [block for block in structure.blocks if block.type == "table"]
    assert tables
    table = tables[0]
    assert table.header == ["Indicator", "Weight"]
    assert table.rows[0] == ["Completion", "40%"]
    assert all(
        "Indicator" not in block.text
        for block in structure.blocks
        if block.type == "paragraph"
    )


def test_docx_converter_structure_blocks_keep_order(tmp_path) -> None:
    from docx import Document

    target = tmp_path / "policy.docx"
    doc = Document()
    doc.add_heading("绩效考核办法", level=1)
    doc.add_paragraph("本办法适用于全体员工。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "权重"
    table.cell(1, 0).text = "完成率"
    table.cell(1, 1).text = "40%"
    doc.add_paragraph("后续说明。")
    doc.save(target)

    result = DocxMarkdownConverter().convert(str(target))
    structure = result.structure
    assert structure is not None
    assert [block.type for block in structure.blocks] == ["heading", "paragraph", "table", "paragraph"]
    assert structure.blocks[0].level == 1
    table_block = structure.blocks[2]
    assert table_block.header == ["指标", "权重"]
    assert table_block.rows == [["完成率", "40%"]]


def test_docx_converter_detects_localized_heading_style(tmp_path) -> None:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    target = tmp_path / "localized.docx"
    doc = Document()
    style = doc.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = doc.add_paragraph("本地化标题", style=style)
    assert paragraph.text == "本地化标题"
    doc.save(target)

    result = DocxMarkdownConverter().convert(str(target))
    structure = result.structure
    assert structure is not None
    assert structure.blocks[0].type == "heading"
    assert structure.blocks[0].level == 1


def test_text_converter_emits_list_and_table_blocks(tmp_path) -> None:
    target = tmp_path / "policy.md"
    target.write_text(
        "## 备份策略\n\n- 每日备份\n- 每周校验\n\n| 项目 | 频率 |\n|---|---|\n| 全量 | 每日 |\n",
        encoding="utf-8",
    )
    result = TextMarkdownConverter().convert(str(target))
    structure = result.structure
    assert structure is not None
    assert any(block.type == "list" and block.items == ["每日备份", "每周校验"] for block in structure.blocks)
    table = next(block for block in structure.blocks if block.type == "table")
    assert table.header == ["项目", "频率"]
    assert table.rows == [["全量", "每日"]]