import pytest
from openpyxl import Workbook

from app.core.errors import AppError
from app.ingestion.structure import BLOCK_HEADING, BLOCK_PARAGRAPH, BLOCK_TABLE
from app.ingestion.tabular import XlsxSemanticConverter, _format_value


def _workbook(path) -> Workbook:
    workbook = Workbook()
    return workbook


def test_multi_sheet_produces_one_section_per_sheet(tmp_path) -> None:
    target = tmp_path / "人员台账.xlsx"
    workbook = Workbook()
    roster = workbook.active
    roster.title = "花名册"
    roster.append(["姓名", "部门", "状态"])
    roster.append(["张三", "研发部", "在职"])
    second = workbook.create_sheet("部门预算")
    second.append(["部门", "预算"])
    second.append(["研发部", "100"])
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    headings = [block.text for block in result.structure.blocks if block.type == BLOCK_HEADING]
    assert any("花名册" in text for text in headings)
    assert any("部门预算" in text for text in headings)
    assert [table.sheet for table in result.tables] == ["花名册", "部门预算"]


def test_header_inference_and_semantic_rows(tmp_path) -> None:
    target = tmp_path / "roster.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "花名册"
    sheet.append(["姓名", "部门", "状态"])
    sheet.append(["张三", "研发部", "在职"])
    sheet.append(["李四", "财务部", "试用"])
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    assert "字段：姓名、部门、状态" in result.markdown
    assert "姓名=张三；部门=研发部；状态=在职" in result.markdown
    assert "姓名=李四；部门=财务部；状态=试用" in result.markdown
    table = result.tables[0]
    assert table.header == ["姓名", "部门", "状态"]
    assert table.rows == [["张三", "研发部", "在职"], ["李四", "财务部", "试用"]]


def test_ambiguous_first_row_falls_back_to_generated_headers(tmp_path) -> None:
    target = tmp_path / "numbers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["1", "2"])
    sheet.append(["3", "4"])
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    assert "字段：列1、列2" in result.markdown
    assert "列1=1；列2=2" in result.markdown
    assert "列1=3；列2=4" in result.markdown


def test_table_grid_is_metadata_only_not_indexed(tmp_path) -> None:
    target = tmp_path / "roster.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "花名册"
    sheet.append(["姓名", "部门"])
    sheet.append(["张三", "研发部"])
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    types = [block.type for block in result.structure.blocks]
    assert BLOCK_TABLE not in types
    assert types.count(BLOCK_PARAGRAPH) == 1
    # The raw grid is kept as metadata instead of a second indexed copy.
    assert "|" not in result.markdown
    assert result.tables[0].rows == [["张三", "研发部"]]


def test_merged_cells_are_filled_down(tmp_path) -> None:
    target = tmp_path / "regions.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "区域"
    sheet.append(["大区", "城市"])
    sheet.merge_cells("A2:A3")
    sheet["A2"] = "华东"
    sheet["B2"] = "上海"
    sheet["B3"] = "杭州"
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    assert "大区=华东；城市=上海" in result.markdown
    assert "大区=华东；城市=杭州" in result.markdown


def test_formula_value_falls_back_to_formula_text(tmp_path) -> None:
    target = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet.append(["产品", "单价"])
    sheet.append(["A", "=1+1"])
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    assert "=1+1" in result.markdown


def test_format_value_helpers() -> None:
    assert _format_value(None, "=SUM(A1:A2)") == "=SUM(A1:A2)"
    assert _format_value(40.0) == "40"
    assert _format_value(0.5) == "0.5"
    assert _format_value(True) == "是"
    import datetime as dt

    assert _format_value(dt.date(2026, 3, 1)) == "2026-03-01"


def test_long_table_is_grouped_with_repeated_header(tmp_path) -> None:
    target = tmp_path / "long.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["序号", "金额"])
    for index in range(1, 61):
        sheet.append([str(index), str(index * 10)])
    workbook.save(target)

    converter = XlsxSemanticConverter(rows_per_group=25)
    result = converter.convert(str(target))
    assert result.markdown.count("字段：序号、金额") == 3
    assert "序号=1；金额=10" in result.markdown
    assert "序号=60；金额=600" in result.markdown
    # every row appears exactly once
    assert result.markdown.count("序号=30；金额=300") == 1


def test_hidden_and_empty_sheets_are_skipped(tmp_path) -> None:
    target = tmp_path / "hidden.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "可见"
    visible.append(["名称", "值"])
    visible.append(["甲", "1"])
    hidden = workbook.create_sheet("隐藏")
    hidden.append(["名称", "值"])
    hidden.append(["乙", "2"])
    hidden.sheet_state = "hidden"
    workbook.create_sheet("空表")
    workbook.save(target)

    result = XlsxSemanticConverter().convert(str(target))
    assert [table.sheet for table in result.tables] == ["可见"]
    assert "乙" not in result.markdown


def test_empty_workbook_is_rejected(tmp_path) -> None:
    target = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(target)
    with pytest.raises(AppError):
        XlsxSemanticConverter().convert(str(target))


def test_corrupt_xlsx_reports_clear_error(tmp_path) -> None:
    target = tmp_path / "broken.xlsx"
    target.write_bytes(b"not an excel file")
    with pytest.raises(AppError):
        XlsxSemanticConverter().convert(str(target))