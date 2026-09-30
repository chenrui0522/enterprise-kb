"""Chat attachments API + staffing orchestrator from chat."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.chat.tools.attachments import save_chat_xlsx
from app.chat.tools.staffing_orchestrator import (
    apply_tool_action,
    clear_tool_state,
    handle_staffing_message,
    update_staffing_state,
)
from app.core.config import get_settings
from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import (
    PERM_CHAT_USE,
    PERM_STAFFING_READ,
    PERM_STAFFING_WRITE,
)
from app.identity.deps import get_current_principal, require_permission, tenant_from_principal
from app.identity.principal import Principal
from app.models.entity import Conversation
from app.models.identity import Project
from app.models.staffing import StaffingAttendanceFact
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401
import app.models.staffing  # noqa: F401


def _build_sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "项目日报表"
    ws["A3"] = "项目名称；2515JS壹号智能西安项目    项目经理：王亮"
    ws["A4"] = "序号"
    ws["B4"] = "日期"
    ws["C4"] = "现场施工人员/人数"
    ws["D4"] = "当前阶段"
    ws["J4"] = "现场人员（人）"
    ws["J5"] = "机械安装"
    ws["K6"] = "我司人员\n姓名"
    ws["K7"] = "姓名"
    ws["A8"] = 1
    ws["B8"] = date(2026, 1, 11)
    ws["C8"] = "【公司人员】：0人； 【外包人员】：机械安装3人"
    ws["D8"] = "机械安装"
    ws["K8"] = "赵鑫磊、王亮"
    ws["A9"] = 2
    ws["B9"] = date(2026, 3, 9)
    ws["C9"] = "【公司人员】：1人 王亮 【外包人员】："
    ws["D9"] = "电气安装"
    ws["K9"] = "赵鑫磊、王亮"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _principal(*, perms: tuple[str, ...], project_ids: tuple[str, ...] = ()) -> Principal:
    return Principal(
        user_id="u-staff",
        tenant_id="autley",
        username="liu",
        display_name="刘珊",
        site="taiyuan",
        clearance="general",
        permissions=perms,
        project_ids=project_ids,
        staffing_org_ok=True,
    )


@pytest.fixture
def storage_root(tmp_path, monkeypatch):
    root = tmp_path / "storage"
    root.mkdir()
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(root))
    get_settings.cache_clear()
    yield root
    get_settings.cache_clear()


def test_save_chat_xlsx_rejects_non_excel(storage_root) -> None:
    with pytest.raises(AppError) as ei:
        save_chat_xlsx(user_id="u1", filename="notes.txt", data=b"hello")
    assert ei.value.status_code == 400
    with pytest.raises(AppError) as word:
        save_chat_xlsx(user_id="u1", filename="日报.docx", data=b"PK\x03\x04word")
    assert word.value.status_code == 400
    assert "PDF" in word.value.message


def test_save_chat_xlsx_accepts_pdf(storage_root) -> None:
    meta = save_chat_xlsx(user_id="u1", filename="2515项目日报.pdf", data=b"%PDF-1.4 sample")
    assert meta["filename"] == "2515项目日报.pdf"
    stored = Path(meta["storage_key"])
    if not stored.is_absolute():
        stored = storage_root / stored
    assert stored.read_bytes().startswith(b"%PDF")


def test_chat_attachments_forbidden_without_chat_perm(storage_root) -> None:
    from app.api.routers import chat as chat_module

    principal = _principal(perms=())
    app = FastAPI()

    from fastapi.responses import JSONResponse

    @app.exception_handler(AppError)
    async def _app_error(_request, exc: AppError):
        return JSONResponse(status_code=exc.status_code, content={"detail": exc.message})

    app.include_router(chat_module.router, prefix="/api/v1")
    app.dependency_overrides[get_current_principal] = lambda: principal
    app.dependency_overrides[tenant_from_principal] = lambda: principal.tenant_id

    def _deny():
        raise AppError("缺少所需权限", status_code=403)

    app.dependency_overrides[require_permission(PERM_CHAT_USE)] = _deny
    client = TestClient(app)
    response = client.post(
        "/api/v1/chat/attachments",
        files={
            "file": (
                "2515项目日报.xlsx",
                _build_sample_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 403


def test_chat_attachments_upload_ok(storage_root) -> None:
    from app.api.routers import chat as chat_module

    principal = _principal(perms=(PERM_CHAT_USE,))
    app = FastAPI()
    app.include_router(chat_module.router, prefix="/api/v1")
    app.dependency_overrides[get_current_principal] = lambda: principal
    app.dependency_overrides[tenant_from_principal] = lambda: principal.tenant_id
    app.dependency_overrides[require_permission(PERM_CHAT_USE)] = lambda: principal
    client = TestClient(app)
    response = client.post(
        "/api/v1/chat/attachments",
        files={
            "file": (
                "2515项目日报.xlsx",
                _build_sample_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201
    body = response.json()
    assert body["filename"] == "2515项目日报.xlsx"
    assert body["id"]


def test_tool_state_clears_on_new_conversation() -> None:
    conv = Conversation(tenant_id="autley", title="t", created_by="u1")
    update_staffing_state(conv, batch_id="b1", project_id="p1")
    assert conv.tool_state and conv.tool_state.get("active_tool") == "staffing"
    clear_tool_state(conv)
    assert conv.tool_state is None


@pytest.mark.asyncio
async def test_import_confirm_summary_and_void_requires_confirm(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'staff.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)

    project_id = "proj2515"
    principal = _principal(
        perms=(PERM_CHAT_USE, PERM_STAFFING_READ, PERM_STAFFING_WRITE),
        project_ids=(project_id,),
    )

    async with factory() as session:
        session.add(
            Project(
                id=project_id,
                tenant_id="autley",
                code="2515",
                name="2515JS壹号",
                status="active",
            )
        )
        conversation = Conversation(
            tenant_id="autley", title="人员投入", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)

        data = _build_sample_xlsx()
        result = await handle_staffing_message(
            session,
            principal=principal,
            conversation=conversation,
            message="导入项目日报",
            attachment={"filename": "2515项目日报.xlsx", "data": data},
        )
        assert result["card"]["type"] == "staffing_batch"
        batch_id = result["card"]["batch_id"]
        assert conversation.tool_state["staffing"]["batch_id"] == batch_id

        # Text confirm must not mutate via message alone
        with pytest.raises(AppError) as ei:
            await handle_staffing_message(
                session,
                principal=principal,
                conversation=conversation,
                message="确认入库",
                attachment=None,
            )
        assert ei.value.status_code == 400

        collision = {}
        for c in result["card"].get("collisions") or []:
            if c.get("person_name"):
                collision[c["person_name"]] = {"action": "split"}
        confirmed = await apply_tool_action(
            session,
            principal=principal,
            conversation=conversation,
            action="ack_and_confirm",
            payload={
                "batch_id": batch_id,
                "project_id": project_id,
                "name_kind_collision": collision or None,
            },
        )
        assert confirmed["card"]["type"] == "staffing_summary"
        assert confirmed["card"]["person_count"] >= 1
        assert "person_day_total" in confirmed["card"]
        first = confirmed["card"]["people"][0] if confirmed["card"]["people"] else {}
        assert "dates" in first
        assert "stint_count" in first
        assert "stints" in first
        assert first["stint_count"] == len(first["stints"])
        fact_count = await session.scalar(
            select(func.count()).select_from(StaffingAttendanceFact).where(
                StaffingAttendanceFact.project_id == project_id,
                StaffingAttendanceFact.status == "active",
            )
        )
        assert fact_count and fact_count > 0

        with pytest.raises(AppError) as ei2:
            await apply_tool_action(
                session,
                principal=principal,
                conversation=conversation,
                action="void_person",
                payload={"person_names": ["王亮"], "confirmed": False},
            )
        assert "二次确认" in ei2.value.message

        still = await session.scalar(
            select(func.count()).select_from(StaffingAttendanceFact).where(
                StaffingAttendanceFact.project_id == project_id,
                StaffingAttendanceFact.status == "active",
            )
        )
        assert still == fact_count

        # Staffing path returns summary card — not RAG citations
        query = await handle_staffing_message(
            session,
            principal=principal,
            conversation=conversation,
            message="有谁在场几天",
            attachment=None,
        )
        assert query["card"]["type"] == "staffing_summary"
        assert "citations" not in query

    await engine.dispose()
    get_settings.cache_clear()
