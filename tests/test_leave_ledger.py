"""Unit tests for leave-ledger classify / parse / rules / export."""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.core.errors import AppError
from app.identity.constants import (
    PERM_LEAVE_LEDGER_READ,
    PERM_LEAVE_LEDGER_WRITE,
    ROLE_PERMISSIONS,
)
from app.leave_ledger.classify import (
    classify_by_filename,
    classify_source,
    looks_like_ooxml_xlsx,
)
from app.leave_ledger.export import build_ledger_xlsx
from app.leave_ledger.parse import (
    LeaveRecord,
    OvertimeDay,
    ParsedSources,
    PunchRecord,
    TravelRecord,
    _parse_date,
    parse_leave,
    parse_overtime,
)
from app.leave_ledger.rules import compute_ledger
from app.leave_ledger.service import _open_blocking_warnings, _ot_resolutions_from_resolved


def _xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_leave_ledger_perms_on_roles() -> None:
    assert PERM_LEAVE_LEDGER_READ in ROLE_PERMISSIONS["admin"]
    assert PERM_LEAVE_LEDGER_WRITE in ROLE_PERMISSIONS["admin"]
    assert PERM_LEAVE_LEDGER_READ in ROLE_PERMISSIONS["editor"]
    assert PERM_LEAVE_LEDGER_WRITE in ROLE_PERMISSIONS["editor"]
    assert PERM_LEAVE_LEDGER_READ in ROLE_PERMISSIONS["reader"]
    assert PERM_LEAVE_LEDGER_WRITE not in ROLE_PERMISSIONS["reader"]


def test_classify_by_filename() -> None:
    assert classify_by_filename("【审批】出差申请（太原公司）.xlsx") == "travel"
    assert classify_by_filename("【审批】加班申请（法定假日）.xlsx") == "overtime"
    assert classify_by_filename("【审批】请假申请（太原公司）.xlsx") == "leave"
    assert classify_by_filename("上下班打卡_日报_202606.xlsx") == "punch"


def test_classify_rejects_non_xlsx() -> None:
    assert not looks_like_ooxml_xlsx(b"not-zip")
    with pytest.raises(AppError):
        classify_source("unknown.csv", b"hello")


def test_classify_by_headers_overtime() -> None:
    data = _xlsx_bytes(
        ["审批编号", "申请人", "加班事由", "加班明细", "加班时长", "当前审批状态"],
        [["1", "甲", "端午", "2026/6/19 10小时", "10小时", "已通过"]],
    )
    assert classify_source("export.xlsx", data) == "overtime"


def test_parse_overtime_splits_detail_days() -> None:
    data = _xlsx_bytes(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "开始时间",
            "结束时间",
            "加班明细",
            "当前审批状态",
        ],
        [
            [
                "1",
                "冯江伟",
                "产品中心/软件部",
                "2026/6/19 08:00",
                "2026/6/21 18:00",
                "2026/6/19 10小时\n2026/6/20 10小时\n2026/6/21 10小时",
                "已通过",
            ]
        ],
    )
    days = parse_overtime(data)
    assert len(days) == 3
    assert [d.work_date for d in days] == [
        date(2026, 6, 19),
        date(2026, 6, 20),
        date(2026, 6, 21),
    ]


def test_project_travel_30_days_gives_2() -> None:
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="甲",
                department="产品中心/软件部",
                center="产品中心",
                category="项目类出差（入场）",
                start=date(2026, 5, 1),
                end=date(2026, 5, 30),
                status="已通过",
            )
        ],
        punch=[
            PunchRecord(
                person_name="甲",
                work_date=date(2026, 5, d),
                punch_time=datetime(2026, 5, d, 9, 0),
                status="正常",
                address="外地现场",
            )
            for d in range(1, 31)
        ],
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原", "公司"],
        holidays=set(),
    )
    assert len(result.rows) == 1
    assert result.rows[0].travel_comp_days == 2.0


def test_project_travel_under_30_no_credit() -> None:
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="乙",
                department="产品中心",
                center="产品中心",
                category="项目类出差（入场）",
                start=date(2026, 5, 1),
                end=date(2026, 5, 20),
                status="已通过",
            )
        ]
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    assert result.rows[0].travel_comp_days == 0.0


def test_hq_punch_truncates_travel() -> None:
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="丙",
                department="产品中心",
                center="产品中心",
                category="项目类出差（入场）",
                start=date(2026, 5, 1),
                end=date(2026, 5, 30),
                status="已通过",
            )
        ],
        punch=[
            PunchRecord(
                person_name="丙",
                work_date=date(2026, 5, 28),
                punch_time=datetime(2026, 5, 28, 9, 0),
                status="正常",
                address="太原市小店区公司园区",
            )
        ],
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原", "公司"],
        holidays=set(),
    )
    assert result.rows[0].actual_travel_end == "2026-05-28"
    # 1..28 = 28 days < 30
    assert result.rows[0].travel_comp_days == 0.0


def test_non_project_follows_ot_with_punch() -> None:
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="丁",
                department="产品中心",
                center="产品中心",
                category="非项目类出差",
                start=date(2026, 6, 1),
                end=date(2026, 6, 10),
                status="已通过",
            )
        ],
        overtime=[
            OvertimeDay(
                applicant="丁",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 19),
                hours=10,
                status="已通过",
            )
        ],
        punch=[
            PunchRecord(
                person_name="丁",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 8, 0),
                status="正常",
                address="现场",
            ),
            PunchRecord(
                person_name="丁",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 18, 0),
                status="正常",
                address="现场",
            ),
        ],
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays={date(2026, 6, 19)},
    )
    assert result.rows[0].travel_comp_days == 0.0
    assert result.rows[0].ot_comp_days == 1.0


def test_half_day_ot_warns() -> None:
    parsed = ParsedSources(
        overtime=[
            OvertimeDay(
                applicant="戊",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 19),
                hours=10,
                status="已通过",
            )
        ],
        punch=[
            PunchRecord(
                person_name="戊",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 8, 0),
                status="正常",
                address="现场",
            ),
            PunchRecord(
                person_name="戊",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 10, 0),
                status="正常",
                address="现场",
            ),
        ],
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    codes = {w.code for w in result.warnings}
    assert "ot_half_day_mismatch" in codes
    assert result.rows[0].ot_comp_days == 0.0


def test_project_travel_ot_only_on_holiday() -> None:
    """Rule 2: inside project travel, only statutory-holiday OT days earn credit."""
    punches = []
    for d in (10, 19):
        punches.extend(
            [
                PunchRecord(
                    person_name="壬",
                    work_date=date(2026, 6, d),
                    punch_time=datetime(2026, 6, d, 8, 0),
                    status="正常",
                    address="现场",
                ),
                PunchRecord(
                    person_name="壬",
                    work_date=date(2026, 6, d),
                    punch_time=datetime(2026, 6, d, 18, 0),
                    status="正常",
                    address="现场",
                ),
            ]
        )
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="壬",
                department="产品中心",
                center="产品中心",
                category="项目类出差（入场）",
                start=date(2026, 6, 1),
                end=date(2026, 6, 30),
                status="已通过",
            )
        ],
        overtime=[
            OvertimeDay(
                applicant="壬",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 10),
                hours=10,
                status="已通过",
            ),
            OvertimeDay(
                applicant="壬",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 19),
                hours=10,
                status="已通过",
            ),
        ],
        punch=punches,
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays={date(2026, 6, 19)},
    )
    row = result.rows[0]
    assert row.travel_comp_days == 2.0
    assert row.ot_comp_days == 1.0
    assert [d["date"] for d in row.ot_days if d["credit"] > 0] == ["2026-06-19"]


def test_ot_resolutions_write_back_credits() -> None:
    parsed = ParsedSources(
        overtime=[
            OvertimeDay(
                applicant="庚",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 19),
                hours=10,
                status="已通过",
            ),
            OvertimeDay(
                applicant="庚",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 20),
                hours=10,
                status="已通过",
            ),
            OvertimeDay(
                applicant="庚",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 21),
                hours=10,
                status="已通过",
            ),
        ]
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
        ot_resolutions={
            ("庚", "2026-06-19"): "accept_full",
            ("庚", "2026-06-20"): "accept_half",
            ("庚", "2026-06-21"): "exclude",
        },
    )
    assert result.rows[0].ot_comp_days == 1.5
    by_date = {d["date"]: d["credit"] for d in result.rows[0].ot_days}
    assert by_date == {
        "2026-06-19": 1.0,
        "2026-06-20": 0.5,
        "2026-06-21": 0.0,
    }
    assert not any(w.code.startswith("ot_") for w in result.warnings)


def test_segments_match_aggregates() -> None:
    parsed = ParsedSources(
        travel=[
            TravelRecord(
                applicant="辛",
                department="产品中心",
                center="产品中心",
                category="项目类出差（入场）",
                start=date(2026, 5, 1),
                end=date(2026, 5, 30),
                status="已通过",
            )
        ],
        overtime=[
            OvertimeDay(
                applicant="辛",
                department="产品中心",
                center="产品中心",
                work_date=date(2026, 6, 19),
                hours=10,
                status="已通过",
            )
        ],
        leave=[
            LeaveRecord(
                applicant="辛",
                department="产品中心",
                center="产品中心",
                leave_kind="调休假",
                start=date(2026, 7, 1),
                end=date(2026, 7, 1),
                duration_text="1天",
                status="已通过",
            )
        ],
        punch=[
            PunchRecord(
                person_name="辛",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 8, 0),
                status="正常",
                address="现场",
            ),
            PunchRecord(
                person_name="辛",
                work_date=date(2026, 6, 19),
                punch_time=datetime(2026, 6, 19, 18, 0),
                status="正常",
                address="现场",
            ),
        ],
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    row = result.rows[0]
    assert len(row.travel_segments) == 1
    assert row.travel_segments[0]["credit"] == 2.0
    assert sum(s["credit"] for s in row.travel_segments) == row.travel_comp_days
    assert len(row.ot_days) == 1 and row.ot_days[0]["credit"] == 1.0
    assert sum(d["credit"] for d in row.ot_days) == row.ot_comp_days
    assert len(row.leave_segments) == 1
    assert row.leave_segments[0]["used_days"] == 1.0


def test_missing_source_warns() -> None:
    result = compute_ledger(
        ParsedSources(),
        present_roles={"travel", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    assert any(w.code == "missing_source_overtime" for w in result.warnings)


def test_parse_unpadded_leave_dates() -> None:
    assert _parse_date("2026/7/1") == date(2026, 7, 1)
    assert _parse_date("2026/6/19 9:00") == date(2026, 6, 19)
    data = _xlsx_bytes(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "假种",
            "开始时间",
            "结束时间",
            "请假时长",
            "当前审批状态",
        ],
        [
            [
                "1",
                "甲",
                "产品中心/软件部",
                "调休假",
                "2026/7/1",
                "2026/7/2",
                "2天",
                "已通过",
            ]
        ],
    )
    leaves = parse_leave(data)
    assert len(leaves) == 1
    assert leaves[0].start == date(2026, 7, 1)
    assert leaves[0].end == date(2026, 7, 2)


def test_leave_segments_fill_dates() -> None:
    parsed = ParsedSources(
        leave=[
            LeaveRecord(
                applicant="癸",
                department="产品中心",
                center="产品中心",
                leave_kind="调休假",
                start=date(2026, 7, 1),
                end=None,
                duration_text="2天",
                status="已通过",
            )
        ]
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    seg = result.rows[0].leave_segments[0]
    assert seg["start"] == "2026-07-01"
    assert seg["end"] == "2026-07-02"
    assert seg["used_days"] == 2.0
    assert result.rows[0].leave_start == "2026-07-01"
    assert result.rows[0].leave_end == "2026-07-02"


def test_only_tiaoxiu_leave_counts() -> None:
    parsed = ParsedSources(
        leave=[
            LeaveRecord(
                applicant="己",
                department="产品中心",
                center="产品中心",
                leave_kind="调休假",
                start=date(2026, 7, 1),
                end=date(2026, 7, 1),
                duration_text="1天",
                status="已通过",
            ),
            LeaveRecord(
                applicant="己",
                department="产品中心",
                center="产品中心",
                leave_kind="年假",
                start=date(2026, 7, 2),
                end=date(2026, 7, 2),
                duration_text="1天",
                status="已通过",
            ),
        ]
    )
    result = compute_ledger(
        parsed,
        present_roles={"travel", "overtime", "leave", "punch"},
        hq_keywords=["太原"],
        holidays=set(),
    )
    assert result.rows[0].used_comp_days == 1.0


def test_export_headers() -> None:
    data = build_ledger_xlsx(
        [
            {
                "person_name": "甲",
                "center": "产品中心",
                "department": "产品中心/软件部",
                "travel_segments": [
                    {
                        "system_start": "2026-05-01",
                        "system_end": "2026-05-30",
                        "actual_start": "2026-05-01",
                        "actual_end": "2026-05-30",
                        "credit": 2,
                    }
                ],
                "ot_days": [
                    {"date": "2026-06-19", "credit": 1.0},
                    {"date": "2026-06-20", "credit": 0.0},
                ],
                "leave_segments": [
                    {"start": "2026-07-01", "end": "2026-07-01", "used_days": 1}
                ],
                "travel_comp_days": 2,
                "ot_comp_days": 1,
                "total_comp_days": 3,
                "used_comp_days": 1,
                "remaining_comp_days": 2,
                "remark": "项目出差30天→2天",
            }
        ]
    )
    assert looks_like_ooxml_xlsx(data)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.cell(1, 1).value == "姓名"
    assert ws.cell(1, 4).value == "出差日期（系统提报）"
    assert ws.cell(2, 4).value == "起始日"
    assert ws.cell(1, 12).value == "合计可调休天数"
    assert ws.cell(1, 16).value == "合计已休天数"
    assert ws.cell(3, 1).value == "甲"
    assert ws.cell(3, 8).value == 2
    assert ws.cell(3, 9).value == "2026-06-19"
    assert ws.cell(3, 10).value == "2026-06-19"
    assert ws.cell(3, 11).value == 1.0
    assert ws.cell(3, 12).value == 3
    # Zero-credit OT day must not occupy a row; person block is 1 data row.
    assert ws.cell(4, 9).value in (None, "")
    # Merges include identity columns across header and person block as applicable.
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:A2" in merged
    assert "D1:E1" in merged


def test_ot_resolutions_from_resolved_helper() -> None:
    mapped = _ot_resolutions_from_resolved(
        {
            "ot_missing_punch": {"甲:2026-06-19": "accept_full"},
            "ot_half_day_mismatch": {"乙:2026-06-20": "accept_half", "丙:2026-06-21": "exclude"},
        }
    )
    assert mapped[("甲", "2026-06-19")] == "accept_full"
    assert mapped[("乙", "2026-06-20")] == "accept_half"
    assert mapped[("丙", "2026-06-21")] == "exclude"


def test_open_warnings_block_until_resolved() -> None:
    warnings = [
        {
            "code": "missing_source_overtime",
            "message": "缺少加班",
            "detail": {"role": "overtime"},
            "blocking": True,
        },
        {
            "code": "ot_half_day_mismatch",
            "message": "半天",
            "detail": {"person": "甲", "date": "2026-06-19"},
            "blocking": True,
        },
    ]
    assert len(_open_blocking_warnings(warnings, {})) == 2
    resolved = {
        "missing_sources": {"overtime": "accept"},
        "ot_half_day_mismatch": {"甲:2026-06-19": "exclude"},
    }
    assert _open_blocking_warnings(warnings, resolved) == []
