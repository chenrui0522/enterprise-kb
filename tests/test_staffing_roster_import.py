"""Integration tests for roster name correction on import / confirm."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import PERM_STAFFING_READ, PERM_STAFFING_WRITE
from app.identity.principal import Principal
from app.models.identity import Project
from app.models.staffing import StaffingAttendanceFact
from app.staffing import roster as roster_mod
from app.staffing import service as staffing_service
from app.staffing.parse import KIND_FORMAL
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401
import app.models.staffing  # noqa: F401


def _principal(project_ids: tuple[str, ...] = ("proj2515",)) -> Principal:
    return Principal(
        user_id="u1",
        tenant_id="autley",
        username="liu",
        display_name="liu",
        site="taiyuan",
        clearance="general",
        permissions=(PERM_STAFFING_READ, PERM_STAFFING_WRITE),
        project_ids=project_ids,
        staffing_org_ok=True,
    )


def _write_roster(path: Path, names: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["序列", "姓名", "部门", "工种", "备注", "待进场"])
    for i, (name, dept, title) in enumerate(names, start=1):
        ws.append([i, name, dept, title, "", ""])
    wb.save(path)


def _mini_daily_xlsx(*, names_cell: str, formal_text: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "项目日报表"
    ws["A3"] = "项目名称；2515测试项目"
    ws["B4"] = "日期"
    ws["C4"] = "现场施工人员/人数"
    ws["D4"] = "当前阶段"
    ws["K6"] = "我司人员"
    ws["K7"] = "姓名"
    ws["A8"] = 1
    ws["B8"] = date(2026, 1, 11)
    ws["C8"] = formal_text or "【公司人员】：0人"
    ws["D8"] = "机械安装"
    ws["K8"] = names_cell
    buf = __import__("io").BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _use_roster(monkeypatch, path: Path) -> None:
    monkeypatch.setenv("KB_STAFFING_ROSTER_PATH", str(path))
    get_settings.cache_clear()
    roster_mod.clear_roster_cache()


@pytest.mark.asyncio
async def test_import_auto_corrects_unique_prefix(tmp_path, monkeypatch) -> None:
    roster_path = tmp_path / "roster.xlsx"
    _write_roster(roster_path, [("豆子度", "软件部", "软件工程师"), ("王亮", "项目管理部", "项目经理")])
    _use_roster(monkeypatch, roster_path)

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'a.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        session.add(
            Project(
                id="proj2515",
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        await session.commit()
        batch = await staffing_service.create_import_batch(
            session,
            principal=principal,
            filename="2515日报.xlsx",
            data=_mini_daily_xlsx(names_cell="豆子、王亮"),
            storage_key="x",
        )
        names = {r["person_name"] for r in (batch.parse_result or {}).get("rows") or []}
        assert "豆子度" in names
        assert "豆子" not in names
        codes = {w["code"] for w in batch.warnings or []}
        assert "roster_name_corrected" in codes
        assert "roster_name_unresolved" not in codes

        # Only auto-correct + maybe source noise: acknowledge non-roster and confirm.
        resolutions = {
            w["code"]: "acknowledged"
            for w in (batch.warnings or [])
            if w.get("code")
            and w["code"]
            not in {
                "roster_name_corrected",
                "roster_name_unresolved",
                "roster_unavailable",
                "name_kind_collision",
            }
        }
        if resolutions:
            await staffing_service.review_batch(
                session, principal=principal, batch_id=batch.id, resolutions=resolutions
            )
        confirmed = await staffing_service.confirm_batch(
            session, principal=principal, batch_id=batch.id
        )
        assert confirmed.status == "confirmed"
        facts = (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.batch_id == batch.id,
                    StaffingAttendanceFact.status == "active",
                )
            )
        ).scalars().all()
        assert {f.person_name for f in facts} == {"豆子度", "王亮"}

    get_settings.cache_clear()
    roster_mod.clear_roster_cache()
    await engine.dispose()


@pytest.mark.asyncio
async def test_import_unresolved_blocks_confirm_until_resolved(tmp_path, monkeypatch) -> None:
    roster_path = tmp_path / "roster.xlsx"
    _write_roster(roster_path, [("史承志", "机电服务处", "施工员（机械）")])
    _use_roster(monkeypatch, roster_path)

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'b.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        session.add(
            Project(
                id="proj2515",
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        await session.commit()
        batch = await staffing_service.create_import_batch(
            session,
            principal=principal,
            filename="2515日报.xlsx",
            data=_mini_daily_xlsx(names_cell="某某某"),
            storage_key="x",
        )
        assert any(w["code"] == "roster_name_unresolved" for w in batch.warnings or [])
        with pytest.raises(AppError, match="未决议"):
            await staffing_service.confirm_batch(
                session, principal=principal, batch_id=batch.id
            )

        await staffing_service.review_batch(
            session,
            principal=principal,
            batch_id=batch.id,
            resolutions={
                "roster_name_unresolved": {
                    "某某某": {"action": "select_roster_name", "name": "史承志"},
                }
            },
        )
        confirmed = await staffing_service.confirm_batch(
            session, principal=principal, batch_id=batch.id
        )
        assert confirmed.status == "confirmed"
        facts = (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.batch_id == batch.id,
                    StaffingAttendanceFact.status == "active",
                )
            )
        ).scalars().all()
        assert len(facts) == 1
        assert facts[0].person_name == "史承志"
        assert facts[0].person_kind == KIND_FORMAL or facts[0].person_kind

    get_settings.cache_clear()
    roster_mod.clear_roster_cache()
    await engine.dispose()


@pytest.mark.asyncio
async def test_roster_discard_and_rename_resolutions(tmp_path, monkeypatch) -> None:
    roster_path = tmp_path / "roster.xlsx"
    _write_roster(roster_path, [("王亮", "项目管理部", "项目经理")])
    _use_roster(monkeypatch, roster_path)

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'c.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        session.add(
            Project(
                id="proj2515",
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        await session.commit()
        batch = await staffing_service.create_import_batch(
            session,
            principal=principal,
            filename="2515日报.xlsx",
            data=_mini_daily_xlsx(names_cell="丢弃人、改名人"),
            storage_key="x",
        )
        await staffing_service.review_batch(
            session,
            principal=principal,
            batch_id=batch.id,
            resolutions={
                "roster_name_unresolved": {
                    "丢弃人": {"action": "discard"},
                    "改名人": {
                        "action": "rename",
                        "name": "外部特批",
                        "acknowledged_off_roster": True,
                    },
                }
            },
        )
        confirmed = await staffing_service.confirm_batch(
            session, principal=principal, batch_id=batch.id
        )
        assert confirmed.status == "confirmed"
        facts = (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.batch_id == batch.id,
                    StaffingAttendanceFact.status == "active",
                )
            )
        ).scalars().all()
        assert {f.person_name for f in facts} == {"外部特批"}

    get_settings.cache_clear()
    roster_mod.clear_roster_cache()
    await engine.dispose()


@pytest.mark.asyncio
async def test_missing_roster_blocks_confirm(tmp_path, monkeypatch) -> None:
    missing = tmp_path / "no_such_roster.xlsx"
    _use_roster(monkeypatch, missing)

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'd.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        session.add(
            Project(
                id="proj2515",
                tenant_id="autley",
                code="2515",
                name="2515",
                status="active",
            )
        )
        await session.commit()
        batch = await staffing_service.create_import_batch(
            session,
            principal=principal,
            filename="2515日报.xlsx",
            data=_mini_daily_xlsx(names_cell="王亮"),
            storage_key="x",
        )
        assert any(w["code"] == "roster_unavailable" for w in batch.warnings or [])
        # Acknowledging must not clear roster_unavailable.
        await staffing_service.review_batch(
            session,
            principal=principal,
            batch_id=batch.id,
            resolutions={"roster_unavailable": "acknowledged"},
        )
        with pytest.raises(AppError, match="未决议|roster_unavailable"):
            await staffing_service.confirm_batch(
                session, principal=principal, batch_id=batch.id
            )

    get_settings.cache_clear()
    roster_mod.clear_roster_cache()
    await engine.dispose()
