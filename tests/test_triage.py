from types import SimpleNamespace

import pytest

from app.core.config import Settings
from app.core.errors import AppError, UpstreamError
from app.ingestion.converters import (
    ConversionResult,
    DocxMarkdownConverter,
    HybridPdfConverter,
    PDFMarkdownConverter,
)
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.storage import FileDocumentStorage
from app.ingestion.triage import (
    DOC_KIND_PDF_COMPLEX,
    DOC_KIND_PDF_MIXED,
    DOC_KIND_PDF_SCANNED,
    DOC_KIND_PDF_TEXT,
    plan_candidates,
    triage_file,
)


class FakeOcrConverter:
    """Stands in for MinerU: records the pages it was asked to OCR."""

    def __init__(self, text: str = "OCR 识别内容") -> None:
        self.text = text
        self.calls: list[str] = []

    def cache_signature(self) -> dict:
        return {"converter": "fake-ocr", "version": "1"}

    async def convert_async(self, file_path: str) -> ConversionResult:
        self.calls.append(file_path)
        return ConversionResult(markdown=self.text, page_count=None)


def _write_pdf(path, pages: list[str | None]) -> None:
    import fitz

    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text, fontsize=12)
    document.save(path)
    document.close()


def _write_table_pdf(path) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Quarterly Metrics", fontsize=18)
    page.insert_text((72, 104), "The table below lists weights and results.", fontsize=11)
    left, top, col_w, row_h, cols, rows = 72, 140, 90, 20, 2, 3
    for index in range(rows + 1):
        y = top + index * row_h
        page.draw_line(fitz.Point(left, y), fitz.Point(left + cols * col_w, y))
    for index in range(cols + 1):
        x = left + index * col_w
        page.draw_line(fitz.Point(x, top), fitz.Point(x, top + rows * row_h))
    cells = [["Indicator", "Weight"], ["Completion", "40%"], ["Quality", "60%"]]
    for row in range(rows):
        for column in range(cols):
            page.insert_text(
                fitz.Point(left + 6 + column * col_w, top + 14 + row * row_h),
                cells[row][column],
                fontsize=11,
            )
    document.save(path)
    document.close()


def _settings(**overrides) -> Settings:
    values = {"mineru_url": "", "docling_url": ""}
    values.update(overrides)
    return Settings(**values)


def test_triage_detects_text_pdf_and_is_reproducible(tmp_path) -> None:
    pdf = tmp_path / "manual.pdf"
    _write_pdf(pdf, ["Warranty covers twelve months for ZB-100."])
    first = triage_file(pdf, "manual.pdf")
    second = triage_file(pdf, "manual.pdf")
    assert first.kind == DOC_KIND_PDF_TEXT
    assert first.as_report() == second.as_report()
    assert first.scanned_pages == []
    assert first.has_text_layer


def test_triage_flags_scanned_pdf(tmp_path) -> None:
    pdf = tmp_path / "scan.pdf"
    _write_pdf(pdf, [None, None])
    result = triage_file(pdf, "scan.pdf")
    assert result.kind == DOC_KIND_PDF_SCANNED
    assert result.scanned_pages == [1, 2]


def test_triage_flags_mixed_scan(tmp_path) -> None:
    pdf = tmp_path / "mixed.pdf"
    _write_pdf(pdf, ["Page one carries a text layer.", None])
    result = triage_file(pdf, "mixed.pdf")
    assert result.kind == DOC_KIND_PDF_MIXED
    assert result.scanned_pages == [2]
    assert result.is_mixed_scan


def test_triage_flags_table_page(tmp_path) -> None:
    pdf = tmp_path / "table.pdf"
    _write_table_pdf(pdf)
    result = triage_file(pdf, "table.pdf")
    assert result.kind == DOC_KIND_PDF_COMPLEX
    assert result.has_tables


def test_triage_survives_broken_pdf(tmp_path) -> None:
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"not a real pdf")
    result = triage_file(pdf, "broken.pdf")
    assert result.kind == DOC_KIND_PDF_TEXT
    assert "分诊探测失败" in result.reason


def test_plan_prefers_ocr_for_scanned_and_tables(tmp_path) -> None:
    settings = _settings(mineru_url="http://mineru.test")
    scanned = triage_file(_write_and_return(tmp_path, "scan.pdf", [None]), "scan.pdf")
    labels = [label for label, _ in plan_candidates("scan.pdf", "auto", settings, scanned)]
    assert labels[0] == "ocr"
    assert labels[-1] == "text"

    table = tmp_path / "table.pdf"
    _write_table_pdf(table)
    complex_triage = triage_file(table, "table.pdf")
    labels = [label for label, _ in plan_candidates("table.pdf", "auto", settings, complex_triage)]
    assert labels[0] == "ocr"


def test_plan_keeps_text_pdf_local(tmp_path) -> None:
    settings = _settings(mineru_url="http://mineru.test")
    pdf = tmp_path / "manual.pdf"
    _write_pdf(pdf, ["Warranty covers twelve months for ZB-100."])
    triage = triage_file(pdf, "manual.pdf")
    labels = [label for label, _ in plan_candidates("manual.pdf", "auto", settings, triage)]
    assert labels[0] == "text"


def test_plan_uses_docling_when_available(tmp_path) -> None:
    settings = _settings(mineru_url="", docling_url="http://docling.test")
    pdf = tmp_path / "manual.pdf"
    _write_pdf(pdf, ["Warranty covers twelve months for ZB-100."])
    triage = triage_file(pdf, "manual.pdf")
    labels = [
        label
        for label, _ in plan_candidates(
            "manual.pdf", "auto", settings, triage, docling_candidate=True
        )
    ]
    assert labels == ["docling", "text"]


def test_plan_routes_xlsx_through_semantic_converter(tmp_path) -> None:
    settings = _settings()
    labels = [label for label, _ in plan_candidates("报表.xlsx", "auto", settings)]
    assert labels[0] == "xlsx"
    assert "markitdown" in labels


def test_plan_requires_mineru_for_images() -> None:
    with pytest.raises(AppError, match="MinerU"):
        plan_candidates("手册.png", "auto", _settings())
    labels = [label for label, _ in plan_candidates("手册.png", "auto", _settings(mineru_url="http://m"))]
    assert labels == ["ocr"]


def _write_and_return(tmp_path, name: str, pages: list[str | None]):
    target = tmp_path / name
    _write_pdf(target, pages)
    return target


@pytest.mark.asyncio
async def test_hybrid_converter_keeps_text_pages_and_ocrs_others(tmp_path) -> None:
    pdf = tmp_path / "mixed.pdf"
    _write_pdf(pdf, ["First page body content.", None])
    ocr = FakeOcrConverter("OCR content of page two")
    converter = HybridPdfConverter(ocr, [2])
    result = await converter.convert_async(str(pdf))
    assert "First page body content" in result.markdown
    assert "OCR content of page two" in result.markdown
    assert "<!-- PAGE:1 -->" in result.markdown
    assert "<!-- PAGE:2 -->" in result.markdown
    assert result.ocr_pages == [2]
    assert len(ocr.calls) == 1


@pytest.mark.asyncio
async def test_pipeline_falls_back_when_first_candidate_fails(tmp_path, monkeypatch) -> None:
    from app.ingestion.mineru import MinerUConverter

    pdf = tmp_path / "table.pdf"
    _write_table_pdf(pdf)

    async def broken_ocr(self, file_path):  # noqa: ARG001
        raise UpstreamError("MinerU 服务不可用")

    monkeypatch.setattr(MinerUConverter, "convert_async", broken_ocr)
    storage = FileDocumentStorage(str(tmp_path / "docs"))
    settings = _settings(mineru_url="http://mineru.test")
    pipeline = IngestionPipeline(
        store=object(), embedder=object(), storage=storage, settings=settings
    )
    document = SimpleNamespace(filename="table.pdf")
    version = SimpleNamespace(parse_mode="auto")
    result = await pipeline._convert_to_markdown(document, version, pdf)

    assert "Completion" in result.markdown
    meta = pipeline._last_conversion_meta
    assert meta["label"] == "text"
    assert meta["fallback_used"] is True
    assert meta["attempts"] == 2
    assert pipeline._last_failures[0]["label"] == "ocr"


@pytest.mark.asyncio
async def test_pipeline_raises_when_every_candidate_fails(tmp_path, monkeypatch) -> None:
    from app.ingestion.mineru import MinerUConverter

    pdf = tmp_path / "scan.pdf"
    _write_pdf(pdf, [None])

    async def broken_ocr(self, file_path):  # noqa: ARG001
        raise UpstreamError("MinerU 服务不可用")

    monkeypatch.setattr(MinerUConverter, "convert_async", broken_ocr)
    storage = FileDocumentStorage(str(tmp_path / "docs"))
    settings = _settings(mineru_url="http://mineru.test")
    pipeline = IngestionPipeline(
        store=object(), embedder=object(), storage=storage, settings=settings
    )
    document = SimpleNamespace(filename="scan.pdf")
    version = SimpleNamespace(parse_mode="auto")
    with pytest.raises(AppError):
        await pipeline._convert_to_markdown(document, version, pdf)
    assert len(pipeline._last_failures) == 2


def test_docx_converter_attaches_images_to_heading(tmp_path) -> None:
    import io

    from docx import Document
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (120, 90), (200, 30, 30)).save(buffer, format="PNG")

    target = tmp_path / "manual.docx"
    doc = Document()
    doc.add_heading("安装步骤", level=1)
    doc.add_paragraph("按下图连接电源。")
    doc.add_picture(io.BytesIO(buffer.getvalue()))
    doc.add_heading("维护", level=1)
    doc.add_paragraph("每季度检查一次。")
    doc.save(target)

    result = DocxMarkdownConverter().convert(str(target))
    assert result.structure is not None
    assert len(result.images) == 1
    assert result.images[0].heading_path == "安装步骤"
    assert result.images[0].data == buffer.getvalue()


def test_pdf_converter_still_reads_text(tmp_path) -> None:
    pdf = tmp_path / "plain.pdf"
    _write_pdf(pdf, ["Plain text layer content."])
    result = PDFMarkdownConverter().convert(str(pdf))
    assert "Plain text layer content" in result.markdown
@pytest.mark.asyncio
async def test_broken_file_is_not_retried_through_ocr(tmp_path, monkeypatch) -> None:
    """A corrupt file must fail fast instead of being sent to a slow parser."""
    from app.ingestion.mineru import MinerUConverter

    calls: list[str] = []

    async def fake_ocr(self, file_path):  # noqa: ARG001
        calls.append(str(file_path))
        return ConversionResult(markdown="不该被使用", page_count=None)

    monkeypatch.setattr(MinerUConverter, "convert_async", fake_ocr)

    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf at all")
    storage = FileDocumentStorage(str(tmp_path / "docs"))
    pipeline = IngestionPipeline(
        store=object(),
        embedder=object(),
        storage=storage,
        settings=_settings(mineru_url="http://mineru.test"),
    )
    document = SimpleNamespace(filename="broken.pdf")
    version = SimpleNamespace(parse_mode="auto")
    with pytest.raises(AppError):
        await pipeline._convert_to_markdown(document, version, broken)
    assert calls == []
    assert pipeline._last_failures[0]["label"] == "text"
    assert pipeline._last_plan == ["text", "ocr"]
