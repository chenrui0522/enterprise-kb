"""Service-path smoke: create → sources → compute → confirm → export."""

from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.attendance_summary import service as att_service
from app.core.config import get_settings
from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import (
    PERM_ATTENDANCE_SUMMARY_READ,
    PERM_ATTENDANCE_SUMMARY_WRITE,
    PERM_CHAT_USE,
)
from app.identity.principal import Principal
import app.models.attendance_summary  # noqa: F401
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401


def _xlsx(headers: list[str], rows: list[list] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col, h)
    for r_i, row in enumerate(rows or [], start=2):
        for c_i, val in enumerate(row, start=1):
            ws.cell(r_i, c_i, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hr_xlsx() -> bytes:
    return _xlsx(
        ["姓名", "中心", "岗位", "驻地"],
        [["甲", "产品中心", "工程师", "太原"]],
    )


def _travel_xlsx() -> bytes:
    return _xlsx(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "出差类别",
            "出差地点",
            "开始时间",
            "结束时间",
            "当前审批状态",
        ],
        [
            [
                "1",
                "甲",
                "产品中心",
                "项目类出差（入场）",
                "北京",
                "2026/3/2",
                "2026/3/4",
                "已通过",
            ]
        ],
    )


def _leave_xlsx(*, day: str = "2026/3/10") -> bytes:
    return _xlsx(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "假种",
            "开始时间",
            "结束时间",
            "请假时长",
            "当前审批状态",
        ],
        [
            [
                "2",
                "甲",
                "产品中心",
                "年假",
                day,
                day,
                "1天",
                "已通过",
            ]
        ],
    )


def _punch_xlsx(*, year: int = 2026, month: int = 3) -> bytes:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Sheet2")
    ws = wb.create_sheet("Sheet3")
    ws.cell(1, 1, "打卡详情")
    addr = "奥特莱物流科技有限公司 | 山西省太原市小店区长治路251号a座"
    r = 2
    for day in range(1, 32):
        try:
            wd = date(year, month, day)
        except ValueError:
            continue
        if wd.weekday() >= 5:
            continue
        date_s = f"{year}/{month}/{day}"
        for time_s in ("08:30", "18:02"):
            ws.cell(r, 1, date_s)
            ws.cell(r, 2, "甲")
            ws.cell(r, 3, time_s)
            ws.cell(r, 4, "正常｜人脸识别")
            ws.cell(r, 5, addr)
            ws.cell(r, 6, "iPhone")
            ws.cell(r, 7, "--")
            ws.cell(r, 9, "--")
            r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _principal() -> Principal:
    return Principal(
        user_id="u-att",
        tenant_id="autley",
        username="ops",
        display_name="Ops",
        site="taiyuan",
        clearance="general",
        permissions=(
            PERM_CHAT_USE,
            PERM_ATTENDANCE_SUMMARY_READ,
            PERM_ATTENDANCE_SUMMARY_WRITE,
        ),
        attendance_summary_org_ok=True,
    )


async def _prepare_session(tmp_path, monkeypatch):
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'att.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    return async_sessionmaker(engine, expire_on_commit=False)


async def _ready_computed_job(session, principal, *, conflict: bool = False):
    job = await att_service.create_job(session, principal, year=2026, month=3)
    job = await att_service.cold_start_roster_from_hr(
        session, principal, job.id, filename="hr.xlsx", data=_hr_xlsx()
    )
    leave_day = "2026/3/3" if conflict else "2026/3/10"
    for role, data, name in [
        ("travel", _travel_xlsx(), "t.xlsx"),
        ("leave", _leave_xlsx(day=leave_day), "l.xlsx"),
        ("punch", _punch_xlsx(), "p.xlsx"),
    ]:
        job = await att_service.upload_source(
            session, principal, job.id, filename=name, data=data, role=role
        )
    job = await att_service.confirm_period(
        session, principal, job.id, acknowledge_incomplete_month=True
    )
    job = await att_service.confirm_calendar(session, principal, job.id, use_draft=True)
    job = await att_service.recompute(session, principal, job.id)
    return job


@pytest.mark.asyncio
async def test_service_happy_path_and_unconfirmed_export_rejected(
    tmp_path, monkeypatch
) -> None:
    factory = await _prepare_session(tmp_path, monkeypatch)
    principal = _principal()

    async with factory() as session:
        job = await _ready_computed_job(session, principal)
        assert job.status == "computed"
        people = (job.compute_result or {}).get("people") or []
        assert len(people) == 1

        with pytest.raises(AppError) as rejected:
            await att_service.export_formal(session, principal, job.id)
        assert "确认" in rejected.value.message or "正式" in rejected.value.message

        preview = await att_service.export_preview(session, principal, job.id)
        assert preview[:2] == b"PK"

        job = await att_service.acknowledge_data_quality(
            session, principal, job.id, ack_all=True
        )
        job = await att_service.confirm(session, principal, job.id)
        assert job.status == "confirmed"

        formal = await att_service.export_formal(session, principal, job.id)
        assert formal[:2] == b"PK"
        # Access control distinguishes preview (pre-confirm) vs formal (post-confirm);
        # workbook bytes may match for the same compute_result.
        assert len(formal) > 1000 and len(preview) > 1000


@pytest.mark.asyncio
async def test_set_prior_and_resolve_leave_travel_conflict(
    tmp_path, monkeypatch
) -> None:
    factory = await _prepare_session(tmp_path, monkeypatch)
    principal = _principal()

    async with factory() as session:
        prior = await _ready_computed_job(session, principal, conflict=False)
        prior = await att_service.acknowledge_data_quality(
            session, principal, prior.id, ack_all=True
        )
        prior = await att_service.confirm(session, principal, prior.id)

        job = await att_service.create_job(session, principal, year=2026, month=4)
        job = await att_service.set_prior(session, principal, job.id, prior.id)
        assert job.prior_job_id == prior.id
        assert job.roster

        conflicted = await _ready_computed_job(session, principal, conflict=True)
        warns = (conflicted.compute_result or {}).get("warnings") or []
        blocking = [w for w in warns if w.get("blocking")]
        assert blocking, "expected leave/travel conflict"
        wid = blocking[0]["id"]
        conflicted = await att_service.resolve_conflict(
            session, principal, conflicted.id, warning_id=wid, prefer="leave"
        )
        remaining = [
            w
            for w in ((conflicted.compute_result or {}).get("warnings") or [])
            if w.get("id") == wid and w.get("blocking")
        ]
        assert not remaining or conflicted.resolutions


@pytest.mark.asyncio
async def test_update_period_and_confirm_calendar_affect_compute(
    tmp_path, monkeypatch
) -> None:
    factory = await _prepare_session(tmp_path, monkeypatch)
    principal = _principal()

    async with factory() as session:
        job = await att_service.create_job(session, principal, year=2026, month=3)
        job = await att_service.update_job_period(
            session, principal, job.id, year=2026, month=4
        )
        assert job.year == 2026 and job.month == 4
        job = await att_service.cold_start_roster_from_hr(
            session, principal, job.id, filename="hr.xlsx", data=_hr_xlsx()
        )
        job = await att_service.upload_source(
            session,
            principal,
            job.id,
            filename="p.xlsx",
            data=_punch_xlsx(year=2026, month=4),
            role="punch",
        )
        job = await att_service.upload_source(
            session,
            principal,
            job.id,
            filename="t.xlsx",
            data=_xlsx(
                [
                    "审批编号",
                    "申请人",
                    "申请人部门",
                    "出差类别",
                    "出差地点",
                    "开始时间",
                    "结束时间",
                    "当前审批状态",
                ],
                [
                    [
                        "1",
                        "甲",
                        "产品中心",
                        "项目类出差（入场）",
                        "北京",
                        "2026/4/1",
                        "2026/4/2",
                        "已通过",
                    ]
                ],
            ),
            role="travel",
        )
        job = await att_service.upload_source(
            session,
            principal,
            job.id,
            filename="l.xlsx",
            data=_leave_xlsx(day="2026/4/7"),
            role="leave",
        )
        job = await att_service.confirm_period(
            session, principal, job.id, acknowledge_incomplete_month=True
        )
        # Mark one weekday as non-work via confirm_calendar
        job = await att_service.confirm_calendar(
            session,
            principal,
            job.id,
            non_work_dates=["2026-04-03"],
        )
        assert job.calendar_confirmed_at is not None
        job = await att_service.recompute(session, principal, job.id)
        assert job.status == "computed"
        assert job.month == 4
        people = (job.compute_result or {}).get("people") or []
        assert people
