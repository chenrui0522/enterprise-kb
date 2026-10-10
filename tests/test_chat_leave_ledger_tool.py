"""Smoke tests for leave-ledger chat orchestrator."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.chat.tools.leave_ledger_orchestrator import (
    apply_leave_ledger_action,
    handle_leave_ledger_message,
    update_leave_ledger_state,
)
from app.chat.tools.registry import TOOL_LEAVE_LEDGER
from app.core.config import get_settings
from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import (
    PERM_CHAT_USE,
    PERM_LEAVE_LEDGER_READ,
    PERM_LEAVE_LEDGER_WRITE,
)
from app.identity.principal import Principal
from app.leave_ledger import service as leave_service
from app.models.entity import Conversation
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401
import app.models.leave_ledger  # noqa: F401


def _xlsx_bytes(headers: list[str], rows: list[list] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col, h)
    for r_i, row in enumerate(rows or [], start=2):
        for c_i, val in enumerate(row, start=1):
            ws.cell(r_i, c_i, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _travel_xlsx() -> bytes:
    return _xlsx_bytes(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "出差类别",
            "出差事由",
            "开始时间",
            "结束时间",
            "当前审批状态",
        ],
        [
            [
                "1",
                "甲",
                "产品中心/软件部",
                "项目类出差（入场）",
                "现场",
                "2026/5/1",
                "2026/5/3",
                "已通过",
            ]
        ],
    )


def _ambiguous_xlsx() -> bytes:
    return _xlsx_bytes(["列A", "列B"], [["x", "y"]])


def _overtime_xlsx() -> bytes:
    return _xlsx_bytes(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "开始时间",
            "结束时间",
            "加班明细",
            "加班事由",
            "当前审批状态",
        ],
        [
            [
                "2",
                "甲",
                "产品中心/软件部",
                "2026/6/19 08:00",
                "2026/6/19 18:00",
                "2026/6/19 10小时",
                "端午",
                "已通过",
            ]
        ],
    )


def _leave_xlsx() -> bytes:
    return _xlsx_bytes(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "假种",
            "开始时间",
            "结束时间",
            "请假时长",
            "当前审批状态",
        ],
        [
            [
                "3",
                "甲",
                "产品中心/软件部",
                "调休",
                "2026/7/1",
                "2026/7/1",
                "1天",
                "已通过",
            ]
        ],
    )


def _principal(*, write: bool = True) -> Principal:
    perms = [PERM_CHAT_USE, PERM_LEAVE_LEDGER_READ]
    if write:
        perms.append(PERM_LEAVE_LEDGER_WRITE)
    return Principal(
        user_id="u-leave",
        tenant_id="autley",
        username="liu",
        display_name="刘珊",
        site="taiyuan",
        clearance="general",
        permissions=tuple(perms),
        leave_ledger_org_ok=True,
    )


@pytest.mark.asyncio
async def test_incremental_upload_same_job_and_role_pick(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        first = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="做调休台账",
            attachment={"filename": "出差申请.xlsx", "data": _travel_xlsx()},
            auto_note="已按调休台账处理。",
        )
        assert first["card"]["type"] == "leave_ledger_job"
        job_id = first["card"]["job_id"]
        assert conversation.tool_state["leave_ledger"]["job_id"] == job_id
        travel = next(r for r in first["card"]["roles"] if r["role"] == "travel")
        assert travel["filled"] is True

        ambiguous = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="再传一个",
            attachment={"filename": "mystery.xlsx", "data": _ambiguous_xlsx()},
        )
        assert ambiguous["card"]["need_role_pick"] is True
        assert conversation.tool_state["leave_ledger"].get("pending_attachment")

        # Replace pending with a parseable overtime workbook so pick_role can succeed
        # (unclassified garbage cannot satisfy any role parser).
        from app.chat.tools.attachments import save_chat_xlsx

        overtime_meta = save_chat_xlsx(
            user_id=principal.user_id,
            filename="source.xlsx",
            data=_overtime_xlsx(),
        )
        state = dict(conversation.tool_state or {})
        leave = dict(state.get("leave_ledger") or {})
        leave["pending_attachment"] = {
            "id": overtime_meta["id"],
            "filename": overtime_meta["filename"],
            "storage_key": overtime_meta["storage_key"],
        }
        state["leave_ledger"] = leave
        conversation.tool_state = state

        picked = await apply_leave_ledger_action(
            session,
            principal=principal,
            conversation=conversation,
            action="leave_ledger.pick_role:overtime",
            payload={"job_id": job_id},
        )
        assert picked["card"]["job_id"] == job_id
        ot = next(r for r in picked["card"]["roles"] if r["role"] == "overtime")
        assert ot["filled"] is True
        assert not conversation.tool_state["leave_ledger"].get("pending_attachment")

        second = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="补请假",
            attachment={"filename": "请假申请.xlsx", "data": _leave_xlsx()},
        )
        assert second["card"]["job_id"] == job_id
        leave_role = next(r for r in second["card"]["roles"] if r["role"] == "leave")
        assert leave_role["filled"] is True

        # Text confirm must not confirm
        with pytest.raises(AppError) as ei:
            await handle_leave_ledger_message(
                session,
                principal=principal,
                conversation=conversation,
                message="确认台账",
                attachment=None,
            )
        assert ei.value.status_code == 400

        await apply_leave_ledger_action(
            session,
            principal=principal,
            conversation=conversation,
            action="leave_ledger.accept_missing:punch",
            payload={"job_id": job_id, "role": "punch"},
        )

        with pytest.raises(AppError) as no_second:
            await apply_leave_ledger_action(
                session,
                principal=principal,
                conversation=conversation,
                action="leave_ledger.confirm",
                payload={"job_id": job_id},
            )
        assert no_second.value.status_code == 400
        assert "二次确认" in no_second.value.message

        job_card = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="进度？",
            attachment=None,
        )
        assert job_card["card"]["job_id"] == job_id
        if job_card["card"].get("open_blocking_count", 0) == 0:
            confirmed = await apply_leave_ledger_action(
                session,
                principal=principal,
                conversation=conversation,
                action="leave_ledger.confirm",
                payload={"job_id": job_id, "confirmed": True},
            )
            assert confirmed["card"]["status"] == "confirmed"
            exported = await apply_leave_ledger_action(
                session,
                principal=principal,
                conversation=conversation,
                action="leave_ledger.export",
                payload={"job_id": job_id},
            )
            assert exported["card"].get("export_ready") is True
        else:
            assert job_card["card"]["status"] != "confirmed"


@pytest.mark.asyncio
async def test_ask_current_job_person_and_summary(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-ask.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        seeded = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="上传",
            attachment={"filename": "出差申请.xlsx", "data": _travel_xlsx()},
        )
        job_id = seeded["card"]["job_id"]
        job = await leave_service.get_job(session, principal.tenant_id, job_id)
        job.compute_result = {
            "rows": [
                {
                    "person_name": "甲",
                    "travel_comp_days": 2,
                    "ot_comp_days": 1,
                    "total_comp_days": 3,
                    "used_comp_days": 1,
                    "remaining_comp_days": 2,
                    "remark": "demo",
                }
            ],
            "totals": {
                "person_count": 1,
                "total_comp_days": 3,
                "used_comp_days": 1,
                "remaining_comp_days": 2,
            },
            "warnings": [],
        }
        await session.commit()

        rem = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="甲剩余调休多少",
            attachment=None,
        )
        assert "剩余调休 2 天" in rem["content"]
        assert rem["card"]["job_id"] == job_id

        summary = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="多少人汇总",
            attachment=None,
        )
        assert "共 1 人" in summary["content"]

        missing = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="乙出差可调休几天",
            attachment=None,
        )
        assert "未在当前台账结果中找到" in missing["content"]


@pytest.mark.asyncio
async def test_ask_without_job_does_not_invent(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-ask0.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        result = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="张三剩余调休几天",
            attachment=None,
        )
        assert "还没有调休台账计算结果" in result["content"]
        assert result["card"] is None


@pytest.mark.asyncio
async def test_multi_attach_same_turn_two_roles(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-multi.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        result = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="一次两份",
            attachments=[
                {"filename": "出差申请.xlsx", "data": _travel_xlsx()},
                {"filename": "加班申请.xlsx", "data": _overtime_xlsx()},
            ],
        )
        assert result["card"]["job_id"]
        by_role = {r["role"]: r["filled"] for r in result["card"]["roles"]}
        assert by_role["travel"] is True
        assert by_role["overtime"] is True
        assert "本轮已收录 2 份" in result["content"]


@pytest.mark.asyncio
async def test_multi_attach_stops_at_ambiguous(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-amb.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        result = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="混合",
            attachments=[
                {"filename": "出差申请.xlsx", "data": _travel_xlsx()},
                {"filename": "mystery.xlsx", "data": _ambiguous_xlsx()},
                {"filename": "加班申请.xlsx", "data": _overtime_xlsx()},
            ],
        )
        assert result["card"]["need_role_pick"] is True
        by_role = {r["role"]: r["filled"] for r in result["card"]["roles"]}
        assert by_role["travel"] is True
        assert by_role["overtime"] is False
        assert "本轮已收录" in result["content"]
        assert "加班申请.xlsx" in result["content"]
        assert conversation.tool_state["leave_ledger"].get("pending_attachment")


@pytest.mark.asyncio
async def test_void_requires_confirm_and_write(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave2.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    writer = _principal(write=True)
    reader = _principal(write=False)

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=writer.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)

        result = await handle_leave_ledger_message(
            session,
            principal=writer,
            conversation=conversation,
            message="上传",
            attachment={"filename": "出差申请.xlsx", "data": _travel_xlsx()},
        )
        job_id = result["card"]["job_id"]

        with pytest.raises(AppError) as ei:
            await apply_leave_ledger_action(
                session,
                principal=writer,
                conversation=conversation,
                action="leave_ledger.void",
                payload={"job_id": job_id},
            )
        assert ei.value.status_code == 400

        with pytest.raises(AppError) as denied:
            await apply_leave_ledger_action(
                session,
                principal=reader,
                conversation=conversation,
                action="leave_ledger.void",
                payload={"job_id": job_id, "confirmed": True},
            )
        assert denied.value.status_code == 403

        voided = await apply_leave_ledger_action(
            session,
            principal=writer,
            conversation=conversation,
            action="leave_ledger.void",
            payload={"job_id": job_id, "confirmed": True},
        )
        assert voided["card"]["status"] == "voided"


@pytest.mark.asyncio
async def test_chat_merge_with_confirmed_prior(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-mrg.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        seeded = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="上传",
            attachment={"filename": "出差申请.xlsx", "data": _travel_xlsx()},
        )
        job_id = seeded["card"]["job_id"]
        job = await leave_service.get_job(session, principal.tenant_id, job_id)
        job.status = "parsed"
        job.compute_result = {
            "rows": [
                {
                    "person_name": "甲",
                    "ot_days": [{"date": "2026-02-01", "credit": 1.0}],
                    "travel_segments": [],
                    "leave_segments": [],
                }
            ]
        }
        job.warnings = []
        await session.commit()

        from app.models.leave_ledger import LeaveLedgerJob, LeaveLedgerSnapshot

        prior = LeaveLedgerJob(
            id="prior1",
            tenant_id="autley",
            status="confirmed",
            created_by=principal.user_id,
            compute_result={"rows": []},
            confirmed_at=None,
            warnings=[],
            resolved_warnings={},
        )
        session.add(prior)
        await session.flush()
        session.add(
            LeaveLedgerSnapshot(
                tenant_id="autley",
                job_id="prior1",
                rows=[
                    {
                        "person_name": "甲",
                        "ot_days": [{"date": "2026-01-05", "credit": 0.5}],
                        "travel_segments": [],
                        "leave_segments": [],
                    }
                ],
            )
        )
        await session.commit()

        with pytest.raises(AppError) as missing:
            await apply_leave_ledger_action(
                session,
                principal=principal,
                conversation=conversation,
                action="leave_ledger.merge",
                payload={"job_id": job_id},
            )
        assert missing.value.status_code == 400

        merged = await apply_leave_ledger_action(
            session,
            principal=principal,
            conversation=conversation,
            action="leave_ledger.merge",
            payload={"job_id": job_id, "prior_job_id": "prior1"},
        )
        assert "已合并" in merged["content"]
        assert merged["card"]["merged_from_job_id"] == "prior1"


@pytest.mark.asyncio
async def test_chat_export_xlsx_not_four_source_slot(tmp_path, monkeypatch) -> None:
    from app.leave_ledger.export import build_ledger_xlsx

    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'leave-exp.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()
    export_bytes = build_ledger_xlsx(
        [
            {
                "person_name": "甲",
                "center": "",
                "department": "",
                "travel_segments": [],
                "ot_days": [{"date": "2026-01-05", "credit": 0.5}],
                "leave_segments": [],
                "total_comp_days": 0.5,
                "used_comp_days": 0.0,
                "remaining_comp_days": 0.5,
                "remark": "",
            }
        ]
    )

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="调休", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_leave_ledger_state(conversation, active_tool=TOOL_LEAVE_LEDGER)

        seeded = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="上传",
            attachment={"filename": "出差申请.xlsx", "data": _travel_xlsx()},
        )
        job_id = seeded["card"]["job_id"]
        job = await leave_service.get_job(session, principal.tenant_id, job_id)
        job.status = "parsed"
        job.compute_result = {
            "rows": [
                {
                    "person_name": "甲",
                    "ot_days": [{"date": "2026-02-01", "credit": 1.0}],
                    "travel_segments": [],
                    "leave_segments": [],
                }
            ]
        }
        job.warnings = []
        await session.commit()

        offered = await handle_leave_ledger_message(
            session,
            principal=principal,
            conversation=conversation,
            message="上月台账",
            attachment={"filename": "上月台账.xlsx", "data": export_bytes},
        )
        assert "未计入四源" in offered["content"]
        assert offered["card"]["pending_prior_filename"]
        by_role = {r["role"]: r["filled"] for r in offered["card"]["roles"]}
        assert by_role["travel"] is True
        assert by_role["overtime"] is False

        done = await apply_leave_ledger_action(
            session,
            principal=principal,
            conversation=conversation,
            action="leave_ledger.confirm_prior_merge",
            payload={"job_id": job_id},
        )
        assert "已导入" in done["content"]
        assert done["card"]["merged_from_job_id"]
        dates = {
            d["date"]: d["credit"]
            for d in (
                await leave_service.get_job(session, principal.tenant_id, job_id)
            ).compute_result["rows"][0]["ot_days"]
        }
        assert dates["2026-01-05"] == 0.5
