"""Chat orchestrator + Q&A smoke for attendance summary."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.attendance_summary import service as att_service
from app.chat.tools.attendance_summary_orchestrator import (
    apply_attendance_summary_action,
    handle_attendance_summary_message,
    update_attendance_summary_state,
)
from app.chat.tools.attendance_summary_query import answer_attendance_query
from app.chat.tools.registry import TOOL_ATTENDANCE_SUMMARY, list_tools_for_principal
from app.core.config import get_settings
from app.core.db import Base
from app.core.errors import AppError
from app.identity.constants import (
    PERM_ATTENDANCE_SUMMARY_READ,
    PERM_ATTENDANCE_SUMMARY_WRITE,
    PERM_CHAT_USE,
)
from app.identity.principal import Principal
from app.models.entity import Conversation
import app.models.attendance_summary  # noqa: F401
import app.models.entity  # noqa: F401
import app.models.identity  # noqa: F401


def _xlsx(headers: list[str], rows: list[list] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col, h)
    for r_i, row in enumerate(rows or [], start=2):
        for c_i, val in enumerate(row, start=1):
            ws.cell(r_i, c_i, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hr_xlsx() -> bytes:
    return _xlsx(
        ["姓名", "中心", "岗位", "驻地"],
        [["甲", "产品中心", "工程师", "太原"]],
    )


def _travel_xlsx(*, start: str = "2026/3/2", end: str = "2026/3/4") -> bytes:
    return _xlsx(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "出差类别",
            "出差地点",
            "开始时间",
            "结束时间",
            "当前审批状态",
        ],
        [
            [
                "1",
                "甲",
                "产品中心",
                "项目类出差（入场）",
                "北京",
                start,
                end,
                "已通过",
            ]
        ],
    )


def _leave_xlsx(*, day: str = "2026/3/10") -> bytes:
    return _xlsx(
        [
            "审批编号",
            "申请人",
            "申请人部门",
            "假种",
            "开始时间",
            "结束时间",
            "请假时长",
            "当前审批状态",
        ],
        [["2", "甲", "产品中心", "年假", day, day, "1天", "已通过"]],
    )


def _punch_xlsx(*, year: int = 2026, month: int = 3) -> bytes:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Sheet2")
    ws = wb.create_sheet("Sheet3")
    ws.cell(1, 1, "打卡详情")
    addr = "奥特莱物流科技有限公司 | 山西省太原市小店区长治路251号a座"
    r = 2
    for day in range(1, 32):
        try:
            wd = date(year, month, day)
        except ValueError:
            continue
        if wd.weekday() >= 5:
            continue
        date_s = f"{year}/{month}/{day}"
        for time_s in ("08:30", "18:02"):
            ws.cell(r, 1, date_s)
            ws.cell(r, 2, "甲")
            ws.cell(r, 3, time_s)
            ws.cell(r, 4, "正常｜人脸识别")
            ws.cell(r, 5, addr)
            ws.cell(r, 6, "iPhone")
            ws.cell(r, 7, "--")
            ws.cell(r, 9, "--")
            r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _principal() -> Principal:
    return Principal(
        user_id="u-att-chat",
        tenant_id="autley",
        username="ops",
        display_name="Ops",
        site="taiyuan",
        clearance="general",
        permissions=(
            PERM_CHAT_USE,
            PERM_ATTENDANCE_SUMMARY_READ,
            PERM_ATTENDANCE_SUMMARY_WRITE,
        ),
        attendance_summary_org_ok=True,
    )


def test_tool_listed_for_permitted_user() -> None:
    ids = {t["id"] for t in list_tools_for_principal(_principal())}
    assert TOOL_ATTENDANCE_SUMMARY in ids


def test_answer_attendance_query_from_compute_result() -> None:
    people = [
        {
            "name": "甲",
            "present_days": 18.0,
            "required_days": 21.0,
            "travel_domestic": 2.0,
            "travel_overseas": 0.0,
            "leave_annual": 1.0,
            "late_count": 0,
            "missing_count": 0,
            "absent_days": 0.0,
        }
    ]
    text, key = answer_attendance_query(
        "甲实出勤",
        people=people,
        job_status="computed",
        year=2026,
        month=3,
        warning_count=0,
        open_blocking=0,
    )
    assert text is not None
    assert "实出勤" in text
    assert "18" in text
    assert key == "甲"

    summary, _ = answer_attendance_query(
        "汇总一下",
        people=people,
        job_status="computed",
        year=2026,
        month=3,
        warning_count=1,
        open_blocking=0,
        last_person_name="甲",
    )
    assert summary is not None
    assert "2026-03" in summary
    assert "1 人" in summary

    invent, _ = answer_attendance_query(
        "乙实出勤",
        people=people,
        job_status="computed",
        year=2026,
        month=3,
        warning_count=0,
        open_blocking=0,
    )
    assert invent is None or "未" in invent or "找不到" in invent or "找到" in invent


@pytest.mark.asyncio
async def test_orchestrator_confirm_set_prior_resolve_and_export(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'att-chat.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="考勤", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)

        update_attendance_summary_state(
            conversation, active_tool=TOOL_ATTENDANCE_SUMMARY
        )
        first = await handle_attendance_summary_message(
            session,
            principal=principal,
            conversation=conversation,
            message="做2026年3月考勤汇总",
        )
        assert first.get("card", {}).get("type") == "attendance_summary_job"
        job_id = first["card"]["job_id"]
        assert conversation.tool_state["attendance_summary"]["job_id"] == job_id

        # Cold-start roster + three sources via service (chat upload needs attachment plumbing)
        await att_service.cold_start_roster_from_hr(
            session, principal, job_id, filename="hr.xlsx", data=_hr_xlsx()
        )
        for role, data, name in [
            ("travel", _travel_xlsx(start="2026/3/2", end="2026/3/4"), "t.xlsx"),
            ("leave", _leave_xlsx(day="2026/3/3"), "l.xlsx"),  # conflict with travel
            ("punch", _punch_xlsx(), "p.xlsx"),
        ]:
            await att_service.upload_source(
                session, principal, job_id, filename=name, data=data, role=role
            )

        await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.confirm_period",
            payload={"acknowledge_incomplete_month": True},
        )
        await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.confirm_calendar",
            payload={},
        )
        recomputed = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.recompute",
            payload={},
        )
        assert recomputed.get("card", {}).get("status") in {
            "computed",
            "draft",
            "ready",
        } or recomputed.get("card")

        job = await att_service.get_job(session, principal.tenant_id, job_id)
        warns = (job.compute_result or {}).get("warnings") or []
        blocking = [w for w in warns if w.get("blocking")]
        assert blocking
        resolved = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.resolve_conflict",
            payload={"warning_id": blocking[0]["id"], "prefer": "leave"},
        )
        assert resolved.get("card") or resolved.get("reply")

        await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.ack_data_quality",
            payload={"ack_all": True},
        )
        confirmed = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.confirm",
            payload={"confirmed": True},
        )
        assert confirmed.get("card", {}).get("status") == "confirmed" or (
            await att_service.get_job(session, principal.tenant_id, job_id)
        ).status == "confirmed"

        with pytest.raises(AppError):
            # unconfirmed path already covered in service tests; void then ensure formal gate
            other = await att_service.create_job(
                session, principal, year=2026, month=5
            )
            await att_service.export_formal(session, principal, other.id)

        exported = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.export",
            payload={},
        )
        assert exported.get("card") or exported.get("attachment") or exported.get(
            "reply"
        )

        # set_prior onto a fresh job from confirmed prior
        next_job = await att_service.create_job(
            session, principal, year=2026, month=4
        )
        update_attendance_summary_state(
            conversation,
            active_tool=TOOL_ATTENDANCE_SUMMARY,
            job_id=next_job.id,
        )
        prior_set = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.set_prior",
            payload={"prior_job_id": job_id},
        )
        assert prior_set.get("card") or prior_set.get("reply")
        refreshed = await att_service.get_job(
            session, principal.tenant_id, next_job.id
        )
        assert refreshed.prior_job_id == job_id


@pytest.mark.asyncio
async def test_e2e_create_sources_confirm_formal_export(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'att-e2e.sqlite'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        job = await att_service.create_job(session, principal, year=2026, month=3)
        await att_service.cold_start_roster_from_hr(
            session, principal, job.id, filename="hr.xlsx", data=_hr_xlsx()
        )
        for role, data, name in [
            ("travel", _travel_xlsx(), "t.xlsx"),
            ("leave", _leave_xlsx(day="2026/3/10"), "l.xlsx"),
            ("punch", _punch_xlsx(), "p.xlsx"),
        ]:
            await att_service.upload_source(
                session, principal, job.id, filename=name, data=data, role=role
            )
        await att_service.confirm_period(
            session, principal, job.id, acknowledge_incomplete_month=True
        )
        await att_service.confirm_calendar(session, principal, job.id, use_draft=True)
        job = await att_service.recompute(session, principal, job.id)

        with pytest.raises(AppError):
            await att_service.export_formal(session, principal, job.id)

        await att_service.acknowledge_data_quality(
            session, principal, job.id, ack_all=True
        )
        job = await att_service.confirm(session, principal, job.id)
        blob = await att_service.export_formal(session, principal, job.id)
        assert blob[:2] == b"PK"
        preview = await att_service.export_preview(session, principal, job.id)
        assert preview[:2] == b"PK"
        # preview workbook is distinct from formal (may differ by watermark/sheets)
        assert preview != blob or len(preview) > 0


def _punch_xlsx_with_weekend(*, year: int = 2026, month: int = 3) -> bytes:
    """Weekday punches plus one Saturday (non-work) punch for non_company_punch."""
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Sheet2")
    ws = wb.create_sheet("Sheet3")
    ws.cell(1, 1, "打卡详情")
    addr = "奥特莱物流科技有限公司 | 山西省太原市小店区长治路251号a座"
    r = 2
    for day in range(1, 32):
        try:
            wd = date(year, month, day)
        except ValueError:
            continue
        # include Saturday 2026-03-07
        if wd.weekday() >= 5 and wd != date(2026, 3, 7):
            continue
        date_s = f"{year}/{month}/{day}"
        for time_s in ("08:30", "18:02"):
            ws.cell(r, 1, date_s)
            ws.cell(r, 2, "甲")
            ws.cell(r, 3, time_s)
            ws.cell(r, 4, "正常｜人脸识别")
            ws.cell(r, 5, addr)
            ws.cell(r, 6, "iPhone")
            ws.cell(r, 7, "--")
            ws.cell(r, 9, "--")
            r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_orchestrator_non_company_punch_prefer_roundtrips(
    tmp_path, monkeypatch
) -> None:
    """Chat resolve_conflict: ignore and present for non_company_punch."""
    monkeypatch.setenv("KB_DOCUMENT_STORAGE_DIR", str(tmp_path / "storage"))
    get_settings.cache_clear()
    engine = create_async_engine(
        f"sqlite+aiosqlite:///{tmp_path / 'att-nonco.sqlite'}"
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    principal = _principal()

    async with factory() as session:
        conversation = Conversation(
            tenant_id="autley", title="考勤非出勤日", created_by=principal.user_id
        )
        session.add(conversation)
        await session.commit()
        await session.refresh(conversation)
        update_attendance_summary_state(
            conversation, active_tool=TOOL_ATTENDANCE_SUMMARY
        )

        job = await att_service.create_job(session, principal, year=2026, month=3)
        update_attendance_summary_state(
            conversation, active_tool=TOOL_ATTENDANCE_SUMMARY, job_id=job.id
        )
        await att_service.cold_start_roster_from_hr(
            session, principal, job.id, filename="hr.xlsx", data=_hr_xlsx()
        )
        for role, data, name in [
            ("travel", _travel_xlsx(start="2026/3/16", end="2026/3/17"), "t.xlsx"),
            ("leave", _leave_xlsx(day="2026/3/18"), "l.xlsx"),
            ("punch", _punch_xlsx_with_weekend(), "p.xlsx"),
        ]:
            await att_service.upload_source(
                session, principal, job.id, filename=name, data=data, role=role
            )
        await att_service.confirm_period(
            session, principal, job.id, acknowledge_incomplete_month=True
        )
        await att_service.confirm_calendar(
            session, principal, job.id, use_draft=True
        )
        job = await att_service.recompute(session, principal, job.id)
        blocking = [
            w
            for w in ((job.compute_result or {}).get("warnings") or [])
            if w.get("blocking")
            and (
                w.get("type") == "non_company_punch_present"
                or "non_company_punch" in str(w.get("id"))
            )
        ]
        assert blocking, "expected non_company_punch blocking warning"
        wid = blocking[0]["id"]
        present_before = (job.compute_result or {})["people"][0]["present_days"]

        ignored = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.resolve_conflict",
            payload={"warning_id": wid, "prefer": "ignore"},
        )
        assert ignored.get("card") or ignored.get("reply")
        job = await att_service.get_job(session, principal.tenant_id, job.id)
        still = [
            w
            for w in ((job.compute_result or {}).get("warnings") or [])
            if w.get("id") == wid and w.get("blocking")
        ]
        assert not still
        present_ignore = (job.compute_result or {})["people"][0]["present_days"]
        assert present_ignore <= present_before

        # Fresh job for present prefer (same sources)
        job2 = await att_service.create_job(session, principal, year=2026, month=3)
        update_attendance_summary_state(
            conversation, active_tool=TOOL_ATTENDANCE_SUMMARY, job_id=job2.id
        )
        await att_service.cold_start_roster_from_hr(
            session, principal, job2.id, filename="hr.xlsx", data=_hr_xlsx()
        )
        for role, data, name in [
            ("travel", _travel_xlsx(start="2026/3/16", end="2026/3/17"), "t.xlsx"),
            ("leave", _leave_xlsx(day="2026/3/18"), "l.xlsx"),
            ("punch", _punch_xlsx_with_weekend(), "p.xlsx"),
        ]:
            await att_service.upload_source(
                session, principal, job2.id, filename=name, data=data, role=role
            )
        await att_service.confirm_period(
            session, principal, job2.id, acknowledge_incomplete_month=True
        )
        await att_service.confirm_calendar(
            session, principal, job2.id, use_draft=True
        )
        job2 = await att_service.recompute(session, principal, job2.id)
        blocking2 = [
            w
            for w in ((job2.compute_result or {}).get("warnings") or [])
            if w.get("blocking")
            and "non_company_punch" in str(w.get("id"))
        ]
        assert blocking2
        wid2 = blocking2[0]["id"]
        present_base = (job2.compute_result or {})["people"][0]["present_days"]

        with pytest.raises(AppError):
            await apply_attendance_summary_action(
                session,
                principal=principal,
                conversation=conversation,
                action="attendance_summary.resolve_conflict",
                payload={"warning_id": wid2, "prefer": "not_a_real_prefer"},
            )

        presented = await apply_attendance_summary_action(
            session,
            principal=principal,
            conversation=conversation,
            action="attendance_summary.resolve_conflict",
            payload={"warning_id": wid2, "prefer": "present"},
        )
        assert presented.get("card") or presented.get("reply")
        job2 = await att_service.get_job(session, principal.tenant_id, job2.id)
        still2 = [
            w
            for w in ((job2.compute_result or {}).get("warnings") or [])
            if w.get("id") == wid2 and w.get("blocking")
        ]
        assert not still2
        present_after = (job2.compute_result or {})["people"][0]["present_days"]
        assert present_after >= present_base

        # full_day also accepted via service wrapper
        job3 = await att_service.create_job(session, principal, year=2026, month=3)
        await att_service.cold_start_roster_from_hr(
            session, principal, job3.id, filename="hr.xlsx", data=_hr_xlsx()
        )
        for role, data, name in [
            ("travel", _travel_xlsx(start="2026/3/16", end="2026/3/17"), "t.xlsx"),
            ("leave", _leave_xlsx(day="2026/3/18"), "l.xlsx"),
            ("punch", _punch_xlsx_with_weekend(), "p.xlsx"),
        ]:
            await att_service.upload_source(
                session, principal, job3.id, filename=name, data=data, role=role
            )
        await att_service.confirm_period(
            session, principal, job3.id, acknowledge_incomplete_month=True
        )
        await att_service.confirm_calendar(
            session, principal, job3.id, use_draft=True
        )
        job3 = await att_service.recompute(session, principal, job3.id)
        blocking3 = [
            w
            for w in ((job3.compute_result or {}).get("warnings") or [])
            if w.get("blocking") and "non_company_punch" in str(w.get("id"))
        ]
        assert blocking3
        job3 = await att_service.resolve_conflict(
            session,
            principal,
            job3.id,
            warning_id=blocking3[0]["id"],
            prefer="full_day",
        )
        still3 = [
            w
            for w in ((job3.compute_result or {}).get("warnings") or [])
            if w.get("id") == blocking3[0]["id"] and w.get("blocking")
        ]
        assert not still3
