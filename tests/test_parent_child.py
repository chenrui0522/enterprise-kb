from types import SimpleNamespace

import pytest

from app.ingestion.chunker import ChunkContext
from app.ingestion.parent_child import (
    CHUNK_KIND_CHILD,
    CHUNK_KIND_IMAGE,
    CHUNK_KIND_TABLE_ROW,
    CHUNK_KIND_TABLE_SUMMARY,
    ChunkParams,
    build_chunk_plan,
    params_for,
)
from app.ingestion.structure import (
    BLOCK_HEADING,
    BLOCK_IMAGE,
    BLOCK_PARAGRAPH,
    BLOCK_TABLE,
    Block,
    DocumentStructure,
    TableArtifact,
)
from app.ingestion.tabular import XlsxSemanticConverter


def _context() -> ChunkContext:
    return ChunkContext(doc_id="doc1", version_id="v1", tenant_id="t1", title="员工手册")


def _structure(*blocks: Block) -> DocumentStructure:
    return DocumentStructure(blocks=list(blocks))


def test_parents_follow_the_heading_tree() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="第一章 总则", level=1, page=1),
        Block(type=BLOCK_PARAGRAPH, text="本手册适用于全体员工。", page=1),
        Block(type=BLOCK_HEADING, text="1.1 考勤", level=2, page=2),
        Block(type=BLOCK_PARAGRAPH, text="员工需按时打卡。", page=2),
    )
    plan = build_chunk_plan(
        structure, _context(), doc_type="generic", params=params_for("generic")
    )
    heading_paths = [parent.heading_path for parent in plan.parents]
    assert heading_paths == ["第一章 总则", "第一章 总则 > 1.1 考勤"]
    assert all(parent.child_count > 0 for parent in plan.parents)
    assert all(child.chunk_kind == CHUNK_KIND_CHILD for child in plan.children)


def test_children_never_cross_their_parent() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="第一章", level=1),
        Block(type=BLOCK_PARAGRAPH, text="第一段内容。" * 40),
        Block(type=BLOCK_HEADING, text="第二章", level=1),
        Block(type=BLOCK_PARAGRAPH, text="第二段内容。" * 40),
    )
    plan = build_chunk_plan(
        structure,
        _context(),
        doc_type="generic",
        params=ChunkParams(parent_size=300, child_size=120, child_overlap=20),
    )
    parents = {parent.id: parent for parent in plan.parents}
    assert len(plan.parents) >= 2
    for child in plan.children:
        parent = parents[child.parent_id]
        assert parent.text[:20] in child.text or child.text[:20] in parent.text
        assert parent.heading_path in {child.heading_path, parent.heading_path}


def test_long_subtree_splits_into_second_level_parents() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="安装手册", level=1),
        *[
            Block(type=BLOCK_PARAGRAPH, text=f"第{index}段说明文字，内容足够长以便触发二级父块切分。" * 6)
            for index in range(1, 9)
        ],
    )
    plan = build_chunk_plan(
        structure,
        _context(),
        doc_type="generic",
        params=ChunkParams(parent_size=400, child_size=150, child_overlap=20),
    )
    assert len(plan.parents) > 1
    assert {parent.heading_path for parent in plan.parents} == {"安装手册"}
    assert len({parent.id for parent in plan.parents}) == len(plan.parents)


def test_table_produces_row_and_summary_chunks_only() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="报销标准", level=1, page=3),
        Block(type=BLOCK_PARAGRAPH, text="下表列出差旅报销标准。", page=3),
        Block(
            type=BLOCK_TABLE,
            header=["项目", "上限"],
            rows=[["住宿", "500"], ["交通", "实报实销"], ["餐补", "100"]],
            text="| 项目 | 上限 |",
            page=3,
        ),
    )
    plan = build_chunk_plan(
        structure, _context(), doc_type="generic", params=params_for("generic")
    )
    row_chunks = [child for child in plan.children if child.chunk_kind == CHUNK_KIND_TABLE_ROW]
    summary_chunks = [
        child for child in plan.children if child.chunk_kind == CHUNK_KIND_TABLE_SUMMARY
    ]
    assert len(row_chunks) == 3
    assert [chunk.row_index for chunk in row_chunks] == [1, 2, 3]
    assert all(chunk.table_id for chunk in row_chunks)
    assert "项目=住宿；上限=500" in row_chunks[0].text
    assert len(summary_chunks) == 1
    assert "共 3 行数据" in summary_chunks[0].text
    # The grid never becomes a second, whole-table chunk.
    assert all("| 项目 |" not in child.text for child in plan.children)
    assert plan.tables[0].row_count == 3


def test_spreadsheet_rows_are_indexed_once(tmp_path) -> None:
    from openpyxl import Workbook

    target = tmp_path / "roster.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "花名册"
    sheet.append(["姓名", "部门"])
    sheet.append(["张三", "研发部"])
    sheet.append(["李四", "财务部"])
    workbook.save(target)

    converted = XlsxSemanticConverter().convert(str(target))
    plan = build_chunk_plan(
        converted.structure,
        _context(),
        doc_type="table",
        params=params_for("table"),
        tables=converted.tables,
    )
    row_chunks = [child for child in plan.children if child.chunk_kind == CHUNK_KIND_TABLE_ROW]
    assert len(row_chunks) == 2
    assert [chunk.row_index for chunk in row_chunks] == [1, 2]
    assert "姓名=张三；部门=研发部" in row_chunks[0].text
    assert plan.parents == []
    # Each data row is indexed exactly once as a row chunk (the table summary
    # may quote it as an example, but never as a second row chunk).
    assert sum(1 for child in row_chunks if "姓名=张三" in child.text) == 1


def test_table_documents_disable_parent_child() -> None:
    structure = _structure(
        Block(type=BLOCK_PARAGRAPH, text="说明文字。"),
        Block(type=BLOCK_TABLE, header=["A"], rows=[["1"]], text="| A |"),
    )
    plan = build_chunk_plan(
        structure, _context(), doc_type="table", params=params_for("table")
    )
    assert plan.parents == []
    assert any(child.chunk_kind == CHUNK_KIND_TABLE_ROW for child in plan.children)
    assert all(child.parent_id == "" for child in plan.children)


def test_images_are_self_contained_chunks() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="接线图", level=1),
        Block(type=BLOCK_IMAGE, text="电源接线示意图", image_id="a" * 32, page=7),
    )
    plan = build_chunk_plan(
        structure, _context(), doc_type="generic", params=params_for("generic")
    )
    images = [child for child in plan.children if child.chunk_kind == CHUNK_KIND_IMAGE]
    assert len(images) == 1
    assert images[0].parent_id == ""
    assert images[0].image_id == "a" * 32
    assert images[0].heading_path == "接线图"


def test_policy_children_keep_clause_metadata() -> None:
    structure = _structure(
        Block(type=BLOCK_HEADING, text="差旅制度", level=1),
        Block(type=BLOCK_PARAGRAPH, text="第一条 员工出差需提前申请。"),
        Block(type=BLOCK_PARAGRAPH, text="第二条 住宿标准见附表。"),
    )
    plan = build_chunk_plan(
        structure, _context(), doc_type="policy", params=params_for("policy")
    )
    clause_children = [child for child in plan.children if child.clause_no]
    assert {child.clause_no for child in clause_children} == {"第一条", "第二条"}
    assert all(child.chunk_type == "policy_clause" for child in clause_children)
    assert all(child.parent_id for child in clause_children)


def test_params_scale_with_settings_and_can_be_disabled() -> None:
    base = params_for("generic")
    scaled = params_for("generic", SimpleNamespace(chunk_size=1600, chunk_overlap=160))
    assert scaled.child_size == base.child_size * 2
    assert scaled.parent_size == base.parent_size * 2

    disabled = params_for(
        "generic", SimpleNamespace(chunk_size=800, chunk_overlap=80, parent_child_enabled=False)
    )
    assert disabled.parent_child is False

    table = params_for("table", SimpleNamespace(chunk_size=800, chunk_overlap=80))
    assert table.parent_child is False


def test_table_artifact_is_used_for_spreadsheets(tmp_path) -> None:
    structure = _structure(Block(type=BLOCK_PARAGRAPH, text="预算说明。"))
    plan = build_chunk_plan(
        structure,
        _context(),
        doc_type="table",
        params=params_for("table"),
        tables=[
            TableArtifact(
                sheet="部门预算",
                header=["部门", "预算"],
                rows=[["研发部", "100"], ["财务部", "80"]],
                source="xlsx",
            )
        ],
    )
    assert len(plan.tables) == 1
    assert plan.tables[0].name == "部门预算"
    summary = next(
        child for child in plan.children if child.chunk_kind == CHUNK_KIND_TABLE_SUMMARY
    )
    assert "部门预算" in summary.text
    assert "部门、预算" in summary.text
    rows = [child for child in plan.children if child.chunk_kind == CHUNK_KIND_TABLE_ROW]
    assert len(rows) == 2
    assert "部门=研发部；预算=100" in rows[0].text


def test_chunker_version_label_is_applied() -> None:
    structure = _structure(Block(type=BLOCK_PARAGRAPH, text="内容。"))
    plan = build_chunk_plan(
        structure, _context(), doc_type="generic", params=params_for("generic")
    )
    assert all(child.chunker_version == "structure-v2:generic" for child in plan.children)
    assert all(parent.chunker_version == "structure-v2:generic" for parent in plan.parents)
    assert [child.chunk_index for child in plan.children] == list(range(len(plan.children)))