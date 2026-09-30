"""Unit tests for staffing daily-report parsing (internal staff only)."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.core.errors import AppError
from app.staffing.parse import (
    KIND_CONTRACT,
    KIND_FORMAL,
    extract_project_codes,
    is_ooxml_xlsx,
    parse_daily_report,
)


def _build_sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "项目日报表"
    ws["A3"] = "项目名称；2515JS壹号智能西安项目    项目经理：王亮"
    ws["A4"] = "序号"
    ws["B4"] = "日期"
    ws["C4"] = "现场施工人员/人数"
    ws["D4"] = "当前阶段"
    ws["J4"] = "现场人员（人）"
    ws["J5"] = "机械安装"
    ws["K6"] = "我司人员\n姓名"
    ws["K7"] = "姓名"

    # Day 1: formal empty, contract names (机械)
    ws["A8"] = 1
    ws["B8"] = date(2026, 1, 11)
    ws["C8"] = "【公司人员】：0人； 【外包人员】：机械安装3人"
    ws["D8"] = "机械安装"
    ws["K8"] = "赵鑫磊、王亮、高建、张梦翔"

    # Day 2: formal 王亮 + outsource ignored
    ws["A9"] = 2
    ws["B9"] = date(2026, 3, 9)
    ws["C9"] = "【公司人员】：1人 王亮 【外包人员】：电气安装人员2人（刘福亮、汪万里）"
    ws["D9"] = "电气安装"
    ws["K9"] = "赵鑫磊、王亮、高建、张梦翔"

    # Day 3: formal 冯江伟、马越
    ws["A10"] = 3
    ws["B10"] = date(2026, 4, 11)
    ws["C10"] = "【公司人员】：2人，冯江伟、马越 、 【外包人员】：电气安装人员0人"
    ws["D10"] = "软件调试"
    ws["K10"] = "赵鑫磊、王亮、高建、张梦翔"

    # Ambiguous token
    ws["A11"] = 4
    ws["B11"] = date(2026, 7, 1)
    ws["C11"] = "【公司人员】：1人，王磊工 【外包人员】："
    ws["D11"] = "试运行"
    ws["K11"] = ""

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_rejects_non_ooxml() -> None:
    assert is_ooxml_xlsx(b"not-a-zip") is False
    with pytest.raises(AppError) as ei:
        parse_daily_report(b"\x87\x7d\x1c\x32garbage", filename="bad.xlsx")
    assert ei.value.status_code == 415


def test_extract_codes_from_filename() -> None:
    assert "2515" in extract_project_codes("2515项目日报260721.xlsx")


def test_suspected_name_split_warns_for_prefix_fragment() -> None:
    from app.staffing.parse import _suspected_name_split_pairs, parse_daily_report

    assert ("豆子", "豆子度") in _suspected_name_split_pairs(["豆子度", "豆子", "王亮"])
    assert _suspected_name_split_pairs(["张三", "李四"]) == []

    wb = Workbook()
    ws = wb.active
    ws["B4"] = "日期"
    ws["C4"] = "现场施工人员/人数"
    ws["K6"] = "我司人员\n姓名"
    ws["B8"] = date(2026, 8, 17)
    ws["C8"] = "【公司人员】：1人，豆子 【外包人员】："
    ws["K8"] = ""
    ws["B9"] = date(2026, 8, 18)
    ws["C9"] = "【公司人员】：1人，豆子度 【外包人员】："
    ws["K9"] = ""
    buf = io.BytesIO()
    wb.save(buf)
    result = parse_daily_report(buf.getvalue(), filename="2515项目日报.xlsx")
    assert any(
        w.code == "suspected_name_split"
        and (w.detail or {}).get("short_name") == "豆子"
        and (w.detail or {}).get("long_name") == "豆子度"
        for w in result.warnings
    )


def test_parse_internal_only_and_kinds() -> None:
    data = _build_sample_xlsx()
    result = parse_daily_report(data, filename="2515项目日报.xlsx")
    assert "2515" in result.filename_codes

    # Outsource names must never appear
    names = {r.person_name for r in result.rows}
    assert "刘福亮" not in names
    assert "汪万里" not in names

    # Contract names from K on Jan 11
    jan = [r for r in result.rows if r.work_date == "2026-01-11"]
    assert {r.person_name for r in jan} == {"赵鑫磊", "王亮", "高建", "张梦翔"}
    assert all(r.person_kind == KIND_CONTRACT for r in jan)
    assert all(r.stage == "机械安装" for r in jan)

    # Mar 9: 王亮 appears as both formal (C) and contract (K) → two rows that day
    mar_wang = [
        r.person_kind
        for r in result.rows
        if r.work_date == "2026-03-09" and r.person_name == "王亮"
    ]
    assert set(mar_wang) == {KIND_FORMAL, KIND_CONTRACT}
    assert all(r.stage == "电气安装" for r in result.rows if r.work_date == "2026-03-09")
    mar_contract = {
        r.person_name
        for r in result.rows
        if r.work_date == "2026-03-09" and r.person_kind == KIND_CONTRACT
    }
    assert "赵鑫磊" in mar_contract

    # Formal pair on Apr 11
    apr = {
        (r.person_name, r.person_kind)
        for r in result.rows
        if r.work_date == "2026-04-11"
    }
    assert ("冯江伟", KIND_FORMAL) in apr
    assert ("马越", KIND_FORMAL) in apr
    assert all(r.stage == "软件调试" for r in result.rows if r.work_date == "2026-04-11")

    # Ambiguous 王磊工 still counted + warned
    assert any(r.person_name == "王磊工" for r in result.rows)
    assert any(w.code == "ambiguous_person_token" for w in result.warnings)

    # 王亮 appears as both kinds across batch → collision warning
    assert any(
        w.code == "name_kind_collision" and (w.detail or {}).get("person_name") == "王亮"
        for w in result.warnings
    )

def test_desktop_sample_if_present() -> None:
    path = Path(r"c:\Users\chenr\Desktop\项目日报1.xlsx")
    if not path.exists():
        pytest.skip("desktop sample not present")
    data = path.read_bytes()
    result = parse_daily_report(data, filename=path.name)
    names = {r.person_name for r in result.rows}
    assert "刘福亮" not in names
    assert "冯江伟" in names or "王亮" in names or "赵鑫磊" in names
    # Days on site for 冯江伟 if present
    if "冯江伟" in names:
        days = {r.work_date for r in result.rows if r.person_name == "冯江伟"}
        assert len(days) >= 1
