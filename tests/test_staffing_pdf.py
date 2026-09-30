"""Excel-exported daily-report PDFs rebuild into the existing xlsx parser."""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from app.core.errors import AppError
from app.ingestion.storage import FileDocumentStorage
from app.staffing.parse import (
    UNSUPPORTED_UPLOAD_MESSAGE,
    _as_date,
    parse_daily_report,
)
from app.staffing.pdf_grid import (
    UnreadableStaffingPdf,
    extract_pdf_grid,
    pdf_to_xlsx_bytes,
)
from tests.test_staffing_parse import _build_sample_xlsx


def _ruled_pdf(pages: list[list[list[str]]]) -> bytes:
    doc = pymupdf.open()
    for grid in pages:
        cols = max((len(row) for row in grid), default=1)
        widths = []
        for col in range(cols):
            longest = 1
            for row in grid:
                if col < len(row) and row[col]:
                    longest = max(longest, len(row[col]))
            widths.append(max(36, longest * 8 + 10))
        row_h = 18
        page = doc.new_page(width=sum(widths) + 24, height=row_h * max(len(grid), 1) + 24)
        x = 12
        for col in range(cols):
            y = 12
            for row in grid:
                rect = pymupdf.Rect(x, y, x + widths[col], y + row_h)
                page.draw_rect(rect, color=(0, 0, 0), width=0.4)
                text = row[col] if col < len(row) else ""
                if text:
                    page.insert_textbox(
                        pymupdf.Rect(rect.x0 + 1, rect.y0 + 1, rect.x1 - 1, rect.y1 - 1),
                        text,
                        fontsize=8,
                        fontname="china-s",
                    )
                y += row_h
            x += widths[col]
    return doc.tobytes()


def _image_only_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(20, 20, 120, 80), color=(0, 0, 0), fill=(0.2, 0.2, 0.2))
    return doc.tobytes()


def _sheet_grid(data: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    sheet = workbook.active
    grid: list[list[str]] = []
    for row_index in range(1, (sheet.max_row or 0) + 1):
        row: list[str] = []
        for col_index in range(1, (sheet.max_column or 0) + 1):
            value = sheet.cell(row_index, col_index).value
            if isinstance(value, datetime):
                row.append(f"{value.year}/{value.month}/{value.day}")
            elif isinstance(value, date):
                row.append(f"{value.year}/{value.month}/{value.day}")
            elif value is None:
                row.append("")
            else:
                row.append(str(value).replace("\n", " "))
        grid.append(row)
    return grid


def _facts(data: bytes, filename: str) -> set[tuple[str, str, str]]:
    result = parse_daily_report(data, filename=filename)
    return {(row.work_date, row.person_name, row.person_kind) for row in result.rows}


def test_rejects_pdf_without_text_layer() -> None:
    with pytest.raises(AppError) as caught:
        extract_pdf_grid(_image_only_pdf())
    assert caught.value.status_code == 415
    assert "文字层" in caught.value.message


def test_keeps_multirow_header_and_blank_cells() -> None:
    grid = extract_pdf_grid(
        _ruled_pdf(
            [
                [
                    ["项目日报", "", ""],
                    ["日期", "现场施工人员/人数", "我司人员"],
                    ["2026/1/11", "【公司人员】：2人，冯江伟、马越", ""],
                ]
            ]
        )
    )
    assert grid[0] == ["项目日报", "", ""]
    assert grid[1][0] == "日期"
    assert grid[1][1] == "现场施工人员/人数"
    assert grid[2][2] == ""
    assert "冯江伟" in grid[2][1]


def test_stitches_later_page_and_drops_repeated_header() -> None:
    header = [
        ["项目日报表", ""],
        ["日期", "现场施工人员/人数"],
    ]
    grid = extract_pdf_grid(
        _ruled_pdf(
            [
                header + [["2026/1/11", "【公司人员】：1人，冯江伟"]],
                header + [["2026/3/9", "【公司人员】：1人，马越"]],
            ]
        )
    )
    flat = [cell for row in grid for cell in row]
    assert flat.count("日期") == 1
    assert any("冯江伟" in cell for cell in flat)
    assert any("马越" in cell for cell in flat)
    assert len(grid) == 4


def test_page_without_date_column_is_unreadable() -> None:
    with pytest.raises(UnreadableStaffingPdf):
        extract_pdf_grid(
            _ruled_pdf(
                [
                    [
                        ["日期", "现场施工人员/人数"],
                        ["2026/1/11", "【公司人员】：1人，冯江伟"],
                    ],
                    [
                        ["姓名", "备注"],
                        ["赵鑫磊", "机械安装"],
                    ],
                ]
            )
        )


def test_date_string_becomes_workbook_date() -> None:
    data = pdf_to_xlsx_bytes(
        _ruled_pdf([[["日期", "备注"], ["2026/1/11", "下周到场"]]])
    )
    sheet = load_workbook(io.BytesIO(data)).active
    assert _as_date(sheet["A2"].value) == date(2026, 1, 11)
    assert sheet["B2"].value == "下周到场"


def test_pdf_matches_source_xlsx_and_stored_bytes_stay_pdf(tmp_path: Path) -> None:
    workbook = _build_sample_xlsx()
    pdf = _ruled_pdf([_sheet_grid(workbook)])
    storage = FileDocumentStorage(tmp_path)
    key = storage.store_staffing("batch-1", "2515项目日报.pdf", pdf)
    assert storage.read_bytes(key).startswith(b"%PDF")
    assert _facts(pdf, "2515项目日报.pdf") == _facts(workbook, "2515项目日报.xlsx")
    names = {name for _day, name, _kind in _facts(pdf, "2515项目日报.pdf")}
    assert "刘福亮" not in names
    assert "汪万里" not in names


def test_pdf_without_required_columns_warns_and_writes_no_rows() -> None:
    result = parse_daily_report(
        _ruled_pdf([[["标题", "说明"], ["只有文字", "没有日报列"]]]),
        filename="notes.pdf",
    )
    assert result.rows == []
    assert any(warning.code == "unreadable_workbook" for warning in result.warnings)


def test_word_and_other_files_are_rejected() -> None:
    with pytest.raises(AppError) as word:
        parse_daily_report(b"PK\x03\x04not-a-workbook", filename="日报.docx")
    assert word.value.status_code == 415
    assert word.value.message == UNSUPPORTED_UPLOAD_MESSAGE

    with pytest.raises(AppError) as other:
        parse_daily_report(b"hello", filename="notes.txt")
    assert other.value.status_code == 415
    assert other.value.message == UNSUPPORTED_UPLOAD_MESSAGE
