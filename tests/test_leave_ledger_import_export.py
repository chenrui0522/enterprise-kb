"""Round-trip and reject tests for official export reverse-parse."""

from __future__ import annotations

import pytest

from app.core.errors import AppError
from app.leave_ledger.export import build_ledger_xlsx
from app.leave_ledger.import_export import (
    looks_like_official_export,
    parse_official_export_xlsx,
)


def _sample_rows() -> list[dict]:
    return [
        {
            "person_name": "甲",
            "center": "运营",
            "department": "计划",
            "travel_segments": [
                {
                    "system_start": "2026-06-01",
                    "system_end": "2026-06-30",
                    "actual_start": "2026-06-01",
                    "actual_end": "2026-06-30",
                    "credit": 2.0,
                }
            ],
            "ot_days": [
                {"date": "2026-06-19", "credit": 1.0, "holiday": True},
            ],
            "leave_segments": [
                {"start": "2026-07-01", "end": "2026-07-01", "used_days": 1.0},
            ],
            "travel_comp_days": 2.0,
            "ot_comp_days": 1.0,
            "total_comp_days": 3.0,
            "used_comp_days": 1.0,
            "remaining_comp_days": 2.0,
            "remark": "跨月",
        },
        {
            "person_name": "乙",
            "center": "",
            "department": "",
            "travel_segments": [],
            "ot_days": [{"date": "2026-05-01", "credit": 0.5}],
            "leave_segments": [],
            "travel_comp_days": 0.0,
            "ot_comp_days": 0.5,
            "total_comp_days": 0.5,
            "used_comp_days": 0.0,
            "remaining_comp_days": 0.5,
            "remark": "",
        },
    ]


def test_looks_like_and_round_trip() -> None:
    rows = _sample_rows()
    data = build_ledger_xlsx(rows)
    assert looks_like_official_export(data)
    parsed = parse_official_export_xlsx(data)
    assert len(parsed) == 2
    a = next(p for p in parsed if p["person_name"] == "甲")
    assert len(a["travel_segments"]) == 1
    assert a["travel_segments"][0]["actual_start"] == "2026-06-01"
    assert float(a["travel_segments"][0]["credit"]) == 2.0
    assert len(a["ot_days"]) == 1
    assert a["ot_days"][0]["date"] == "2026-06-19"
    assert float(a["ot_days"][0]["credit"]) == 1.0
    assert len(a["leave_segments"]) == 1
    assert float(a["total_comp_days"]) == 3.0
    assert float(a["remaining_comp_days"]) == 2.0
    b = next(p for p in parsed if p["person_name"] == "乙")
    assert float(b["ot_comp_days"]) == 0.5


def test_reject_non_template() -> None:
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名"
    ws["B1"] = "随便"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    assert not looks_like_official_export(data)
    with pytest.raises(AppError) as ei:
        parse_official_export_xlsx(data)
    assert ei.value.status_code == 400
