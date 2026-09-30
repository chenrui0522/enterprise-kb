"""Staffing import / confirm / void / summary services."""

from __future__ import annotations

from datetime import date, datetime, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError, NotFoundError
from app.core.logging import get_logger, log_event, preview_text, timed_event
from app.identity.principal import Principal
from app.models.entity import AuditEvent, new_id
from app.models.identity import Project
from app.models.staffing import StaffingAttendanceFact, StaffingImportBatch
from app.staffing.match import match_project, match_to_dict
from app.staffing.parse import _suspected_name_split_pairs, parse_daily_report, parse_result_to_dict
from app.staffing.roster import (
    classify_name_against_roster,
    correct_names_against_roster,
    get_roster,
)
from app.staffing.stints import annotate_stints_with_stages, build_stints, stage_day_counts

logger = get_logger("staffing")


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _actor(principal: Principal) -> str:
    return principal.username or principal.user_id


async def _add_audit(
    session: AsyncSession,
    *,
    tenant_id: str,
    actor: str,
    action: str,
    resource_type: str,
    resource_id: str | None,
    detail: dict | None,
) -> None:
    session.add(
        AuditEvent(
            tenant_id=tenant_id,
            actor=actor,
            action=action,
            resource_type=resource_type,
            resource_id=resource_id,
            detail=detail,
        )
    )


async def ensure_project_access(principal: Principal, project_id: str) -> None:
    if project_id in principal.project_ids:
        return
    if "projects:manage" in principal.permissions:
        return
    raise AppError("无权访问该项目", status_code=403)


async def create_import_batch(
    session: AsyncSession,
    *,
    principal: Principal,
    filename: str,
    data: bytes,
    storage_key: str,
    batch_id: str | None = None,
) -> StaffingImportBatch:
    with timed_event(logger, "staffing.import.parse") as meta:
        result = parse_daily_report(data, filename=filename)
        meta["row_count"] = len(result.rows)
        meta["warning_count"] = len(result.warnings)
        meta["filename"] = preview_text(filename, limit=80)

    with timed_event(logger, "staffing.import.match_project") as meta:
        match = await match_project(
            session,
            principal.tenant_id,
            principal.project_ids,
            result.filename_codes,
        )
        meta["status"] = match.status
        meta["codes"] = result.filename_codes

    from dataclasses import asdict

    parse_dict = parse_result_to_dict(result)
    corrected_rows, roster_warnings = correct_names_against_roster(parse_dict["rows"])
    parse_dict["rows"] = corrected_rows
    corrected_from = {
        (w.get("detail") or {}).get("from_name")
        for w in roster_warnings
        if w.get("code") == "roster_name_corrected"
    }
    warnings = [
        asdict(w)
        for w in result.warnings
        if not (
            w.code == "suspected_name_split"
            and (w.detail or {}).get("short_name") in corrected_from
        )
    ]
    warnings.extend(roster_warnings)

    blocking_codes = {
        w.get("code")
        for w in warnings
        if isinstance(w, dict) and w.get("code") != "roster_name_corrected"
    }
    needs_review = bool(blocking_codes) or match.status != "matched"
    if match.status == "unmatched":
        warnings.append(
            {
                "code": "project_unmatched",
                "message": "未能从文件名/标题匹配到已授权项目",
                "row": None,
                "detail": {"codes": result.filename_codes},
            }
        )
        needs_review = True
    elif match.status == "ambiguous":
        warnings.append(
            {
                "code": "project_ambiguous",
                "message": "匹配到多个已授权项目，需人工选定",
                "row": None,
                "detail": {"candidates": match.candidates},
            }
        )
        needs_review = True

    batch = StaffingImportBatch(
        id=batch_id or new_id(),
        tenant_id=principal.tenant_id,
        project_id=match.project_id,
        status="needs_review" if needs_review else "parsed",
        filename=filename,
        storage_key=storage_key,
        uploaded_by=principal.user_id,
        warnings=warnings,
        parse_result=parse_dict,
        resolved_warnings={},
        match_info=match_to_dict(match),
    )
    session.add(batch)
    await _add_audit(
        session,
        tenant_id=principal.tenant_id,
        actor=_actor(principal),
        action="staffing.import",
        resource_type="staffing_batch",
        resource_id=batch.id,
        detail={
            "filename": filename,
            "warning_count": len(warnings),
            "match": match.status,
            "person_days": len(parse_dict.get("rows") or []),
        },
    )
    await session.commit()
    await session.refresh(batch)

    if needs_review:
        log_event(
            logger,
            "staffing.review.pending",
            event="staffing.review.pending",
            batch_id=batch.id,
            issue_codes=sorted({w["code"] for w in warnings}),
        )
    return batch


async def get_batch(
    session: AsyncSession, tenant_id: str, batch_id: str
) -> StaffingImportBatch:
    batch = await session.get(StaffingImportBatch, batch_id)
    if batch is None or batch.tenant_id != tenant_id:
        raise NotFoundError("导入批次不存在")
    return batch


async def review_batch(
    session: AsyncSession,
    *,
    principal: Principal,
    batch_id: str,
    project_id: str | None = None,
    resolutions: dict[str, Any] | None = None,
    person_renames: dict[str, str] | None = None,
) -> StaffingImportBatch:
    batch = await get_batch(session, principal.tenant_id, batch_id)
    if batch.status in {"confirmed", "voided"}:
        raise AppError("已确认或已作废的批次不可再复核", status_code=409)

    if project_id:
        await ensure_project_access(principal, project_id)
        proj = await session.get(Project, project_id)
        if proj is None or proj.tenant_id != principal.tenant_id:
            raise AppError("项目不存在", status_code=404)
        batch.project_id = project_id

    resolved = dict(batch.resolved_warnings or {})
    if resolutions:
        resolved.update(resolutions)
    batch.resolved_warnings = resolved

    if resolutions and "roster_name_unresolved" in resolutions:
        _apply_roster_resolutions_to_parse_result(batch, resolutions["roster_name_unresolved"])

    if person_renames and batch.parse_result:
        rows = batch.parse_result.get("rows") or []
        for row in rows:
            old = row.get("person_name")
            if old in person_renames:
                row["person_name"] = person_renames[old]
        batch.parse_result = {**batch.parse_result, "rows": rows}
        # Drop split warnings that were fixed by renaming.
        kept = []
        for warning in batch.warnings or []:
            if not isinstance(warning, dict):
                kept.append(warning)
                continue
            if warning.get("code") != "suspected_name_split":
                kept.append(warning)
                continue
            detail = warning.get("detail") if isinstance(warning.get("detail"), dict) else {}
            short = detail.get("short_name")
            long = detail.get("long_name")
            if short in person_renames or long in person_renames.values() or long in person_renames:
                continue
            kept.append(warning)
        batch.warnings = kept

    # Determine remaining open warnings
    open_codes = _open_warning_codes_from(batch.warnings or [], resolved, batch.project_id)

    if batch.project_id and not open_codes:
        batch.status = "parsed"
    else:
        batch.status = "needs_review"

    await _add_audit(
        session,
        tenant_id=principal.tenant_id,
        actor=_actor(principal),
        action="staffing.review.resolve",
        resource_type="staffing_batch",
        resource_id=batch.id,
        detail={"project_id": batch.project_id, "resolved": list(resolved.keys())},
    )
    await session.commit()
    await session.refresh(batch)
    return batch


def _collision_resolutions(resolved: dict[str, Any]) -> dict[str, Any]:
    raw = resolved.get("name_kind_collision")
    return raw if isinstance(raw, dict) else {}


def _roster_unresolved_resolutions(resolved: dict[str, Any]) -> dict[str, Any]:
    raw = resolved.get("roster_name_unresolved")
    return raw if isinstance(raw, dict) else {}


def _is_roster_unresolved_resolved(warnings: list, resolved: dict[str, Any]) -> bool:
    unresolved = [
        w
        for w in warnings
        if isinstance(w, dict) and w.get("code") == "roster_name_unresolved"
    ]
    if not unresolved:
        return True
    by_token = _roster_unresolved_resolutions(resolved)
    for w in unresolved:
        detail = w.get("detail") if isinstance(w.get("detail"), dict) else {}
        token = detail.get("token")
        entry = by_token.get(token) if token else None
        if not isinstance(entry, dict):
            return False
        action = entry.get("action")
        if action == "discard":
            continue
        if action in {"select_roster_name", "rename"} and (entry.get("name") or "").strip():
            continue
        return False
    return True


def _apply_roster_resolutions_to_parse_result(
    batch: StaffingImportBatch, roster_res: Any
) -> None:
    """Rewrite or drop pending rows for resolved roster tokens."""
    if not isinstance(roster_res, dict) or not batch.parse_result:
        return
    roster = get_roster()
    rows_out: list[dict[str, Any]] = []
    for row in list(batch.parse_result.get("rows") or []):
        name = (row.get("person_name") or "").strip()
        entry = roster_res.get(name)
        if not isinstance(entry, dict):
            rows_out.append(row)
            continue
        action = entry.get("action")
        if action == "discard":
            continue
        new_name = (entry.get("name") or "").strip()
        if action == "select_roster_name":
            if new_name not in roster:
                raise AppError(
                    f"选定姓名「{new_name}」不在花名册中",
                    status_code=400,
                )
            rows_out.append({**row, "person_name": new_name})
        elif action == "rename":
            if not new_name:
                raise AppError("改名不能为空", status_code=400)
            updated = {**row, "person_name": new_name}
            if new_name not in roster:
                entry["acknowledged_off_roster"] = True
            rows_out.append(updated)
        else:
            raise AppError(f"未知花名册决议：{action}", status_code=400)
    batch.parse_result = {**batch.parse_result, "rows": rows_out}


def _rows_after_roster_resolution(
    rows: list[dict[str, Any]], resolved: dict[str, Any]
) -> list[dict[str, Any]]:
    """Safety net on confirm: apply unresolved resolutions if rows still have tokens."""
    by_token = _roster_unresolved_resolutions(resolved)
    if not by_token:
        return rows
    out: list[dict[str, Any]] = []
    for r in rows:
        name = (r.get("person_name") or "").strip()
        entry = by_token.get(name)
        if not isinstance(entry, dict):
            out.append(r)
            continue
        action = entry.get("action")
        if action == "discard":
            continue
        new_name = (entry.get("name") or "").strip()
        if action in {"select_roster_name", "rename"} and new_name:
            out.append({**r, "person_name": new_name})
        else:
            out.append(r)
    return out


def _is_collision_resolved(warnings: list, resolved: dict[str, Any]) -> bool:
    collisions = [
        w
        for w in warnings
        if isinstance(w, dict) and w.get("code") == "name_kind_collision"
    ]
    if not collisions:
        return True
    by_name = _collision_resolutions(resolved)
    for w in collisions:
        detail = w.get("detail") if isinstance(w.get("detail"), dict) else {}
        name = detail.get("person_name")
        entry = by_name.get(name) if name else None
        if not isinstance(entry, dict):
            return False
        action = entry.get("action")
        if action == "split":
            continue
        if action == "merge" and entry.get("keep_kind") in {
            "internal_formal",
            "internal_contract",
        }:
            continue
        return False
    return True


def _open_warning_codes_from(
    warnings: list, resolved: dict[str, Any], project_id: str | None
) -> set[str]:
    open_codes: set[str] = set()
    for w in warnings:
        if not isinstance(w, dict):
            continue
        code = w.get("code")
        if not code:
            continue
        if code == "roster_name_corrected":
            # Informational auto-fix; never blocks confirm.
            continue
        if code == "roster_unavailable":
            # Must re-import after configuring roster; acknowledge is not enough.
            open_codes.add(code)
            continue
        if code in {"project_unmatched", "project_ambiguous"} and project_id:
            continue
        if code == "name_kind_collision":
            if not _is_collision_resolved(warnings, resolved):
                open_codes.add(code)
            continue
        if code == "roster_name_unresolved":
            if not _is_roster_unresolved_resolved(warnings, resolved):
                open_codes.add(code)
            continue
        if code not in resolved:
            open_codes.add(code)
    return open_codes


def _open_warning_codes(batch: StaffingImportBatch) -> set[str]:
    return _open_warning_codes_from(
        batch.warnings or [],
        batch.resolved_warnings or {},
        batch.project_id,
    )


def _rows_after_collision_resolution(
    rows: list[dict[str, Any]], resolved: dict[str, Any]
) -> list[dict[str, Any]]:
    by_name = _collision_resolutions(resolved)
    out: list[dict[str, Any]] = []
    for r in rows:
        name = r.get("person_name")
        kind = r.get("person_kind") or "internal_formal"
        entry = by_name.get(name) if name else None
        if isinstance(entry, dict) and entry.get("action") == "merge":
            if kind != entry.get("keep_kind"):
                continue
        out.append(r)
    return out


async def confirm_batch(
    session: AsyncSession, *, principal: Principal, batch_id: str
) -> StaffingImportBatch:
    batch = await get_batch(session, principal.tenant_id, batch_id)
    if batch.status == "confirmed":
        raise AppError("批次已确认", status_code=409)
    if batch.status == "voided":
        raise AppError("批次已作废", status_code=409)
    if not batch.project_id:
        raise AppError("请先选定项目后再确认", status_code=400)
    await ensure_project_access(principal, batch.project_id)
    open_codes = _open_warning_codes(batch)
    if open_codes:
        raise AppError(
            f"仍有未决议告警：{', '.join(sorted(open_codes))}",
            status_code=400,
        )

    rows = _rows_after_collision_resolution(
        list((batch.parse_result or {}).get("rows") or []),
        batch.resolved_warnings or {},
    )
    rows = _rows_after_roster_resolution(rows, batch.resolved_warnings or {})
    dates = {r["work_date"] for r in rows}
    date_objs = [date.fromisoformat(d) for d in dates]

    with timed_event(logger, "staffing.import.confirm") as meta:
        voided = 0
        if date_objs:
            existing = (
                await session.execute(
                    select(StaffingAttendanceFact).where(
                        StaffingAttendanceFact.tenant_id == principal.tenant_id,
                        StaffingAttendanceFact.project_id == batch.project_id,
                        StaffingAttendanceFact.status == "active",
                        StaffingAttendanceFact.work_date.in_(date_objs),
                    )
                )
            ).scalars().all()
            for fact in existing:
                fact.status = "voided"
                fact.void_reason = "superseded_by_batch"
                fact.voided_by = principal.user_id
                fact.voided_at = _utcnow()
                voided += 1

        # dedupe by date + person + kind
        seen: set[tuple[str, str, str]] = set()
        written = 0
        for r in rows:
            kind = r.get("person_kind") or "internal_formal"
            key = (r["work_date"], r["person_name"], kind)
            if key in seen:
                continue
            seen.add(key)
            session.add(
                StaffingAttendanceFact(
                    tenant_id=principal.tenant_id,
                    project_id=batch.project_id,
                    batch_id=batch.id,
                    work_date=date.fromisoformat(r["work_date"]),
                    person_name=r["person_name"],
                    person_kind=kind,
                    status="active",
                    stage=r.get("stage"),
                    source_row=r.get("source_row"),
                )
            )
            written += 1

        batch.status = "confirmed"
        batch.confirmed_at = _utcnow()
        batch.confirmed_by = principal.user_id
        meta["facts_written"] = written
        meta["dates_voided_overlap"] = voided
        meta["batch_id"] = batch.id

        await _add_audit(
            session,
            tenant_id=principal.tenant_id,
            actor=_actor(principal),
            action="staffing.confirm",
            resource_type="staffing_batch",
            resource_id=batch.id,
            detail={
                "project_id": batch.project_id,
                "facts_written": written,
                "dates_voided_overlap": voided,
            },
        )
        await session.commit()
        await session.refresh(batch)
    return batch


KIND_LABEL = {
    "internal_formal": "正式我司",
    "internal_contract": "机电服务处",
}

# Formal first, then contract; within kind by name.
_KIND_SORT = {"internal_formal": 0, "internal_contract": 1}


async def project_summary(
    session: AsyncSession, *, principal: Principal, project_id: str
) -> dict[str, Any]:
    await ensure_project_access(principal, project_id)
    proj = await session.get(Project, project_id)
    facts = (
        await session.execute(
            select(StaffingAttendanceFact).where(
                StaffingAttendanceFact.tenant_id == principal.tenant_id,
                StaffingAttendanceFact.project_id == project_id,
                StaffingAttendanceFact.status == "active",
            )
        )
    ).scalars().all()

    by_person: dict[tuple[str, str], dict[str, Any]] = {}
    for f in facts:
        key = (f.person_name, f.person_kind)
        entry = by_person.setdefault(
            key,
            {
                "person_name": f.person_name,
                "person_kind": f.person_kind,
                "dates": set(),
                "date_stages": {},
            },
        )
        iso = f.work_date.isoformat()
        entry["dates"].add(iso)
        stage = (f.stage or "").strip() or None
        if stage and not entry["date_stages"].get(iso):
            entry["date_stages"][iso] = stage
        elif stage and entry["date_stages"].get(iso) and entry["date_stages"][iso] != stage:
            # Keep first seen; conflicting stages on same day are rare after confirm.
            pass
        elif iso not in entry["date_stages"]:
            entry["date_stages"][iso] = stage

    people = []
    all_dates: set[str] = set()
    by_kind_days: dict[str, set[str]] = {
        "internal_formal": set(),
        "internal_contract": set(),
    }
    by_kind_people: dict[str, set[str]] = {
        "internal_formal": set(),
        "internal_contract": set(),
    }
    person_day_total = 0
    for (name, kind), entry in sorted(
        by_person.items(),
        key=lambda item: (_KIND_SORT.get(item[0][1], 9), item[0][0]),
    ):
        dates = sorted(entry["dates"])
        date_stages = {
            d: entry["date_stages"].get(d) for d in dates
        }
        person_day_total += len(dates)
        all_dates.update(dates)
        by_kind_people.setdefault(kind, set()).add(name)
        # person-days per kind: count this person's days under this kind
        by_kind_days.setdefault(kind, set())
        # use synthetic keys for person-day totals per kind
        for d in dates:
            by_kind_days[kind].add(f"{name}:{d}")
        stints = annotate_stints_with_stages(build_stints(dates), date_stages)
        stages = stage_day_counts(date_stages)
        people.append(
            {
                "person_name": name,
                "person_kind": kind,
                "person_kind_label": KIND_LABEL.get(kind, kind),
                "days_on_site": len(dates),
                "dates": dates,
                "date_stages": date_stages,
                "stages": stages,
                "stint_count": len(stints),
                "stints": stints,
            }
        )

    sorted_dates = sorted(all_dates)
    name_split_suspects = [
        {
            "short_name": short,
            "long_name": long,
            "message": (
                f"「{short}」与「{long}」疑似同一人被拆成两条"
                "（常见于 PDF 换行/抽表），请复核后合并"
            ),
        }
        for short, long in _suspected_name_split_pairs(
            [p["person_name"] for p in people]
        )
    ]
    roster = get_roster()
    roster_name_suspects: list[dict[str, Any]] = []
    seen_tokens: set[str] = set()
    for p in people:
        token = p["person_name"]
        if token in seen_tokens:
            continue
        seen_tokens.add(token)
        hit = classify_name_against_roster(token, roster=roster)
        if hit["status"] == "exact":
            continue
        if hit["status"] == "unavailable":
            roster_name_suspects.append(
                {
                    "token": token,
                    "status": "unavailable",
                    "person_kind": p["person_kind"],
                    "person_kind_label": p["person_kind_label"],
                    "message": "花名册不可用，无法校验该姓名",
                    "candidates": [],
                }
            )
            continue
        if hit["status"] == "auto":
            roster_name_suspects.append(
                {
                    "token": token,
                    "status": "auto",
                    "to_name": hit["to_name"],
                    "person_kind": p["person_kind"],
                    "person_kind_label": p["person_kind_label"],
                    "message": (
                        f"「{token}」可按花名册自动纠正为「{hit['to_name']}」"
                        f"（{p['person_kind_label']}）"
                    ),
                    "candidates": hit.get("candidates") or [],
                }
            )
        else:
            roster_name_suspects.append(
                {
                    "token": token,
                    "status": "unresolved",
                    "person_kind": p["person_kind"],
                    "person_kind_label": p["person_kind_label"],
                    "message": (
                        f"「{token}」无法在花名册中唯一确认"
                        f"（{p['person_kind_label']}），请人工裁定"
                    ),
                    "candidates": hit.get("candidates") or [],
                }
            )
    return {
        "project_id": project_id,
        "project_code": proj.code if proj else "",
        "project_name": proj.name if proj else "",
        "people": people,
        "person_count": len(people),
        "person_day_total": person_day_total,
        "date_from": sorted_dates[0] if sorted_dates else None,
        "date_to": sorted_dates[-1] if sorted_dates else None,
        "by_kind": {
            "formal": {
                "person_count": len(by_kind_people.get("internal_formal") or ()),
                "person_day_total": len(by_kind_days.get("internal_formal") or ()),
            },
            "contract": {
                "person_count": len(by_kind_people.get("internal_contract") or ()),
                "person_day_total": len(by_kind_days.get("internal_contract") or ()),
            },
        },
        "name_split_suspects": name_split_suspects,
        "roster_name_suspects": roster_name_suspects,
        "label": "在场天数",
    }


async def repair_project_names_from_roster(
    session: AsyncSession, *, principal: Principal, project_id: str
) -> dict[str, Any]:
    """Auto-rename active facts with a unique roster prefix/suffix match.

    Applies to both 正式我司 and 机电服务处. Returns applied renames and still-
    unresolved tokens (for manual adjudication).
    """
    await ensure_project_access(principal, project_id)
    roster = get_roster()
    if not roster:
        raise AppError(
            "人员信息花名册不可用，请配置 KB_STAFFING_ROSTER_PATH 后重试",
            status_code=400,
        )

    facts = (
        await session.execute(
            select(StaffingAttendanceFact).where(
                StaffingAttendanceFact.tenant_id == principal.tenant_id,
                StaffingAttendanceFact.project_id == project_id,
                StaffingAttendanceFact.status == "active",
            )
        )
    ).scalars().all()
    tokens = sorted({(f.person_name or "").strip() for f in facts if f.person_name})
    applied: list[dict[str, str]] = []
    unresolved: list[dict[str, Any]] = []
    for token in tokens:
        hit = classify_name_against_roster(token, roster=roster)
        if hit["status"] == "exact":
            continue
        if hit["status"] == "auto":
            result = await merge_person_names(
                session,
                principal=principal,
                project_id=project_id,
                from_name=token,
                to_name=hit["to_name"],
            )
            applied.append(
                {
                    "from_name": token,
                    "to_name": hit["to_name"],
                    "renamed": int(result.get("renamed") or 0),
                    "voided_duplicates": int(result.get("voided_duplicates") or 0),
                }
            )
        else:
            unresolved.append(
                {
                    "token": token,
                    "candidates": hit.get("candidates") or [],
                    "message": f"「{token}」无法在花名册中唯一确认，请人工裁定",
                }
            )

    await _add_audit(
        session,
        tenant_id=principal.tenant_id,
        actor=_actor(principal),
        action="staffing.repair_roster_names",
        resource_type="project",
        resource_id=project_id,
        detail={"applied": applied, "unresolved_count": len(unresolved)},
    )
    await session.commit()
    return {
        "project_id": project_id,
        "applied": applied,
        "unresolved": unresolved,
        "applied_count": len(applied),
        "unresolved_count": len(unresolved),
    }


async def merge_person_names(
    session: AsyncSession,
    *,
    principal: Principal,
    project_id: str,
    from_name: str,
    to_name: str,
) -> dict[str, Any]:
    """Rename active facts from ``from_name`` to ``to_name`` within a project.

    When the long name already has an active fact on the same day+kind, the
    short-name fact is voided instead of violating the batch uniqueness.
    """
    await ensure_project_access(principal, project_id)
    short = (from_name or "").strip()
    long = (to_name or "").strip()
    if not short or not long or short == long:
        raise AppError("合并姓名无效", status_code=400)

    short_facts = list(
        (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.tenant_id == principal.tenant_id,
                    StaffingAttendanceFact.project_id == project_id,
                    StaffingAttendanceFact.status == "active",
                    StaffingAttendanceFact.person_name == short,
                )
            )
        ).scalars().all()
    )
    if not short_facts:
        raise AppError(f"未找到姓名「{short}」的有效出勤", status_code=404)

    long_keys = {
        (f.batch_id, f.work_date, f.person_kind)
        for f in (
            await session.execute(
                select(StaffingAttendanceFact).where(
                    StaffingAttendanceFact.tenant_id == principal.tenant_id,
                    StaffingAttendanceFact.project_id == project_id,
                    StaffingAttendanceFact.status == "active",
                    StaffingAttendanceFact.person_name == long,
                )
            )
        ).scalars().all()
    }

    renamed = 0
    voided = 0
    now = _utcnow()
    for fact in short_facts:
        key = (fact.batch_id, fact.work_date, fact.person_kind)
        if key in long_keys:
            fact.status = "voided"
            fact.void_reason = "merged_name_duplicate"
            fact.voided_by = principal.user_id
            fact.voided_at = now
            voided += 1
            continue
        fact.person_name = long
        long_keys.add(key)
        renamed += 1

    await _add_audit(
        session,
        tenant_id=principal.tenant_id,
        actor=_actor(principal),
        action="staffing.merge_names",
        resource_type="project",
        resource_id=project_id,
        detail={
            "from_name": short,
            "to_name": long,
            "renamed": renamed,
            "voided_duplicates": voided,
        },
    )
    await session.commit()
    return {
        "from_name": short,
        "to_name": long,
        "renamed": renamed,
        "voided_duplicates": voided,
    }


async def list_facts(
    session: AsyncSession,
    *,
    principal: Principal,
    project_id: str,
    include_voided: bool = False,
) -> list[StaffingAttendanceFact]:
    await ensure_project_access(principal, project_id)
    stmt = select(StaffingAttendanceFact).where(
        StaffingAttendanceFact.tenant_id == principal.tenant_id,
        StaffingAttendanceFact.project_id == project_id,
    )
    if not include_voided:
        stmt = stmt.where(StaffingAttendanceFact.status == "active")
    stmt = stmt.order_by(StaffingAttendanceFact.work_date, StaffingAttendanceFact.person_name)
    return list((await session.execute(stmt)).scalars().all())


async def void_facts(
    session: AsyncSession,
    *,
    principal: Principal,
    project_id: str,
    scope: str,
    dates: list[str] | None = None,
    person_names: list[str] | None = None,
    batch_id: str | None = None,
) -> int:
    await ensure_project_access(principal, project_id)
    stmt = select(StaffingAttendanceFact).where(
        StaffingAttendanceFact.tenant_id == principal.tenant_id,
        StaffingAttendanceFact.project_id == project_id,
        StaffingAttendanceFact.status == "active",
    )

    action = "staffing.void.day_person"
    if scope == "person":
        if not person_names:
            raise AppError("按人作废需要 person_names", status_code=400)
        stmt = stmt.where(StaffingAttendanceFact.person_name.in_(person_names))
        action = "staffing.void.person"
    elif scope == "batch":
        if not batch_id:
            raise AppError("按批作废需要 batch_id", status_code=400)
        stmt = stmt.where(StaffingAttendanceFact.batch_id == batch_id)
        action = "staffing.void.batch"
    elif scope == "day_person":
        if not dates or not person_names:
            raise AppError("按日×人作废需要 dates 与 person_names", status_code=400)
        stmt = stmt.where(
            StaffingAttendanceFact.work_date.in_([date.fromisoformat(d) for d in dates]),
            StaffingAttendanceFact.person_name.in_(person_names),
        )
        action = "staffing.void.day_person"
    elif scope == "day":
        if not dates:
            raise AppError("按日作废需要 dates", status_code=400)
        stmt = stmt.where(
            StaffingAttendanceFact.work_date.in_([date.fromisoformat(d) for d in dates])
        )
        action = "staffing.void.day"
    else:
        raise AppError(f"未知作废范围: {scope}", status_code=400)

    with timed_event(logger, "staffing.void") as meta:
        facts = list((await session.execute(stmt)).scalars().all())
        now = _utcnow()
        for f in facts:
            f.status = "voided"
            f.void_reason = f"void_{scope}"
            f.voided_by = principal.user_id
            f.voided_at = now
        if scope == "batch" and batch_id:
            batch = await get_batch(session, principal.tenant_id, batch_id)
            if batch.status == "confirmed":
                batch.status = "voided"
        meta["scope"] = scope
        meta["count"] = len(facts)
        meta["project_id"] = project_id

        await _add_audit(
            session,
            tenant_id=principal.tenant_id,
            actor=_actor(principal),
            action=action,
            resource_type="project",
            resource_id=project_id,
            detail={
                "scope": scope,
                "count": len(facts),
                "dates": dates,
                "person_names": person_names,
                "batch_id": batch_id,
            },
        )
        await session.commit()
    return len(facts)


# Columns match desktop「人员投入汇总表总表」layout (one row per person-day).
_EXPORT_HEADERS = [
    "序号",
    "日期",
    "姓名",
    "项目号",
    "项目阶段",
    "部门",
    "职务",
    "项目名称",
]
_EXPORT_DATE_FORMAT = "mm-dd-yy"
_EXPORT_COL_WIDTHS = {
    "A": 8.125,
    "B": 13.125,
    "C": 12.625,
    "D": 10.25,
    "E": 13.0,
    "F": 14.125,
    "G": 21.875,
    "H": 51.0,
}


async def export_project_xlsx(
    session: AsyncSession, *, principal: Principal, project_id: str
) -> bytes:
    """Build workbook matching 人员投入汇总表总表 (person × day rows)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from app.staffing.roster import get_roster, resolve_department_title

    await ensure_project_access(principal, project_id)
    proj = await session.get(Project, project_id)
    if proj is None or proj.tenant_id != principal.tenant_id:
        raise NotFoundError("项目不存在")

    facts = await list_facts(
        session, principal=principal, project_id=project_id, include_voided=False
    )
    roster = get_roster()
    project_code = proj.code or ""
    # Prefer numeric project code when the workbook stores 项目号 as number.
    try:
        project_no: Any = int(str(project_code).strip()) if str(project_code).strip() else ""
    except ValueError:
        project_no = project_code
    project_name = proj.name or ""

    sorted_facts = sorted(
        facts,
        key=lambda f: (f.work_date, f.person_name or "", _KIND_SORT.get(f.person_kind, 9)),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "人员投入汇总表"
    ws.append(list(_EXPORT_HEADERS))

    header_font = Font(size=11.25)
    header_fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for col in range(1, len(_EXPORT_HEADERS) + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    body_font = Font(size=11.25)
    for idx, f in enumerate(sorted_facts, start=1):
        dept, title = resolve_department_title(
            f.person_name or "", f.person_kind, roster=roster
        )
        ws.append(
            [
                idx,
                f.work_date,
                f.person_name or "",
                project_no,
                (f.stage or "").strip(),
                dept,
                title,
                project_name,
            ]
        )
        for col in range(1, len(_EXPORT_HEADERS) + 1):
            cell = ws.cell(ws.max_row, col)
            cell.font = body_font
            if col == 2:
                cell.number_format = _EXPORT_DATE_FORMAT

    for letter, width in _EXPORT_COL_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
