"""Company roster for export 部门/职务 and import-time name correction.

Expects an xlsx with headers including 姓名、部门、工种 (工种 maps to 职务).
Configured via ``KB_STAFFING_ROSTER_PATH`` (default ``./data/staffing_roster.xlsx``).
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.config import get_settings
from app.core.logging import get_logger

logger = get_logger("staffing.roster")

KIND_CONTRACT = "internal_contract"
MAX_NAME_DELTA = 2
MAX_CANDIDATES_IN_WARNING = 20


@dataclass(frozen=True)
class RosterEntry:
    name: str
    department: str
    title: str
    left: bool = False


def _norm_header(value: Any) -> str:
    return str(value or "").strip().replace(" ", "")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_roster_from_path(path: Path) -> dict[str, RosterEntry]:
    """Parse roster workbook → ``{姓名: RosterEntry}`` (prefer non-离职 rows)."""
    if not path.is_file():
        return {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return {}
        cols = {_norm_header(h): i for i, h in enumerate(header) if h is not None}
        name_i = cols.get("姓名")
        dept_i = cols.get("部门")
        title_i = cols.get("工种") if "工种" in cols else cols.get("职务")
        note_i = cols.get("备注")
        if name_i is None:
            return {}

        out: dict[str, RosterEntry] = {}
        for row in rows_iter:
            if not row:
                continue
            name = _cell_str(row[name_i] if name_i < len(row) else None)
            if not name:
                continue
            note = _cell_str(
                row[note_i] if note_i is not None and note_i < len(row) else None
            )
            left = "离职" in note
            dept = _cell_str(
                row[dept_i] if dept_i is not None and dept_i < len(row) else None
            )
            title = _cell_str(
                row[title_i] if title_i is not None and title_i < len(row) else None
            )
            entry = RosterEntry(name=name, department=dept, title=title, left=left)
            prev = out.get(name)
            if prev is None or (prev.left and not left):
                out[name] = entry
        return out
    finally:
        wb.close()


@lru_cache(maxsize=1)
def _cached_roster(path_str: str, mtime_ns: int) -> dict[str, RosterEntry]:
    return load_roster_from_path(Path(path_str))


def clear_roster_cache() -> None:
    _cached_roster.cache_clear()


def get_roster() -> dict[str, RosterEntry]:
    settings = get_settings()
    raw = (settings.staffing_roster_path or "").strip()
    if not raw:
        return {}
    path = Path(raw)
    if not path.is_file():
        logger.debug("staffing_roster_missing", extra={"path": str(path)})
        return {}
    try:
        mtime_ns = path.stat().st_mtime_ns
    except OSError:
        return {}
    return _cached_roster(str(path.resolve()), mtime_ns)


def resolve_department_title(
    person_name: str,
    person_kind: str | None,
    *,
    roster: dict[str, RosterEntry] | None = None,
) -> tuple[str, str]:
    """Return (部门, 职务). Roster wins; contract kind falls back to 机电服务处."""
    book = roster if roster is not None else get_roster()
    entry = book.get((person_name or "").strip())
    if entry:
        return entry.department or "", entry.title or ""
    if person_kind == KIND_CONTRACT:
        return "机电服务处", ""
    return "", ""


def _prefix_suffix_candidates(token: str, roster_names: list[str]) -> list[str]:
    """Roster full names that uniquely extend ``token`` within ``MAX_NAME_DELTA``."""
    hits: list[str] = []
    for name in roster_names:
        extra = len(name) - len(token)
        if extra < 1 or extra > MAX_NAME_DELTA:
            continue
        if name.startswith(token) or name.endswith(token):
            hits.append(name)
    return hits


def classify_name_against_roster(
    token: str,
    *,
    roster: dict[str, RosterEntry] | None = None,
) -> dict[str, Any]:
    """Classify one display name vs roster.

    Returns ``{status, token, to_name?, candidates?}`` where status is
    ``exact`` | ``auto`` | ``unresolved`` | ``unavailable``.
    """
    book = roster if roster is not None else get_roster()
    name = (token or "").strip()
    if not book:
        return {"status": "unavailable", "token": name, "candidates": []}
    if not name:
        return {"status": "exact", "token": name, "to_name": name, "candidates": []}
    if name in book:
        return {"status": "exact", "token": name, "to_name": name, "candidates": []}
    candidates = _prefix_suffix_candidates(name, list(book.keys()))
    if len(candidates) == 1:
        return {
            "status": "auto",
            "token": name,
            "to_name": candidates[0],
            "candidates": candidates,
        }
    return {
        "status": "unresolved",
        "token": name,
        "candidates": candidates[:MAX_CANDIDATES_IN_WARNING],
    }


def correct_names_against_roster(
    rows: list[dict[str, Any]],
    *,
    roster: dict[str, RosterEntry] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Rewrite ``person_name`` using roster; return ``(rows, warning_dicts)``.

    Empty/missing roster → rows unchanged + blocking ``roster_unavailable``.
    Exact hit → keep. Unique prefix/suffix (len delta 1–2) → auto-correct.
    Otherwise → blocking ``roster_name_unresolved``.
    """
    book = roster if roster is not None else get_roster()
    if not book:
        return list(rows), [
            {
                "code": "roster_unavailable",
                "message": (
                    "人员信息花名册不可用或为空，"
                    "请配置 KB_STAFFING_ROSTER_PATH（默认 ./data/staffing_roster.xlsx）后重新导入"
                ),
                "row": None,
                "detail": {},
            }
        ]

    roster_names = list(book.keys())
    tokens = sorted(
        {
            (r.get("person_name") or "").strip()
            for r in rows
            if (r.get("person_name") or "").strip()
        }
    )
    rename_map: dict[str, str] = {}
    warnings: list[dict[str, Any]] = []

    for token in tokens:
        if token in book:
            continue
        candidates = _prefix_suffix_candidates(token, roster_names)
        if len(candidates) == 1:
            target = candidates[0]
            rename_map[token] = target
            warnings.append(
                {
                    "code": "roster_name_corrected",
                    "message": f"「{token}」已按花名册自动纠正为「{target}」",
                    "row": None,
                    "detail": {"from_name": token, "to_name": target},
                }
            )
        else:
            warnings.append(
                {
                    "code": "roster_name_unresolved",
                    "message": f"「{token}」无法在花名册中唯一确认，请人工复核",
                    "row": None,
                    "detail": {
                        "token": token,
                        "candidates": candidates[:MAX_CANDIDATES_IN_WARNING],
                    },
                }
            )

    if not rename_map:
        return list(rows), warnings

    out: list[dict[str, Any]] = []
    for row in rows:
        name = (row.get("person_name") or "").strip()
        if name in rename_map:
            out.append({**row, "person_name": rename_map[name]})
        else:
            out.append(row)
    return out, warnings
