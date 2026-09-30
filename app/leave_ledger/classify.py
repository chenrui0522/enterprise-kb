"""Classify leave-ledger source workbooks by filename and header fingerprint."""

from __future__ import annotations

import io
import zipfile
from typing import Literal

from openpyxl import load_workbook

from app.core.errors import AppError

Role = Literal["travel", "overtime", "leave", "punch"]

ROLES: tuple[Role, ...] = ("travel", "overtime", "leave", "punch")

_FILENAME_KEYWORDS: list[tuple[Role, tuple[str, ...]]] = [
    ("travel", ("出差申请", "出差")),
    ("overtime", ("加班申请", "加班")),
    ("leave", ("请假申请", "请假")),
    ("punch", ("打卡", "上下班")),
]

_HEADER_FINGERPRINTS: list[tuple[Role, frozenset[str]]] = [
    ("travel", frozenset({"出差类别", "出差事由"})),
    ("overtime", frozenset({"加班明细", "加班时长", "加班事由"})),
    ("leave", frozenset({"假种", "请假类型", "请假时长"})),
    ("punch", frozenset({"打卡地址", "打卡时间", "地点"})),
]


def looks_like_ooxml_xlsx(data: bytes) -> bool:
    if len(data) < 4 or data[:2] != b"PK":
        return False
    try:
        return zipfile.is_zipfile(io.BytesIO(data))
    except Exception:
        return False


def _header_cells(data: bytes, max_rows: int = 5) -> set[str]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        found: set[str] = set()
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                for cell in row:
                    if cell is None:
                        continue
                    text = str(cell).strip()
                    if text:
                        found.add(text)
        return found
    finally:
        wb.close()


def classify_by_filename(filename: str) -> Role | None:
    name = filename or ""
    hits: list[Role] = []
    for role, keys in _FILENAME_KEYWORDS:
        if any(k in name for k in keys):
            hits.append(role)
    if len(hits) == 1:
        return hits[0]
    # Prefer more specific overtime/travel over generic if both matched somehow
    if "加班" in name and "出差" not in name:
        return "overtime"
    if "出差" in name and "加班" not in name:
        return "travel"
    if "请假" in name:
        return "leave"
    if "打卡" in name:
        return "punch"
    return None if len(hits) != 1 else hits[0]


def classify_by_headers(headers: set[str]) -> Role | None:
    scores: list[tuple[int, Role]] = []
    for role, keys in _HEADER_FINGERPRINTS:
        hit = sum(1 for k in keys if any(k in h for h in headers))
        if hit:
            scores.append((hit, role))
    if not scores:
        return None
    scores.sort(key=lambda x: (-x[0], x[1]))
    best_score, best_role = scores[0]
    if len(scores) > 1 and scores[1][0] == best_score:
        return None
    return best_role


def classify_source(filename: str, data: bytes) -> Role:
    if not looks_like_ooxml_xlsx(data):
        raise AppError("无法解析为标准 Excel 工作簿，请另存为 xlsx 后再上传", status_code=400)
    by_name = classify_by_filename(filename)
    if by_name:
        return by_name
    try:
        headers = _header_cells(data)
    except Exception as exc:  # noqa: BLE001
        raise AppError("无法读取 Excel 表头，请另存为标准工作簿", status_code=400) from exc
    by_header = classify_by_headers(headers)
    if by_header:
        return by_header
    raise AppError(
        f"无法识别文件角色（出差/加班/请假/打卡）：{filename}",
        status_code=400,
    )
