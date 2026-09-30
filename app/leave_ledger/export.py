"""Export leave-ledger snapshot rows to Taiyuan multi-row template xlsx."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Human-facing header labels (row 1 / row 2). Machine tests assert these.
HEADER_ROW1 = [
    "姓名",
    "中心",
    "部门",
    "出差日期（系统提报）",
    None,
    "出差日期（实际）",
    None,
    "可调休天数",
    "加班日期",
    None,
    "可调休天数",
    "合计可调休天数",
    "调休日期",
    None,
    "已休调休天数",
    "合计已休天数",
    "剩余调休天数",
    "备注",
]

HEADER_ROW2 = [
    None,
    None,
    None,
    "起始日",
    "结束日",
    "起始日",
    "结束日",
    None,
    "加班日期（开始时间）",
    "加班日期（结束时间）",
    None,
    None,
    "起始日",
    "结束日",
    None,
    None,
    None,
    None,
]

# Kept for tests that still reference a flat list of logical columns.
EXPORT_HEADERS = [
    "姓名",
    "中心",
    "部门",
    "出差日期（系统提报）起始日",
    "出差日期（系统提报）结束日",
    "出差日期（实际）起始日",
    "出差日期（实际）结束日",
    "出差可调休天数",
    "加班日期（开始时间）",
    "加班日期（结束时间）",
    "加班可调休天数",
    "合计可调休天数",
    "调休日期起始日",
    "调休日期结束日",
    "已休调休天数",
    "合计已休天数",
    "剩余调休天数",
    "备注",
]

_YELLOW = PatternFill("solid", fgColor="FFFF00")
_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _cell(ws, row: int, col: int, value: Any = None) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.border = _THIN
    c.alignment = _CENTER


def build_ledger_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "调休台账"

    for col, val in enumerate(HEADER_ROW1, start=1):
        _cell(ws, 1, col, val)
    for col, val in enumerate(HEADER_ROW2, start=1):
        _cell(ws, 2, col, val)

    # Vertical merges for single-column headers; horizontal for grouped headers.
    for col in (1, 2, 3, 8, 11, 12, 15, 16, 17, 18):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=5)
    ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=7)
    ws.merge_cells(start_row=1, start_column=9, end_row=1, end_column=10)
    ws.merge_cells(start_row=1, start_column=13, end_row=1, end_column=14)

    ws.cell(1, 12).fill = _YELLOW
    ws.cell(1, 16).fill = _YELLOW
    for col in range(1, 19):
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(2, col).font = Font(bold=True)

    cursor = 3
    for person in rows:
        travel = list(person.get("travel_segments") or [])
        ot_days = [
            d
            for d in (person.get("ot_days") or [])
            if float(d.get("credit") or 0) > 0
        ]
        leaves = list(person.get("leave_segments") or [])
        # Fallback: synthesize one segment from aggregates when older snapshots lack lists.
        if not travel and (
            person.get("system_travel_start") or person.get("actual_travel_start")
        ):
            travel = [
                {
                    "system_start": person.get("system_travel_start") or "",
                    "system_end": person.get("system_travel_end") or "",
                    "actual_start": person.get("actual_travel_start") or "",
                    "actual_end": person.get("actual_travel_end") or "",
                    "credit": person.get("travel_comp_days") or 0,
                }
            ]
        if not ot_days and float(person.get("ot_comp_days") or 0) > 0:
            # Best-effort: one placeholder row from joined dates if present.
            raw = (person.get("ot_dates") or "").split("、")
            first = (raw[0] or "").replace("(法定)", "").strip() if raw else ""
            if first:
                ot_days = [
                    {
                        "date": first,
                        "credit": person.get("ot_comp_days") or 0,
                    }
                ]
        if not leaves and (
            person.get("leave_start") or float(person.get("used_comp_days") or 0)
        ):
            leaves = [
                {
                    "start": person.get("leave_start") or "",
                    "end": person.get("leave_end") or "",
                    "used_days": person.get("used_comp_days") or 0,
                }
            ]

        n = max(len(travel), len(ot_days), len(leaves), 1)
        start_row = cursor
        end_row = cursor + n - 1

        for i in range(n):
            r = cursor + i
            t = travel[i] if i < len(travel) else {}
            o = ot_days[i] if i < len(ot_days) else {}
            lv = leaves[i] if i < len(leaves) else {}
            ot_date = (o.get("date") or "").replace("(法定)", "")
            values = [
                person.get("person_name") or "" if i == 0 else None,
                person.get("center") or "" if i == 0 else None,
                person.get("department") or "" if i == 0 else None,
                t.get("system_start") or "",
                t.get("system_end") or "",
                t.get("actual_start") or "",
                t.get("actual_end") or "",
                t.get("credit") if t else "",
                ot_date,
                ot_date if ot_date else "",
                o.get("credit") if o else "",
                person.get("total_comp_days") if i == 0 else None,
                lv.get("start") or "",
                lv.get("end") or "",
                lv.get("used_days") if lv else "",
                person.get("used_comp_days") if i == 0 else None,
                person.get("remaining_comp_days") if i == 0 else None,
                person.get("remark") or "" if i == 0 else None,
            ]
            for col, val in enumerate(values, start=1):
                if i > 0 and col in (1, 2, 3, 12, 16, 17, 18):
                    _cell(ws, r, col, None)
                else:
                    _cell(ws, r, col, val)

        if end_row > start_row:
            for col in (1, 2, 3, 12, 16, 17, 18):
                ws.merge_cells(
                    start_row=start_row,
                    start_column=col,
                    end_row=end_row,
                    end_column=col,
                )
            # Re-apply identity/totals on the top-left of each merge.
            _cell(ws, start_row, 1, person.get("person_name") or "")
            _cell(ws, start_row, 2, person.get("center") or "")
            _cell(ws, start_row, 3, person.get("department") or "")
            _cell(ws, start_row, 12, person.get("total_comp_days") or 0)
            _cell(ws, start_row, 16, person.get("used_comp_days") or 0)
            _cell(ws, start_row, 17, person.get("remaining_comp_days") or 0)
            _cell(ws, start_row, 18, person.get("remark") or "")

        cursor = end_row + 1

    for col in range(1, 19):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
