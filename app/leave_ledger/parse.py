"""Parse WeCom leave-ledger source workbooks into normalized records."""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError


def _as_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


_DT_RE = re.compile(
    r"(?P<y>\d{4})[/-](?P<m>\d{1,2})[/-](?P<d>\d{1,2})"
    r"(?:[ T](?P<H>\d{1,2}):(?P<M>\d{2})(?::(?P<S>\d{2}))?)?"
)
_RANGE_RE = re.compile(
    r"(?P<a>\d{4}[/-]\d{1,2}[/-]\d{1,2})"
    r"\s*[-~～至到]\s*"
    r"(?P<b>\d{4}[/-]\d{1,2}[/-]\d{1,2})"
)


def _parse_dt(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial date (openpyxl occasionally yields float)
        try:
            from openpyxl.utils.datetime import from_excel

            dt = from_excel(value)
            if isinstance(dt, datetime):
                return dt
            if isinstance(dt, date):
                return datetime(dt.year, dt.month, dt.day)
        except Exception:
            pass
    text = _as_str(value).replace("年", "/").replace("月", "/").replace("日", "")
    text = text.replace("上午", "").replace("下午", "").strip()
    for fmt in (
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    m = _DT_RE.search(text)
    if not m:
        return None
    try:
        return datetime(
            int(m.group("y")),
            int(m.group("m")),
            int(m.group("d")),
            int(m.group("H") or 0),
            int(m.group("M") or 0),
            int(m.group("S") or 0),
        )
    except ValueError:
        return None


def _parse_date(value: Any) -> date | None:
    dt = _parse_dt(value)
    return dt.date() if dt else None


def _parse_date_range(value: Any) -> tuple[date | None, date | None]:
    """Parse a single date or 'start 至 end' style range."""
    if value is None or value == "":
        return None, None
    start = _parse_date(value)
    text = _as_str(value)
    m = _RANGE_RE.search(text)
    if m:
        return _parse_date(m.group("a")), _parse_date(m.group("b"))
    return start, start if start else None


def _find_header_row(rows: list[tuple[Any, ...]], required: set[str]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows[:15]):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            name = _as_str(cell)
            if name:
                mapping[name] = col
        if required.issubset(mapping.keys()) or all(
            any(req in h for h in mapping) for req in required
        ):
            # Prefer exact keys; fall back to contains
            exact = {k: mapping[k] for k in required if k in mapping}
            if len(exact) == len(required):
                return idx, exact
            soft: dict[str, int] = {}
            for req in required:
                for h, col in mapping.items():
                    if req in h:
                        soft[req] = col
                        break
            if len(soft) == len(required):
                return idx, soft
    raise AppError(f"未找到含 {sorted(required)} 的表头行", status_code=400)


def _sheet_rows(data: bytes) -> list[tuple[Any, ...]]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        best: list[tuple[Any, ...]] = []
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            nonempty = sum(1 for r in rows if any(c is not None and _as_str(c) for c in r))
            if nonempty > sum(1 for r in best if any(c is not None and _as_str(c) for c in r)):
                best = rows
        return best
    finally:
        wb.close()


def _center_from_dept(dept: str) -> str:
    if not dept:
        return ""
    return dept.split("/")[0].strip()


@dataclass
class TravelRecord:
    applicant: str
    department: str
    center: str
    category: str
    start: date | None
    end: date | None
    status: str
    approval_id: str = ""


@dataclass
class OvertimeDay:
    applicant: str
    department: str
    center: str
    work_date: date
    hours: float
    status: str
    reason: str = ""
    approval_id: str = ""


@dataclass
class LeaveRecord:
    applicant: str
    department: str
    center: str
    leave_kind: str
    start: date | None
    end: date | None
    duration_text: str
    status: str
    approval_id: str = ""


@dataclass
class PunchRecord:
    person_name: str
    work_date: date
    punch_time: datetime | None
    status: str
    address: str


@dataclass
class ParsedSources:
    travel: list[TravelRecord] = field(default_factory=list)
    overtime: list[OvertimeDay] = field(default_factory=list)
    leave: list[LeaveRecord] = field(default_factory=list)
    punch: list[PunchRecord] = field(default_factory=list)


_OT_DAY_RE = re.compile(
    r"(?P<date>\d{4}[/-]\d{1,2}[/-]\d{1,2})\s*(?P<hours>\d+(?:\.\d+)?)\s*小时"
)


def parse_travel(data: bytes) -> list[TravelRecord]:
    rows = _sheet_rows(data)
    hdr_i, cols = _find_header_row(
        rows, {"申请人", "申请人部门", "出差类别", "开始时间", "结束时间", "当前审批状态"}
    )
    out: list[TravelRecord] = []
    for row in rows[hdr_i + 1 :]:
        applicant = _as_str(row[cols["申请人"]] if cols["申请人"] < len(row) else None)
        if not applicant:
            continue
        dept = _as_str(row[cols["申请人部门"]])
        out.append(
            TravelRecord(
                applicant=applicant,
                department=dept,
                center=_center_from_dept(dept),
                category=_as_str(row[cols["出差类别"]]),
                start=_parse_date(row[cols["开始时间"]]),
                end=_parse_date(row[cols["结束时间"]]),
                status=_as_str(row[cols["当前审批状态"]]),
                approval_id=_as_str(row[cols["审批编号"]]) if "审批编号" in cols else "",
            )
        )
    return out


def parse_overtime(data: bytes) -> list[OvertimeDay]:
    rows = _sheet_rows(data)
    hdr_i, cols = _find_header_row(
        rows, {"申请人", "申请人部门", "开始时间", "结束时间", "加班明细", "当前审批状态"}
    )
    # optional columns
    reason_col = None
    id_col = None
    for name, col in _header_map(rows[hdr_i]).items():
        if "加班事由" in name:
            reason_col = col
        if name == "审批编号":
            id_col = col
    out: list[OvertimeDay] = []
    for row in rows[hdr_i + 1 :]:
        applicant = _as_str(row[cols["申请人"]] if cols["申请人"] < len(row) else None)
        if not applicant:
            continue
        dept = _as_str(row[cols["申请人部门"]])
        status = _as_str(row[cols["当前审批状态"]])
        detail = _as_str(row[cols["加班明细"]])
        reason = _as_str(row[reason_col]) if reason_col is not None else ""
        approval_id = _as_str(row[id_col]) if id_col is not None else ""
        matches = list(_OT_DAY_RE.finditer(detail.replace("\r", "\n")))
        if matches:
            for m in matches:
                d = _parse_date(m.group("date").replace("-", "/"))
                if not d:
                    continue
                out.append(
                    OvertimeDay(
                        applicant=applicant,
                        department=dept,
                        center=_center_from_dept(dept),
                        work_date=d,
                        hours=float(m.group("hours")),
                        status=status,
                        reason=reason,
                        approval_id=approval_id,
                    )
                )
        else:
            start = _parse_date(row[cols["开始时间"]])
            if start:
                hours_text = ""
                # try 加班时长
                for name, col in _header_map(rows[hdr_i]).items():
                    if "加班时长" in name:
                        hours_text = _as_str(row[col])
                        break
                hm = re.search(r"(\d+(?:\.\d+)?)\s*小时", hours_text)
                hours = float(hm.group(1)) if hm else 0.0
                out.append(
                    OvertimeDay(
                        applicant=applicant,
                        department=dept,
                        center=_center_from_dept(dept),
                        work_date=start,
                        hours=hours,
                        status=status,
                        reason=reason,
                        approval_id=approval_id,
                    )
                )
    return out


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    return {_as_str(c): i for i, c in enumerate(row) if _as_str(c)}


def parse_leave(data: bytes) -> list[LeaveRecord]:
    rows = _sheet_rows(data)
    # Prefer classic columns; fall back when start/end are merged into one field.
    try:
        hdr_i, cols = _find_header_row(
            rows, {"申请人", "申请人部门", "假种", "开始时间", "结束时间", "当前审批状态"}
        )
        range_col = None
    except AppError:
        hdr_i, cols = _find_header_row(
            rows, {"申请人", "申请人部门", "假种", "当前审批状态"}
        )
        range_col = None
        for name, col in _header_map(rows[hdr_i]).items():
            if any(k in name for k in ("请假时间", "休假时间", "起止")):
                range_col = col
                break
        if range_col is None:
            raise AppError(
                "未找到请假开始/结束时间或请假时间列",
                status_code=400,
            )
    duration_col = None
    id_col = None
    start_col = cols.get("开始时间")
    end_col = cols.get("结束时间")
    for name, col in _header_map(rows[hdr_i]).items():
        if "请假时长" in name or name == "时长":
            duration_col = col
        if name == "审批编号":
            id_col = col
        if start_col is None and ("开始" in name and "时间" in name):
            start_col = col
        if end_col is None and ("结束" in name and "时间" in name):
            end_col = col
    out: list[LeaveRecord] = []
    for row in rows[hdr_i + 1 :]:
        applicant = _as_str(row[cols["申请人"]] if cols["申请人"] < len(row) else None)
        if not applicant:
            continue
        kind = _as_str(row[cols["假种"]])
        dept = _as_str(row[cols["申请人部门"]])
        start: date | None = None
        end: date | None = None
        if start_col is not None and end_col is not None:
            start = _parse_date(row[start_col] if start_col < len(row) else None)
            end = _parse_date(row[end_col] if end_col < len(row) else None)
        elif range_col is not None:
            start, end = _parse_date_range(row[range_col] if range_col < len(row) else None)
        out.append(
            LeaveRecord(
                applicant=applicant,
                department=dept,
                center=_center_from_dept(dept),
                leave_kind=kind,
                start=start,
                end=end,
                duration_text=_as_str(row[duration_col]) if duration_col is not None else "",
                status=_as_str(row[cols["当前审批状态"]]),
                approval_id=_as_str(row[id_col]) if id_col is not None else "",
            )
        )
    return out


def parse_punch(data: bytes) -> list[PunchRecord]:
    """Parse WeCom punch daily export.

    Common shape (Sheet3 «打卡详情»): no column headers — each row is
    [日期(+星期), 姓名, 时间, 状态, 地址, 设备, ...].
    Also supports a headered export with 日期/姓名/时间/地点.
    """
    rows = _sheet_rows(data)

    # Try headered layout first
    required_sets = [
        {"日期", "姓名", "时间", "地点"},
        {"日期", "姓名", "打卡时间", "打卡地址"},
        {"日期", "姓名", "时间", "打卡地址"},
    ]
    for req in required_sets:
        try:
            hdr_i, cols = _find_header_row(rows, req)

            def col(*names: str) -> int:
                for n in names:
                    if n in cols:
                        return cols[n]
                for n in names:
                    for h, c in cols.items():
                        if n in h:
                            return c
                raise AppError(f"缺列 {names}", status_code=400)

            date_c = col("日期")
            name_c = col("姓名")
            time_c = col("时间", "打卡时间")
            addr_c = col("地点", "打卡地址")
            status_c = None
            for h, c in _header_map(rows[hdr_i]).items():
                if "状态" in h:
                    status_c = c
                    break
            out: list[PunchRecord] = []
            for row in rows[hdr_i + 1 :]:
                person = _as_str(row[name_c] if name_c < len(row) else None)
                d = _parse_date(row[date_c] if date_c < len(row) else None)
                if not person or not d:
                    continue
                punch_t = _parse_dt(row[time_c] if time_c < len(row) else None)
                if punch_t and punch_t.year < 2000:
                    punch_t = datetime.combine(d, punch_t.time())
                out.append(
                    PunchRecord(
                        person_name=person,
                        work_date=d,
                        punch_time=punch_t,
                        status=_as_str(row[status_c]) if status_c is not None else "",
                        address=_as_str(row[addr_c] if addr_c < len(row) else None),
                    )
                )
            if out:
                return out
        except AppError:
            continue

    # Fixed-column WeCom «打卡详情» layout
    out = []
    for row in rows:
        if not row:
            continue
        first = _as_str(row[0] if len(row) > 0 else None)
        if not first or first in ("打卡详情",):
            continue
        # date cell like "2026/06/30 星期二"
        date_token = first.split()[0] if first else ""
        d = _parse_date(date_token)
        if not d:
            continue
        person = _as_str(row[1] if len(row) > 1 else None)
        if not person:
            continue
        time_text = _as_str(row[2] if len(row) > 2 else None)
        status = _as_str(row[3] if len(row) > 3 else None)
        address = _as_str(row[4] if len(row) > 4 else None)
        if address == "--":
            address = ""
        punch_t = None
        if time_text and time_text != "--":
            # time-only HH:MM
            try:
                hm = datetime.strptime(time_text, "%H:%M")
                punch_t = datetime.combine(d, hm.time())
            except ValueError:
                punch_t = _parse_dt(f"{d.isoformat()} {time_text}")
        out.append(
            PunchRecord(
                person_name=person,
                work_date=d,
                punch_time=punch_t,
                status=status,
                address=address,
            )
        )
    if not out:
        raise AppError("未找到打卡数据行", status_code=400)
    return out


def parsed_to_dict(parsed: ParsedSources) -> dict[str, Any]:
    def _travel(r: TravelRecord) -> dict[str, Any]:
        d = asdict(r)
        d["start"] = r.start.isoformat() if r.start else None
        d["end"] = r.end.isoformat() if r.end else None
        return d

    def _leave(r: LeaveRecord) -> dict[str, Any]:
        d = asdict(r)
        d["start"] = r.start.isoformat() if r.start else None
        d["end"] = r.end.isoformat() if r.end else None
        return d

    return {
        "travel": [_travel(r) for r in parsed.travel],
        "overtime": [
            {**asdict(r), "work_date": r.work_date.isoformat()} for r in parsed.overtime
        ],
        "leave": [_leave(r) for r in parsed.leave],
        "punch": [
            {
                **asdict(r),
                "work_date": r.work_date.isoformat(),
                "punch_time": r.punch_time.isoformat() if r.punch_time else None,
            }
            for r in parsed.punch
        ],
    }
