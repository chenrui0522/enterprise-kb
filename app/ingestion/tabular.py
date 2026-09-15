"""Semantic conversion for tabular documents (XLSX / XLSM).

A spreadsheet is not a document with a big table in it: it is a set of
records. Dumping it as one Markdown table produces huge, low-precision chunks,
so instead every sheet becomes a section and every row becomes a semantic
"field=value" line that is retrievable on its own. The original grid is kept
as metadata (`TableArtifact`) so later chunking strategies can rebuild
row-level chunks without parsing the file again.
"""

from __future__ import annotations

import datetime as dt

from app.core.errors import AppError
from app.core.logging import get_logger
from app.ingestion.converters import ConversionResult, FileConverter
from app.ingestion.structure import (
    BLOCK_HEADING,
    BLOCK_PARAGRAPH,
    Block,
    DocumentStructure,
    TableArtifact,
    structure_to_markdown,
)

logger = get_logger("ingestion.tabular")

MAX_ROWS = 20000
MAX_COLUMNS = 80
DEFAULT_ROWS_PER_GROUP = 25
MAX_HEADER_CHARS = 40
GENERIC_HEADER_PREFIX = "列"


class XlsxSemanticConverter(FileConverter):
    """XLSX -> per-sheet sections with per-row semantic lines."""

    name = "xlsx-semantic"
    version = "1"

    def __init__(self, rows_per_group: int = DEFAULT_ROWS_PER_GROUP) -> None:
        self._rows_per_group = max(int(rows_per_group), 1)

    def cache_signature(self) -> dict:
        return {**super().cache_signature(), "rows_per_group": self._rows_per_group}

    # ------------------------------------------------------------------
    # FileConverter interface
    # ------------------------------------------------------------------
    def convert(self, file_path: str) -> ConversionResult:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - openpyxl ships with markitdown[xlsx]
            raise AppError("服务端缺少 openpyxl 依赖，无法转换 Excel 文件") from exc

        try:
            workbook = load_workbook(file_path, data_only=True)
        except Exception as exc:
            raise AppError(f"无法打开 Excel 文件（仅支持 .xlsx/.xlsm）：{exc}") from exc

        blocks: list[Block] = []
        tables: list[TableArtifact] = []
        try:
            formulas = _formula_lookup(file_path)
            for worksheet in workbook.worksheets:
                if getattr(worksheet, "sheet_state", "visible") != "visible":
                    continue
                grid = _read_grid(worksheet, formulas=formulas.get(worksheet.title))
                if not grid.rows:
                    continue
                sheet_name = _clean_text(worksheet.title) or "Sheet"
                header, data_rows, header_row = _infer_header(grid.rows)
                header, data_rows = _prune_columns(header, data_rows)
                if not data_rows:
                    continue
                blocks.append(
                    Block(
                        type=BLOCK_HEADING,
                        text=f"{sheet_name}（{len(data_rows)} 行数据）",
                        level=2,
                    )
                )
                blocks.extend(_row_blocks(sheet_name, header, data_rows, header_row, self._rows_per_group))
                tables.append(
                    TableArtifact(
                        sheet=sheet_name,
                        header=list(header),
                        rows=[list(row) for row in data_rows],
                        source="xlsx",
                        start_row=header_row + 1,
                    )
                )
        finally:
            try:
                workbook.close()
            except Exception:  # pragma: no cover - defensive
                logger.debug("Could not close workbook", exc_info=True)

        if not blocks:
            raise AppError("未能从 Excel 中提取任何内容（所有工作表为空或被隐藏）")

        structure = DocumentStructure(blocks=blocks)
        markdown = structure_to_markdown(structure)
        return ConversionResult(
            markdown=markdown,
            page_count=None,
            structure=structure,
            tables=tables,
            triage={
                "converter_service": "xlsx-semantic",
                "sheets": len(tables),
                "rows": sum(len(table.rows) for table in tables),
            },
        )


class _Grid:
    def __init__(self, rows: list[list[str]]) -> None:
        self.rows = rows


def _read_grid(worksheet, *, formulas: dict[str, str] | None = None) -> _Grid:
    """Read one sheet into a dense string grid, filling merged cells down."""
    formulas = formulas or {}
    max_row = min(int(worksheet.max_row or 0), MAX_ROWS)
    max_column = min(int(worksheet.max_column or 0), MAX_COLUMNS)
    if max_row <= 0 or max_column <= 0:
        return _Grid([])

    grid = [["" for _ in range(max_column)] for _ in range(max_row)]
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column), start=0
    ):
        for column_index, cell in enumerate(row):
            grid[row_index][column_index] = _format_value(
                cell.value, formulas.get(cell.coordinate)
            )

    for merged in getattr(worksheet, "merged_cells", []) or []:
        try:
            bounds = merged.bounds  # (min_col, min_row, max_col, max_row), 1-based
        except Exception:
            continue
        min_col, min_row, max_col, max_row = bounds
        if min_row < 1 or min_col < 1 or min_row > max_row or min_col > max_col:
            continue
        top_left = grid[min_row - 1][min_col - 1] if min_row - 1 < len(grid) else ""
        if not top_left:
            continue
        for row_index in range(min_row - 1, min(max_row, len(grid))):
            for column_index in range(min_col - 1, min(max_col, len(grid[row_index]))):
                if not grid[row_index][column_index]:
                    grid[row_index][column_index] = top_left

    trimmed = [row for row in grid if any(cell.strip() for cell in row)]
    return _Grid(trimmed)


def _formula_lookup(file_path: str) -> dict[str, dict[str, str]]:
    """Formula text per sheet, keyed by cell coordinate.

    `data_only=True` yields the last value Excel computed, which is what we
    want to index. When a workbook was written by a tool that never computed
    formulas the value is empty, and the formula string is the only content
    left, so it is used as a fallback. The formula workbook is loaded once for
    the whole file.
    """
    lookup: dict[str, dict[str, str]] = {}
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, data_only=False, read_only=True)
    except Exception:
        return lookup
    try:
        for worksheet in workbook.worksheets:
            cells: dict[str, str] = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        cells[cell.coordinate] = value
            if cells:
                lookup[worksheet.title] = cells
    except Exception:
        return lookup
    finally:
        try:
            workbook.close()
        except Exception:  # pragma: no cover - defensive
            logger.debug("Could not close formula workbook", exc_info=True)
    return lookup


def _format_value(value, formula: str | None = None) -> str:
    if value is None:
        return _clean_text(formula or "")
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dt.datetime, dt.date)):
        if isinstance(value, dt.datetime):
            return value.isoformat(sep=" ", timespec="seconds")
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    return _clean_text(str(value))


def _clean_text(value: str) -> str:
    return " ".join(str(value).replace("\n", " ").split())


def _infer_header(rows: list[list[str]]) -> tuple[list[str], list[list[str]], int]:
    """Return (header, data rows, header row index)."""
    if not rows:
        return [], [], 0
    width = max(len(row) for row in rows)
    candidate = rows[0]
    if _looks_like_header(candidate, rows[1] if len(rows) > 1 else None, width):
        header = _unique_names(candidate, width)
        return header, rows[1:], 0
    header = [f"{GENERIC_HEADER_PREFIX}{index + 1}" for index in range(width)]
    return header, rows, 0


def _looks_like_header(row: list[str], following: list[str] | None, width: int) -> bool:
    """Decide whether the first non-empty row is a field-name row.

    Signals: the row is labels only (text, short, unique) and the row below
    either carries a non-label value, has a different width, or reads like
    data (longer cells). A single-cell first row is treated as a title.
    """
    filled = [cell for cell in row if cell.strip()]
    if not filled:
        return False
    if not all(len(cell) <= MAX_HEADER_CHARS for cell in filled):
        return False
    if not all(_is_textual(cell) for cell in filled):
        return False
    if len(set(filled)) != len(filled):
        return False
    # A single-cell first row is usually a title, not a header.
    if len(filled) < 2:
        return False
    if following is None:
        return True
    following_filled = [cell for cell in following if cell.strip()]
    if not following_filled:
        return True
    if any(not _is_textual(cell) for cell in following_filled):
        return True
    if len(following_filled) != len(filled):
        return True
    # All-text table (name/department/status style): labels are usually the
    # shorter, digit-free side of the pair.
    if any(any(char.isdigit() for char in cell) for cell in filled):
        return False
    header_length = sum(len(cell) for cell in filled) / len(filled)
    following_length = sum(len(cell) for cell in following_filled) / len(following_filled)
    return header_length <= following_length


def _is_textual(cell: str) -> bool:
    value = cell.strip()
    if not value:
        return False
    try:
        float(value)
    except ValueError:
        return True
    return False


def _unique_names(row: list[str], width: int) -> list[str]:
    header: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        raw = _clean_text(row[index]) if index < len(row) else ""
        name = raw or f"{GENERIC_HEADER_PREFIX}{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        header.append(name)
    return header


def _prune_columns(header: list[str], rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    keep = [
        index
        for index in range(len(header))
        if _clean_text(header[index]) or any(index < len(row) and row[index].strip() for row in rows)
    ]
    if not keep:
        return header, rows
    pruned_header = [header[index] for index in keep]
    pruned_rows = [
        [(row[index] if index < len(row) else "") for index in keep] for row in rows
    ]
    return pruned_header, pruned_rows


def _row_blocks(
    sheet_name: str,
    header: list[str],
    rows: list[list[str]],
    header_row: int,
    rows_per_group: int,
) -> list[Block]:
    blocks: list[Block] = []
    total = len(rows)
    for start in range(0, total, rows_per_group):
        group = rows[start : start + rows_per_group]
        lines = []
        if total > rows_per_group:
            lines.append(
                f"工作表：{sheet_name}（第 {header_row + start + 2}-{header_row + start + 1 + len(group)} 行）"
            )
        lines.append("字段：" + "、".join(header))
        for offset, row in enumerate(group, start=start):
            pairs = [
                f"{header[index]}={_clean_text(row[index]) if index < len(row) else ''}"
                for index in range(len(header))
                if index < len(row) and row[index].strip()
            ]
            if not pairs:
                continue
            lines.append(f"第 {header_row + offset + 2} 行：" + "；".join(pairs))
        if any(line.startswith("第 ") for line in lines):
            blocks.append(
                Block(
                    type=BLOCK_PARAGRAPH,
                    text="\n".join(lines),
                )
            )
    return blocks


