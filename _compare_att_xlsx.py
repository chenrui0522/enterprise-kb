"""Compare system Taiyuan attendance xlsx vs 正确表; write markdown report."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

SYS = Path(
    r"c:\Users\chenr\Downloads\考勤汇总-2026-8-d27cdf6fbe8640ef9febcf8fb084117e (1).xlsx"
)
CORRECT = Path(
    r"c:\Users\chenr\Desktop\奥特莱多智能体项目资料汇总‘\AI考勤统计\26年8月太原考勤汇总表jm.xlsx"
)
OUT = Path("考勤全面对比-系统vs正确表-d27cdf6f.md")

# Red-box columns as labeled on row 3 after day columns
TALLY_LABELS = [
    "应出勤",
    "实出勤",
    "国内出差",
    "国外出差",
    "居家",
    "事假",
    "年假",
    "病假",
    "产假",
    "婚假",
    "丧假",
    "调休",
    "迟到",
    "缺卡",
    "旷工",
]

LABEL_ALIAS = {
    "产/产检/陪产假": "产假",
    "迟到（次）": "迟到",
    "缺卡（次）": "缺卡",
    "SHI假": "事假",
}


def _norm_sym(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    # normalize common lookalikes / fullwidth yen
    if s in {"√", "✓", "☑"}:
        return "√"
    if s in {"￥", "¥", "￥"}:
        return "￥"
    if s in {"＄", "$"}:
        return "＄"
    if s in {"×", "x", "X", "✕"}:
        return "×"
    if s in {"★", "*"}:
        return "★"
    if s == "■":
        return "■"
    return s


def _num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return 0
    try:
        return float(s)
    except ValueError:
        return s


def _find_sheet(wb):
    for name in wb.sheetnames:
        if "考勤" in name or name in ("Sheet1", "考勤表"):
            return wb[name]
    return wb[wb.sheetnames[0]]


def parse_workbook(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = _find_sheet(wb)
    header_row = None
    day_cols: dict[int, int] = {}
    tally_cols: dict[str, int] = {}

    for r in range(1, min(10, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(60, ws.max_column + 1))]
        if "序号" in vals and any(isinstance(v, datetime) for v in vals):
            header_row = r
            break
    if not header_row:
        raise RuntimeError(f"cannot find header in {path}")

    for c in range(1, ws.max_column + 1):
        val = ws.cell(header_row, c).value
        if isinstance(val, datetime):
            day_cols[val.day] = c
        elif isinstance(val, str) and re.fullmatch(r"\d{1,2}", val.strip()):
            day_cols[int(val.strip())] = c

    for rr in range(header_row, min(header_row + 3, ws.max_row + 1)):
        for c in range(1, ws.max_column + 1):
            lab = ws.cell(rr, c).value
            if not isinstance(lab, str):
                continue
            alias = LABEL_ALIAS.get(lab.strip(), lab.strip())
            if alias in TALLY_LABELS and alias not in tally_cols:
                tally_cols[alias] = c

    # Locate name/side by scanning first person block
    name_col = side_col = None
    for r in range(header_row + 1, min(header_row + 8, ws.max_row + 1)):
        for c in range(1, 8):
            if ws.cell(r, c).value == "上班":
                side_col = c
                # name is nearest non-empty text left of side
                for cc in range(c - 1, 0, -1):
                    v = ws.cell(r, cc).value
                    if isinstance(v, str) and v.strip() and v.strip() not in (
                        "中心",
                        "岗位",
                        "姓名",
                    ):
                        name_col = cc
                        break
                break
        if name_col and side_col:
            break
    if not name_col or not side_col:
        name_col, side_col = 4, 5

    people: dict[str, dict] = {}
    r = header_row + 1
    while r <= ws.max_row:
        side = ws.cell(r, side_col).value
        name = ws.cell(r, name_col).value
        if side == "上班" and name and str(name).strip():
            pname = str(name).strip()
            pname_key = re.sub(r"[（(].*?离职.*?[）)]", "", pname).strip()
            pname_key = re.sub(r"\s+", "", pname_key)
            in_row = r
            out_row = r + 1
            for rr in (r + 1, r + 2):
                if ws.cell(rr, side_col).value in ("下班", "下"):
                    out_row = rr
                    break
            days = {
                day: (
                    _norm_sym(ws.cell(in_row, col).value),
                    _norm_sym(ws.cell(out_row, col).value),
                )
                for day, col in day_cols.items()
            }
            tallies = {lab: _num(ws.cell(in_row, col).value) for lab, col in tally_cols.items()}
            people[pname_key] = {
                "name": pname_key,
                "raw_name": pname,
                "seq": ws.cell(in_row, 1).value,
                "center": ws.cell(in_row, 2).value,
                "post": ws.cell(in_row, 3).value,
                "days": days,
                "tallies": tallies,
            }
            r = out_row + 1
            continue
        r += 1
    return people


def compare(sys_p: dict, cor_p: dict):
    all_names = sorted(set(sys_p) | set(cor_p), key=lambda n: (
        sys_p.get(n, cor_p.get(n, {})).get("seq") or 9999,
        n,
    ))
    only_sys = sorted(set(sys_p) - set(cor_p))
    only_cor = sorted(set(cor_p) - set(sys_p))
    exact = []
    tally_diff = []
    day_only = []
    col_freq = Counter()

    for name in all_names:
        if name not in sys_p or name not in cor_p:
            continue
        s, c = sys_p[name], cor_p[name]
        t_diffs = []
        for lab in TALLY_LABELS:
            sv, cv = s["tallies"].get(lab, 0), c["tallies"].get(lab, 0)
            if sv != cv:
                # float tolerance
                if isinstance(sv, float) or isinstance(cv, float):
                    try:
                        if abs(float(sv) - float(cv)) < 1e-9:
                            continue
                    except (TypeError, ValueError):
                        pass
                t_diffs.append((lab, sv, cv))
                col_freq[lab] += 1
        d_diffs = []
        days = sorted(set(s["days"]) | set(c["days"]))
        for d in days:
            sd = s["days"].get(d, ("", ""))
            cd = c["days"].get(d, ("", ""))
            if sd != cd:
                d_diffs.append((d, sd, cd))
        if not t_diffs and not d_diffs:
            exact.append(name)
        elif t_diffs:
            tally_diff.append((name, t_diffs, d_diffs, s, c))
        else:
            day_only.append((name, d_diffs, s, c))

    return {
        "all_names": all_names,
        "only_sys": only_sys,
        "only_cor": only_cor,
        "exact": exact,
        "tally_diff": tally_diff,
        "day_only": day_only,
        "col_freq": col_freq,
    }


def _md_day_table(d_diffs, limit=40) -> str:
    if not d_diffs:
        return ""
    lines = [
        "",
        "<details><summary>日格差（点击展开）</summary>",
        "",
        "| 日期 | 系统上/下 | 正确上/下 |",
        "|------|-----------|-----------|",
    ]
    for d, sd, cd in d_diffs[:limit]:
        lines.append(f"| 8/{d} | {sd[0] or '空'}/{sd[1] or '空'} | {cd[0] or '空'}/{cd[1] or '空'} |")
    if len(d_diffs) > limit:
        lines.append(f"| … | 另有 {len(d_diffs)-limit} 日 | … |")
    lines += ["", "</details>", ""]
    return "\n".join(lines)


def guess_reasons(name: str, t_diffs, d_diffs) -> list[str]:
    reasons = []
    labs = {x[0] for x in t_diffs}
    if "国内出差" in labs or "国外出差" in labs:
        reasons.append(
            "**出差天数/符号不一致**：常见于驻地裁剪、晚出发≥18:00 计差、出差与打卡混日，或正确表手工改过。"
        )
    if "实出勤" in labs:
        reasons.append("**实出勤不一致**：常随出差/请假/非出勤日打卡决议/缺卡旷工联动变化。")
    if "调休" in labs or "事假" in labs or "年假" in labs or "病假" in labs:
        reasons.append("**假勤分列不一致**：半天时段、假种映射或审批状态过滤差异。")
    if "迟到" in labs or "缺卡" in labs or "旷工" in labs:
        reasons.append("**打卡异常计数不一致**：补卡是否计缺卡、非应出勤日是否计次、离职截止日。")
    if "应出勤" in labs:
        reasons.append("**应出勤天数不一致**：月历（含 8/8、8/22 是否出勤）或入离职裁剪不同。")
    # day symbol patterns
    travel_day_mismatch = [
        d
        for d, sd, cd in d_diffs
        if ("￥" in cd or "＄" in cd) and ("￥" not in sd and "＄" not in sd)
    ]
    if travel_day_mismatch:
        reasons.append(
            f"**日格出差符号**：正确表为出差符号而系统为出勤/空的日期含 8/{',8/'.join(map(str, travel_day_mismatch[:8]))}"
            + ("…" if len(travel_day_mismatch) > 8 else "")
            + "。"
        )
    if not reasons:
        reasons.append("日格或汇总细项差异，需对照打卡/审批源逐日核实。")
    return reasons


def render(result, sys_path: Path, cor_path: Path) -> str:
    lines = []
    lines.append("# 考勤全面对比（系统导出 vs 正确表）")
    lines.append("")
    lines.append("## 对比条件")
    lines.append("")
    lines.append(f"- 系统导出：`{sys_path.name}`")
    lines.append(f"- 正确表：`{cor_path.name}`")
    lines.append("- 目标月：2026-08")
    lines.append("- 箭头：`系统值 → 正确表值`")
    lines.append("")
    lines.append("## 总览")
    lines.append("")
    lines.append("| 项 | 数量 |")
    lines.append("|----|------|")
    n_sys = len(result["all_names"]) - len(result["only_cor"])
    # recount from parse sizes passed via result meta
    lines.append(
        f"| 系统/正确表人数 | {result['n_sys']} / {result['n_cor']} |"
    )
    lines.append(f"| 完全一致（汇总+日格） | {len(result['exact'])} |")
    lines.append(f"| 红框汇总有差 | {len(result['tally_diff'])} |")
    lines.append(f"| 仅日格差 | {len(result['day_only'])} |")
    lines.append(f"| 仅系统有 | {len(result['only_sys'])} |")
    lines.append(f"| 仅正确表有 | {len(result['only_cor'])} |")
    lines.append("")

    # key people quick check
    key = ["周泽江", "李新兵", "杨林瑞", "何继成", "赵中浩", "冯晓虎", "曹海源", "温惠琴", "张锐", "张楠杰", "刘三平"]
    lines.append("### 关键人快检")
    lines.append("")
    for name in key:
        hit = next((x for x in result["tally_diff"] if x[0] == name), None)
        if name in result["exact"]:
            lines.append(f"- {name}: 完全一致")
        elif hit:
            parts = [f"{lab}{sv}→{cv}" for lab, sv, cv in hit[1][:6]]
            lines.append(f"- {name}: " + " ".join(parts))
        elif name in result["only_sys"]:
            lines.append(f"- {name}: 仅系统有")
        elif name in result["only_cor"]:
            lines.append(f"- {name}: 仅正确表有")
        else:
            day_hit = next((x for x in result["day_only"] if x[0] == name), None)
            if day_hit:
                lines.append(f"- {name}: 汇总一致，仅日格 {len(day_hit[1])} 处差")
            else:
                lines.append(f"- {name}: （未在双方名单）")
    lines.append("")

    if result["col_freq"]:
        lines.append("### 红框汇总列出现频次")
        lines.append("")
        lines.append("| 列 | 人数 |")
        lines.append("|----|------|")
        for lab, n in result["col_freq"].most_common():
            lines.append(f"| {lab} | {n} |")
        lines.append("")

    lines.append("## 完全一致人员")
    lines.append("")
    if result["exact"]:
        lines.append("、".join(result["exact"]))
    else:
        lines.append("（无）")
    lines.append("")

    if result["only_sys"] or result["only_cor"]:
        lines.append("## 名单差异")
        lines.append("")
        if result["only_sys"]:
            lines.append("- 仅系统：" + "、".join(result["only_sys"]))
        if result["only_cor"]:
            lines.append("- 仅正确表：" + "、".join(result["only_cor"]))
        lines.append("")

    lines.append("## 红框汇总有差（含原因）")
    lines.append("")
    for name, t_diffs, d_diffs, s, c in result["tally_diff"]:
        seq = s.get("seq") or ""
        lines.append(f"### {seq} {name}".strip())
        lines.append("")
        lines.append("**红框汇总**")
        lines.append("")
        lines.append("| 列 | 系统 | 正确表 |")
        lines.append("|----|------|--------|")
        for lab, sv, cv in t_diffs:
            lines.append(f"| {lab} | {sv} | {cv} |")
        lines.append("")
        lines.append("**主要原因**")
        lines.append("")
        for i, reason in enumerate(guess_reasons(name, t_diffs, d_diffs), 1):
            lines.append(f"{i}. {reason}")
        lines.append(_md_day_table(d_diffs))
        lines.append("")

    if result["day_only"]:
        lines.append("## 仅日格差（红框一致）")
        lines.append("")
        for name, d_diffs, s, c in result["day_only"]:
            seq = s.get("seq") or ""
            lines.append(f"### {seq} {name}".strip())
            lines.append(_md_day_table(d_diffs))
            lines.append("")

    lines.append("## 附录")
    lines.append("")
    lines.append(f"- 生成脚本：`_compare_att_xlsx.py`")
    lines.append(f"- 系统路径：`{sys_path}`")
    lines.append(f"- 正确表路径：`{cor_path}`")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    import argparse
    import json
    import sys as _sys

    from app.attendance_summary.compare_ledger import compare_files, summarize

    parser = argparse.ArgumentParser(description="Compare system vs 正确表 attendance xlsx")
    parser.add_argument("--system", type=Path, default=SYS)
    parser.add_argument("--correct", type=Path, default=CORRECT)
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument(
        "--exemptions",
        type=Path,
        default=None,
        help="JSON exemption list (default: app/attendance_summary/correct_ledger_exemptions.json)",
    )
    parser.add_argument(
        "--fail-on-non-exempt",
        action="store_true",
        help="Exit 1 when non-exempt diffs remain (CI gate)",
    )
    parser.add_argument(
        "--json-summary",
        type=Path,
        default=None,
        help="Optional path to write JSON summary (includes non_exempt_diffs)",
    )
    args = parser.parse_args()

    if not args.system.is_file():
        raise SystemExit(f"missing system: {args.system}")
    if not args.correct.is_file():
        raise SystemExit(f"missing correct: {args.correct}")

    # Legacy markdown report (full narrative)
    sys_p = parse_workbook(args.system)
    cor_p = parse_workbook(args.correct)
    print("parsed", len(sys_p), len(cor_p))
    result = compare(sys_p, cor_p)
    result["n_sys"] = len(sys_p)
    result["n_cor"] = len(cor_p)
    md = render(result, args.system, args.correct)
    args.out.write_text(md, encoding="utf-8")
    print("wrote", args.out.resolve())
    print(
        "exact",
        len(result["exact"]),
        "tally_diff",
        len(result["tally_diff"]),
        "day_only",
        len(result["day_only"]),
    )

    # Exemption-aware gate summary
    cmp = compare_files(
        args.system, args.correct, exemptions_path=args.exemptions
    )
    summary = summarize(cmp)
    print(
        "non_exempt_diffs",
        summary["non_exempt_diffs"],
        "exempt_diffs",
        summary["exempt_diffs"],
        "exact",
        summary["exact"],
    )
    if args.json_summary:
        args.json_summary.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print("summary", args.json_summary.resolve())
    if args.fail_on_non_exempt and summary["non_exempt_diffs"] > 0:
        print("FAIL: non-exempt diffs remain (see sample):")
        for row in summary["non_exempt_sample"][:10]:
            print(" ", row)
        _sys.exit(1)


if __name__ == "__main__":
    main()
