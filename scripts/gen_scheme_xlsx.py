"""One-sheet comparison of cloud, hybrid, and phased rollout."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "intranet-multi-agent-方案对比.xlsx"

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
ADOPT = PatternFill("solid", fgColor="E2EFDA")
HOLD = PatternFill("solid", fgColor="FFF2CC")
NO = PatternFill("solid", fgColor="FCE4D6")


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "方案对比"
    ws["A1"] = "内网多Agent平台 三种方案对比"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:H1")
    ws["A2"] = (
        "已确认：检索片段可发给 DeepSeek。"
        "V1试跑采用档A：云端生成 + 本地CPU检索（零GPU采购）。"
        "真·纯云端(检索也上云)不采用。全内网L20为后续预案。"
        "硬件价为2026-09国内含税参考价。API按Flash高峰公开价粗算，不是账单。"
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 36

    headers = ["方案", "结论", "数据边界", "生成", "检索", "第一年采购", "参考金额", "说明"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    rows = [
        (
            "档A：云端生成+本地CPU检索",
            "V1试跑采用",
            "原文、向量、BM25留内网；问句+命中片段可出网",
            "DeepSeek deepseek-v4-flash",
            "bge-m3 + bge-reranker-v2-m3，Xinference device=cpu",
            "零GPU；沿用现有开发机 / Compose",
            "硬件 0；API按量",
            "与现enterprise-kb默认一致。不买卡即可验证问答、引用、多轮、入库。",
            ADOPT,
        ),
        (
            "档B：真·纯云端",
            "不采用",
            "问句、命中片段、待向量化文本都出网",
            "DeepSeek等公有云API",
            "云端embedding / rerank",
            "可不买GPU，但要接云端检索API",
            "硬件近0；API另计",
            "上线也快，但向量化文本也出网；第一版不走这条。",
            NO,
        ),
        (
            "混合(上GPU检索)",
            "全员入口后按需",
            "同档A边界",
            "仍用DeepSeek Flash",
            "同上，迁到L4 GPU",
            "业务机约6.5万；L4约2.1万按需",
            "约6.5万起；API一年大约几千到一两万",
            "CPU重排跟不上再加L4。不是V1采购项。",
            HOLD,
        ),
        (
            "全内网",
            "预案，触发后再买",
            "生成也不出网",
            "Qwen2.5-32B-Instruct（或Qwen3-32B），备选GLM-4-32B-0414",
            "同上，放在L4",
            "业务机+推理主机+L4+L20",
            "一期硬件约16.9万；含交换机/UPS约18.1万",
            "买的是数据不出网。出网政策收紧或API不稳/不便宜时再启用。",
            HOLD,
        ),
    ]
    for i, row in enumerate(rows, 5):
        *values, fill = row
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = THIN
            cell.alignment = WRAP
            cell.fill = fill
        ws.row_dimensions[i].height = 58

    ws["A10"] = "迭代阶段（按这个买）"
    ws["A10"].font = TITLE_FONT
    phase_headers = ["阶段", "运行方案", "采购", "参考金额(元)", "API", "进入下一步的条件"]
    for c, h in enumerate(phase_headers, 1):
        cell = ws.cell(row=11, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    phases = [
        ("1. V1试跑(当前)", "档A", "零GPU；Compose起Postgres/Redis/Milvus；model-service用CPU；DeepSeek key", 0, "按量，Flash", "功能验收通过，准备扩大范围"),
        ("2. 全员入口", "档A→按需加L4", "业务机32核/128GB/NVMe；L4仅当CPU重排跟不上", 65000, "一年大约几千到一两万", "出网政策收紧，或API不稳、年费明显高于本地"),
        ("2b. 重排跟不上时", "混合+GPU检索", "加L4 24GB，TEI跑bge-m3+bge-reranker-v2-m3", 21000, "不变", "并入上一阶段，不是V1必买"),
        ("3. 转为全内网", "全内网预案", "再加推理主机+L20 48GB，换Qwen2.5-32B-Instruct INT4", 83000, "生成改为0", "阶段2触发条件成立。若尚无L4再加2.1万"),
    ]
    for i, row in enumerate(phases, 12):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = THIN
            cell.alignment = WRAP
            cell.font = BOLD if i == 12 else Font()
            if c == 4 and isinstance(v, int):
                cell.number_format = "#,##0"
        ws.row_dimensions[i].height = 36
        if i == 12:
            for c in range(1, 7):
                ws.cell(row=i, column=c).fill = ADOPT

    ws["A17"] = (
        "API 粗算口径：2026-09-10 后 Flash 高峰价，未命中输入 2 元、输出 8 元 / 百万 tokens。"
        "假设每人每天 8 问、每问 3 次调用、每次约 6 千输入 + 1 千输出，且都发生在工作高峰、缓存未命中。"
        "110 人全员打满大约一年几千到一两万元；上下文拉得很长或改用 Pro 会上去。价格会变，以 DeepSeek 当时定价页为准。"
        "金额不含实施人力、机柜租赁和品牌整机 15%-30% 的维保上浮。"
    )
    ws["A17"].alignment = WRAP
    ws.merge_cells("A17:H17")
    ws.row_dimensions[17].height = 48

    widths = [18, 28, 42, 42, 36, 42, 36, 55]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name("intranet-multi-agent-方案对比-new.xlsx")
        wb.save(alt)
        print(alt)


if __name__ == "__main__":
    main()
