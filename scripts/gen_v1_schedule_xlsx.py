"""Generate V1 12-week development schedule Excel in project-plan template format."""

from __future__ import annotations

import os
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def is_workday(day: date) -> bool:
    """单双休：周日全休；ISO 单周周六上班，双周周六休息。"""
    if day.weekday() == 6:
        return False
    if day.weekday() == 5:
        return day.isocalendar().week % 2 == 1
    return True


def next_workday(day: date) -> date:
    while not is_workday(day):
        day += timedelta(days=1)
    return day


def prev_workday(day: date) -> date:
    while not is_workday(day):
        day -= timedelta(days=1)
    return day


def count_workdays(start: date, end: date) -> int:
    total = 0
    day = start
    while day <= end:
        if is_workday(day):
            total += 1
        day += timedelta(days=1)
    return total


def snap_range(start: date, end: date) -> tuple[date, date, int]:
    snapped_start = next_workday(start)
    snapped_end = prev_workday(end)
    if snapped_end < snapped_start:
        snapped_end = next_workday(end)
    return snapped_start, snapped_end, count_workdays(snapped_start, snapped_end)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "进度计划"

    thin = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    title_font = Font(name="微软雅黑", size=14, bold=True)
    meta_font = Font(name="微软雅黑", size=10)
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    cat_font = Font(name="微软雅黑", size=10, bold=True)
    cell_font = Font(name="微软雅黑", size=10)
    header_fill = PatternFill("solid", fgColor="808080")
    cat_fill = PatternFill("solid", fgColor="D9E2F3")
    done_cat_fill = PatternFill("solid", fgColor="C6EFCE")
    alt_fill = PatternFill("solid", fgColor="F2F2F2")
    done_fill = PatternFill("solid", fgColor="E2EFDA")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    start = date(2026, 9, 1)
    forward_start = date(2026, 10, 8)
    end = date(2026, 12, 31)
    total_days = count_workdays(start, end)
    done_days = count_workdays(start, date(2026, 9, 20))
    remain_days = count_workdays(forward_start, end)

    ws.merge_cells("A1:I1")
    ws["A1"] = (
        "企业知识库/项目部工作台 V1 开发进度计划"
        "（含已完成底座 + 人员投入统计首批）v3 (20260921)"
    )
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    for col in range(1, 10):
        ws.cell(1, col).border = thin

    ws.merge_cells("A2:B2")
    ws["A2"] = "工程名: 企业知识库-项目部工作台V1"
    ws.merge_cells("C2:D2")
    ws["C2"] = "项目经理: 陈锐"
    ws.merge_cells("E2:F2")
    ws["E2"] = (
        f"总开始/结束: {start.year}/{start.month}/{start.day} 至 "
        f"{end.year}/{end.month}/{end.day}"
    )
    ws["G2"] = f"总工期: {total_days} 工作日"
    ws.merge_cells("H2:I2")
    ws["H2"] = "间隔: 7 天"
    for col in range(1, 10):
        c = ws.cell(2, col)
        c.font = meta_font
        c.alignment = left
        c.border = thin

    headers = [
        "序号",
        "项目名称",
        "开始时间",
        "结束时间",
        "天数",
        "责任人",
        "计划更新原因",
        "责任判定",
        "批准人",
    ]
    for i, h in enumerate(headers, 1):
        c = ws.cell(3, i, h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = thin

    def d(y: int, m: int, day: int) -> date:
        return date(y, m, day)

    owner = "陈锐"
    # kind, name, start, end, owner, note
    # note="已完成" 表示底座已交付；空表示 10/8 后计划
    rows: list[tuple] = [
        ("cat", "一、已完成底座·平台骨架与知识问答（enterprise-kb）", None, None, None, "done"),
        (
            "task",
            "Docker Compose 基建：Postgres/Redis/Milvus/MinIO/模型服务与本地运行手册",
            d(2026, 9, 1),
            d(2026, 9, 5),
            owner,
            "已完成",
        ),
        (
            "task",
            "FastAPI 后端骨架、配置、Alembic、文档上传与异步接入 worker",
            d(2026, 9, 1),
            d(2026, 9, 8),
            owner,
            "已完成",
        ),
        (
            "task",
            "LangGraph 多轮问答：改写/检索判定/混合检索/重排/带引用生成 + SSE",
            d(2026, 9, 3),
            d(2026, 9, 12),
            owner,
            "已完成",
        ),
        (
            "task",
            "前端 React 聊天与文档管理初版；仓库发布与安全提示",
            d(2026, 9, 7),
            d(2026, 9, 12),
            owner,
            "已完成",
        ),
        ("cat", "二、已完成底座·入库检索增强", None, None, None, "done"),
        (
            "task",
            "GPU 模型服务（embedding/rerank）与性能参数调优",
            d(2026, 9, 10),
            d(2026, 9, 12),
            owner,
            "已完成",
        ),
        (
            "task",
            "结构感知切片与父子分块 structure-v2（小块检索、大块生成）",
            d(2026, 9, 11),
            d(2026, 9, 15),
            owner,
            "已完成",
        ),
        (
            "task",
            "文档转换分诊：MinerU/Docling/表格语义化、转换报告与解析缓存",
            d(2026, 9, 12),
            d(2026, 9, 15),
            owner,
            "已完成",
        ),
        (
            "task",
            "图片抽取展示、引用挂图与 lightbox",
            d(2026, 9, 13),
            d(2026, 9, 15),
            owner,
            "已完成",
        ),
        ("cat", "三、已完成底座·组织权限与品牌", None, None, None, "done"),
        (
            "task",
            "组织树/多岗/密级/项目专业切片；登录会话；语料可见性过滤",
            d(2026, 9, 16),
            d(2026, 9, 20),
            owner,
            "已完成",
        ),
        (
            "task",
            "前端奥特莱品牌（Logo、登录/侧栏/欢迎空态、品牌色）",
            d(2026, 9, 18),
            d(2026, 9, 20),
            owner,
            "已完成",
        ),
        ("cat", "四、定标与骨架（W1，10/8起）", None, None, None, None),
        (
            "task",
            "V1成功标准定稿；明确不做计划图/指标测算/文档草稿/门禁",
            d(2026, 10, 8),
            d(2026, 10, 9),
            owner,
            "",
        ),
        (
            "task",
            "人员混表+计划表Excel模板；20行样例与3组手算答案",
            d(2026, 10, 8),
            d(2026, 10, 10),
            owner,
            "",
        ),
        (
            "task",
            "personnel模块骨架、数据表与Alembic迁移",
            d(2026, 10, 11),
            d(2026, 10, 14),
            owner,
            "",
        ),
        (
            "task",
            "本机Compose冒烟；目标机/出网/GPU缺口清单",
            d(2026, 10, 13),
            d(2026, 10, 14),
            owner,
            "",
        ),
        ("cat", "五、人员投入-导入（W2）", None, None, None, None),
        (
            "task",
            "混表（出差/调令/在场）解析、必填与日期校验",
            d(2026, 10, 15),
            d(2026, 10, 18),
            owner,
            "",
        ),
        (
            "task",
            "计划表导入；错误行号回传；项目权限校验",
            d(2026, 10, 19),
            d(2026, 10, 21),
            owner,
            "",
        ),
        ("cat", "六、人员投入-去重与冲突（W3）", None, None, None, None),
        (
            "task",
            "完全相同记录合并去重规则",
            d(2026, 10, 22),
            d(2026, 10, 24),
            owner,
            "",
        ),
        (
            "task",
            "时间重叠、跨项目占用打待审标记",
            d(2026, 10, 25),
            d(2026, 10, 28),
            owner,
            "",
        ),
        ("cat", "七、人员投入-统计与差异（W4）", None, None, None, None),
        (
            "task",
            "进场次数、停留天数、累计在场人天口径实现",
            d(2026, 10, 29),
            d(2026, 11, 1),
            owner,
            "",
        ),
        (
            "task",
            "计划对比差异清单；页面禁止「工时」误导文案",
            d(2026, 11, 2),
            d(2026, 11, 4),
            owner,
            "",
        ),
        ("cat", "八、人员投入-审核流（W5）", None, None, None, None),
        (
            "task",
            "待审确认/驳回；跨项目本项目计入天数；审计落库",
            d(2026, 11, 5),
            d(2026, 11, 8),
            owner,
            "",
        ),
        (
            "task",
            "确认后重算汇总与差异",
            d(2026, 11, 9),
            d(2026, 11, 11),
            owner,
            "",
        ),
        ("cat", "九、人员投入-查询页收口（W6）", None, None, None, None),
        (
            "task",
            "输入项目编号：汇总、按人明细、差异清单",
            d(2026, 11, 12),
            d(2026, 11, 15),
            owner,
            "",
        ),
        (
            "task",
            "实际vs计划图表；主路径联调；人员功能冻结",
            d(2026, 11, 16),
            d(2026, 11, 18),
            owner,
            "",
        ),
        ("cat", "十、工作台整合（W7）", None, None, None, None),
        (
            "task",
            "导航挂接「人员投入」与问答；统一当前项目",
            d(2026, 11, 19),
            d(2026, 11, 22),
            owner,
            "",
        ),
        (
            "task",
            "越权用例、401跳转与空态提示",
            d(2026, 11, 23),
            d(2026, 11, 25),
            owner,
            "",
        ),
        ("cat", "十一、生产部署（W8）", None, None, None, None),
        (
            "task",
            "生产env模板、改默认口令、备份与重启恢复演练",
            d(2026, 11, 26),
            d(2026, 11, 29),
            owner,
            "",
        ),
        (
            "task",
            "运维启停手册；目标机首装（若环境就绪）",
            d(2026, 11, 30),
            d(2026, 12, 2),
            owner,
            "",
        ),
        ("cat", "十二、RAG保底与试点语料（W9）", None, None, None, None),
        (
            "task",
            "试点项目文档入库；准备20条验收问句",
            d(2026, 12, 3),
            d(2026, 12, 6),
            owner,
            "",
        ),
        (
            "task",
            "问答通过线验收；权限/引用挡路缺陷修复",
            d(2026, 12, 7),
            d(2026, 12, 9),
            owner,
            "",
        ),
        ("cat", "十三、真项目端到端试跑（W10）", None, None, None, None),
        (
            "task",
            "试点项目灌真实/准真实人员记录与计划并走审核",
            d(2026, 12, 10),
            d(2026, 12, 13),
            owner,
            "",
        ),
        (
            "task",
            "30分钟演示脚本；P0清零；业务试用反馈",
            d(2026, 12, 14),
            d(2026, 12, 16),
            owner,
            "",
        ),
        ("cat", "十四、冻结与交接（W11）", None, None, None, None),
        (
            "task",
            "功能冻结；用户操作说明与运维手册定稿",
            d(2026, 12, 17),
            d(2026, 12, 20),
            owner,
            "",
        ),
        (
            "task",
            "已知问题清单；演示环境可重建验证",
            d(2026, 12, 21),
            d(2026, 12, 23),
            owner,
            "",
        ),
        ("cat", "十五、正式验收（W12）", None, None, None, None),
        (
            "task",
            "按验收清单逐条勾选；只修现场问题",
            d(2026, 12, 24),
            d(2026, 12, 28),
            owner,
            "",
        ),
        (
            "task",
            "交接包交付；后续项一页纸（计划图/指标/文档草稿/门禁）",
            d(2026, 12, 29),
            d(2026, 12, 31),
            owner,
            "",
        ),
    ]

    def fmt(dt: date) -> str:
        return f"{dt.year}/{dt.month}/{dt.day}"

    r = 4
    seq = 1
    task_i = 0
    done_task_count = 0
    plan_task_count = 0
    for kind, name, s, e, who, note in rows:
        is_done = note == "已完成" or note == "done"
        if kind == "cat":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
            cell = ws.cell(r, 1, name)
            cell.font = cat_font
            row_fill = done_cat_fill if is_done else cat_fill
            cell.fill = row_fill
            cell.alignment = left
            for col in range(1, 10):
                c = ws.cell(r, col)
                c.border = thin
                c.fill = row_fill
            ws.row_dimensions[r].height = 22
        else:
            snapped_start, snapped_end, days = snap_range(s, e)
            reason = "已完成" if note == "已完成" else (note or "")
            values = [
                seq,
                name,
                fmt(snapped_start),
                fmt(snapped_end),
                days,
                who,
                reason,
                "",
                "",
            ]
            if note == "已完成":
                fill = done_fill
                done_task_count += 1
            else:
                fill = alt_fill if task_i % 2 else None
                plan_task_count += 1
                task_i += 1
            for col, val in enumerate(values, 1):
                c = ws.cell(r, col, val)
                c.font = cell_font
                c.border = thin
                c.alignment = center if col != 2 else left
                if fill:
                    c.fill = fill
            ws.row_dimensions[r].height = 28
            seq += 1
        r += 1

    ws2 = wb.create_sheet("说明与验收")
    notes = [
        "说明",
        f"1. 整盘工期：2026/9/1～2026/12/31，合计 {total_days} 个工作日（已按单双休扣除休息日）。",
        f"2. 其中已完成底座约 {done_days} 个工作日（绿色行，「计划更新原因」=已完成）；"
        f"10/8 起后续计划约 {remain_days} 个工作日。",
        "3. 一至三阶段对应现有 enterprise-kb 已交付能力：RAG问答、文档接入增强、组织权限、Compose部署、品牌前端。",
        "4. 四至十五阶段为 10/8 起首批业务：人员投入统计 + 生产部署 + RAG试点保底；计划图/指标测算/文档草稿/门禁不在本表。",
        "5. 单双休：每周日休息；ISO 单周周六上班，双周周六休息。落在休息日的起止已改到相邻工作日。",
        "6. 项目经理/责任人：陈锐。",
        "",
        "12/31 验收清单（摘要）",
        "① 输入试点项目编号可出投入汇总与差异清单",
        "② 出差/调令/在场可导入；重复合并；重叠与跨项目须人工确认",
        "③ 次数/天数与手算一致；不把在场天数当工时",
        "④ 未授权不可见；RAG试点可问带引用（底座能力）",
        "⑤ 目标环境可启停备份；非首批功能不作为本版缺陷",
    ]
    for i, line in enumerate(notes, 1):
        ws2.cell(i, 1, line).font = Font(
            name="微软雅黑", size=11, bold=(i == 1 or i == 9)
        )
    ws2.column_dimensions["A"].width = 96

    widths = [6, 56, 12, 12, 8, 10, 14, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 22
    ws.freeze_panes = "A4"

    out = os.path.join(
        os.path.dirname(os.path.dirname(__file__)),
        "docs",
        "V1-开发进度计划-含已完成底座-v3.xlsx",
    )
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(out)
    print(
        f"tasks={seq - 1} (done={done_task_count}, plan={plan_task_count}), "
        f"total_workdays={total_days}, done_span≈{done_days}, remain≈{remain_days}"
    )


if __name__ == "__main__":
    main()
