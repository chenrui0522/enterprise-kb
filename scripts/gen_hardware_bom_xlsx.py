"""Generate intranet multi-agent hardware BOM Excel."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "intranet-multi-agent-hardware-bom.xlsx"

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def style_body(ws, start_row: int, end_row: int, cols: int) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = WRAP


def autosize(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws, start_row: int, headers: list[str], rows: list[tuple]) -> int:
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    style_header(ws, start_row, len(headers))
    for i, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    end = start_row + len(rows)
    style_body(ws, start_row + 1, end, len(headers))
    return end


def main() -> None:
    wb = Workbook()

    # Sheet 1
    ws1 = wb.active
    ws1.title = "选型结论"
    ws1["A1"] = "企业级多 Agent 平台 - 内网本地 LLM 硬件选型结论"
    ws1["A1"].font = TITLE_FONT
    ws1.merge_cells("A1:C1")
    ws1["A2"] = "版本日期: 2026-09-21 | 适用: ~110人公司 / 全员日常入口 / 高峰可短暂排队"
    ws1.merge_cells("A2:C2")
    write_table(
        ws1,
        4,
        ["决策项", "结论", "说明"],
        [
            ("部署路径", "路径2: 内网本地 LLM + 多 Agent", "生成不走公有云 API；控制面/检索/推理分角色"),
            ("用户规模", "约110人，全员日常入口", "按入口级高峰设计，不按大厂集群设计"),
            ("峰值并发假设", "同时在线约25-45人", "多 Agent 每问约2-4次本地 LLM 调用"),
            ("高峰体验", "接受短暂排队(约3-10秒)", "单副本生成卡 + 限流排队，不上双生成副本"),
            ("主生成模型", "Qwen2.5-32B-Instruct（优先）或 Qwen3-32B", "INT4优先，余量够再用FP8；备选 GLM-4-32B-0414"),
            ("推理服务", "vLLM 单卡", "替换现有 DeepSeek 出网位；gpu_memory_utilization 0.90"),
            ("检索模型", "沿用 bge-m3 + bge-reranker-v2-m3", "与现 enterprise-kb 一致，L4 常驻；优先 TEI"),
            ("推荐形态", "2台为主，OCR按需第3台", "机1业务数据 + 机2双卡推理"),
            ("暂缓", "72B多卡 / K8s / Milvus集群 / 双生成副本", "人头与并发撑不起运维成本"),
        ],
    )
    autosize(ws1, [22, 42, 55])
    ws1.row_dimensions[1].height = 24

    # Sheet 2
    ws2 = wb.create_sheet("硬件清单-第一年")
    ws2["A1"] = "第一年推荐采购 / 配置清单（BOM）"
    ws2["A1"].font = TITLE_FONT
    ws2.merge_cells("A1:J1")
    end2 = write_table(
        ws2,
        3,
        ["机位", "角色", "数量", "CPU", "内存", "系统盘/数据盘", "GPU", "运行组件", "优先级", "备注"],
        [
            (
                "机1",
                "业务与数据面",
                1,
                "32核",
                "128GB",
                "系统 512GB SSD + 数据 NVMe 2-4TB",
                "无",
                "API、Worker×2、Postgres、Redis、Milvus standalone、MinIO/文档存储",
                "P0必采",
                "可水平再起一个API实例；语料未到十万页前Milvus单机即可",
            ),
            (
                "机2",
                "推理面(核心)",
                1,
                "16-32核",
                "64-128GB",
                "系统 512GB SSD + 模型缓存 NVMe 1-2TB",
                "见「GPU明细」双卡",
                "Qwen2.5-32B-Instruct(vLLM) + bge-m3/bge-reranker-v2-m3(TEI)",
                "P0必采",
                "整套预算大头；生成与检索分卡",
            ),
            (
                "机3",
                "OCR解析(二期)",
                1,
                "16核",
                "64GB",
                "系统 512GB + 缓存 1TB",
                "1×24GB",
                "MinerU / Docling",
                "P1按需",
                "扫描件/复杂版式一多再上；避免拖垮问答",
            ),
        ],
    )
    autosize(ws2, [8, 16, 8, 12, 12, 36, 22, 48, 12, 40])
    for r in range(4, end2 + 1):
        ws2.row_dimensions[r].height = 55

    # Sheet 3
    ws3 = wb.create_sheet("GPU明细")
    ws3["A1"] = "推理机 GPU 分配"
    ws3["A1"].font = TITLE_FONT
    ws3.merge_cells("A1:G1")
    write_table(
        ws3,
        3,
        ["卡位", "建议显存", "负载", "模型/服务", "精度建议", "并发策略", "备注"],
        [
            (
                "卡A",
                "24GB",
                "检索向量化 + 重排",
                "bge-m3 + bge-reranker-v2-m3",
                "FP16",
                "与现网一致；rerank限流",
                "两模型合计约4.7GB；reranker官方约2GB。优先TEI，避开Xinference显存累积",
            ),
            (
                "卡B",
                "L20 48GB",
                "主生成(单副本)",
                "Qwen2.5-32B-Instruct（优先）或 Qwen3-32B",
                "INT4优先（约16.4GB）；FP8约33GB",
                "vLLM max-num-seqs 8；高峰排队3-10秒",
                "不要与OCR/embed共卡。备选 GLM-4-32B-0414",
            ),
            (
                "卡C(二期)",
                "24GB",
                "OCR / 版面",
                "MinerU(及可选Docling)",
                "pipeline或VLM按需",
                "入库队列与问答隔离",
                "对应机3",
            ),
        ],
    )
    ws3["A8"] = "显存粗算参考(选型用，非精确压测)"
    ws3["A8"].font = BOLD
    write_table(
        ws3,
        9,
        ["模型档", "权重约(INT4/FP16)", "第一年建议", "对本公司"],
        [
            ("7B/14B", "INT4 ~5-10GB / FP16 ~15-30GB", "可作路由/工具小模型", "不建议作为全员入口唯一大脑"),
            ("32B Qwen", "INT4约16.4GB / FP8约33GB", "Qwen2.5-32B-Instruct优先", "L20 48GB可承载，余量给KV"),
            ("70B/72B", "INT4 ~40GB / FP16 ~145GB+", "需多卡", "第一年暂缓"),
        ],
    )
    autosize(ws3, [12, 28, 22, 36, 22, 32, 36])

    # Sheet 4
    ws4 = wb.create_sheet("分期采购")
    ws4["A1"] = "分期建议"
    ws4["A1"].font = TITLE_FONT
    end4 = write_table(
        ws4,
        3,
        ["阶段", "采购内容", "触发条件", "目标"],
        [
            (
                "一期(现在)",
                "机1 + 机2(卡A 24GB + 卡B 48/80GB)",
                "立项启动",
                "全员入口可用；本地32B；检索隔离；排队可感知",
            ),
            (
                "二期",
                "机3 OCR专用卡",
                "扫描件/复杂版式入库明显拖慢问答或Worker堆积",
                "解析与问答算力隔离",
            ),
            (
                "三期(可选)",
                "第二张生成卡做副本 或 升72B评测",
                "午高峰排队经常>10-15秒且业务不可接受；或评测证明32B质量不够",
                "吞吐或质量升级",
            ),
        ],
    )
    autosize(ws4, [14, 42, 48, 40])
    for r in range(4, end4 + 1):
        ws4.row_dimensions[r].height = 50

    # Sheet 5
    ws5 = wb.create_sheet("暂缓与配套原则")
    ws5["A1"] = "明确暂缓项"
    ws5["A1"].font = TITLE_FONT
    write_table(
        ws5,
        3,
        ["类别", "内容", "原因"],
        [
            ("算力", "72B多卡旗舰集群", "110人规模溢价高，日常打不满"),
            ("平台", "第一天K8s大集群", "2-3台角色机 Compose/轻量编排更合适"),
            ("向量", "Milvus分布式集群", "语料未到十万页级前standalone足够"),
            ("高可用", "双生成副本(一期)", "已接受高峰短暂排队"),
            ("混部", "单卡硬塞生成+embed+OCR", "全员入口高峰会互相拖死"),
        ],
    )
    ws5["A10"] = "与硬件配套的软件原则(非实现清单)"
    ws5["A10"].font = BOLD
    write_table(
        ws5,
        11,
        ["原则", "说明"],
        [
            ("全局限流", "按卡B显存设置同时in-flight生成上限"),
            ("可见排队", "提示前方等待，避免静默转圈"),
            ("短链路优先", "入口默认: 路由→检索→一次主答；重Agent限流或异步"),
            ("现有栈映射", "DeepSeek API→vLLM上的Qwen2.5-32B；bge-m3/reranker→L4上TEI；MinerU→机3"),
        ],
    )
    autosize(ws5, [16, 55, 40])

    # Sheet 6
    ws6 = wb.create_sheet("假设与边界")
    ws6["A1"] = "选型假设(变更需重估BOM)"
    ws6["A1"].font = TITLE_FONT
    end6 = write_table(
        ws6,
        3,
        ["假设ID", "内容", "若变化则..."],
        [
            ("A1", "公司约110人，系统为全员日常入口", "人数或使用率显著上升→评估第二张生成卡"),
            ("A2", "高峰可接受约3-10秒排队", "必须近乎秒开→一期就上双生成副本"),
            ("A3", "主生成走内网本地32B", "必须72B质量→多卡与机柜级预算"),
            ("A4", "文档正文留内网；生成也本地", "若生成可出网→可回到云端LLM省掉卡B大头"),
            ("A5", "检索继续本地embedding/rerank", "与现enterprise-kb架构一致"),
            ("A6", "OCR非天天全员扫库", "扫描件为主力入库→提前上机3"),
            ("A7", "GPU生态按NVIDIA CUDA类推断", "若强制国产卡→需单独做推理框架与型号适配表"),
        ],
    )
    autosize(ws6, [10, 42, 48])
    for r in range(4, end6 + 1):
        ws6.row_dimensions[r].height = 35

    # Sheet 7: line-item reference prices (CNY, tax-inclusive channel quotes, Sep 2026)
    ws7 = wb.create_sheet("成本明细")
    ws7["A1"] = "采购成本明细（人民币、含税参考价，2026-09 公开渠道，非正式报价）"
    ws7["A1"].font = TITLE_FONT
    ws7.merge_cells("A1:I1")
    ws7["A2"] = "口径: 国内渠道/电商/整机反推的单台或单卡含税区间。实际以供应商报价、质保年限、是否全新为准。不含实施人力、机柜租赁、软件授权。"
    ws7.merge_cells("A2:I2")
    end7 = write_table(
        ws7,
        4,
        ["类别", "物料", "规格要点", "阶段", "单价低(元)", "单价高(元)", "参考单价(元)", "数量", "小计参考(元)", "依据"],
        [
            ("主机", "机1 业务服务器", "32核 / 128GB / 系统盘+NVMe 2-4TB", "一期", 45000, 85000, 65000, 1, 65000, "品牌入门机架或渠道组装，2026国内常见成交带"),
            ("主机", "机2 GPU机箱(不含卡)", "16-32核 / 128GB / NVMe 1-2TB / 双电 1600W+", "一期", 35000, 60000, 45000, 1, 45000, "4U白牌GPU服务器主机，不含显卡"),
            ("GPU", "卡A 推荐 L4", "24GB / 约72W / 数据中心卡", "一期", 18000, 25000, 21000, 1, 21000, "2026国内预算价综述约1.8-2.5万；适合7x24常驻embed"),
            ("GPU", "卡A 备选 RTX 4090", "24GB 消费级", "备选", 16000, 25000, 20000, 1, 20000, "渠道常见带；功耗约450W，不作为推荐常驻卡"),
            ("GPU", "卡B 推荐 L20", "48GB / 国内推理卡", "一期", 30000, 45000, 38000, 1, 38000, "预算价综述3-4.5万；ZOL八卡L20整机47.5万反推约3.5-4万/卡"),
            ("GPU", "卡B 宽裕 L40S", "48GB Ada", "对照", 60489, 77918, 69000, 1, 69000, "京东自营报道: 2026-03约60489元, 2026-09约77918元"),
            ("GPU", "卡B 对照 RTX 5090", "32GB", "对照", 40000, 47000, 43500, 1, 43500, "SMM 2026-09-16 华南报价约43500元/卡；显存小于48GB推荐档"),
            ("GPU", "卡B 旗舰 A100", "80GB", "对照", 100000, 150000, 120000, 1, 120000, "2026国内预算价综述10-15万"),
            ("GPU", "卡B 旗舰 H20", "96GB", "对照", 100000, 180000, 140000, 1, 140000, "2026国内预算价综述10-18万；供货与合规以当时政策为准"),
            ("主机", "机3 OCR服务器", "16核 / 64GB / 1TB缓存", "二期", 25000, 40000, 32000, 1, 32000, "单卡推理主机"),
            ("GPU", "卡C OCR", "L4 24GB 或同级", "二期", 18000, 25000, 21000, 1, 21000, "与卡A同价带"),
            ("配套", "交换机+UPS(简易)", "千兆/万兆小交换机 + 在线式UPS", "选配", 8000, 20000, 12000, 1, 12000, "小机房自建；已有机房则可为0"),
        ],
    )
    autosize(ws7, [10, 24, 42, 10, 14, 14, 16, 8, 16, 55])
    for r in range(5, end7 + 1):
        for c in (5, 6, 7, 9):
            ws7.cell(row=r, column=c).number_format = "#,##0"
        ws7.row_dimensions[r].height = 32
    ws7.column_dimensions["J"].width = 62

    # Sheet 8: scheme totals
    ws8 = wb.create_sheet("方案总价")
    ws8["A1"] = "方案总价（按上表参考单价加总）"
    ws8["A1"].font = TITLE_FONT
    ws8.merge_cells("A1:H1")
    end8 = write_table(
        ws8,
        3,
        ["方案", "组成", "一期CAPEX低(元)", "一期CAPEX高(元)", "一期参考(元)", "含选配参考(元)", "再加二期OCR参考(元)", "相对推荐"],
        [
            (
                "推荐(采用)",
                "机1 + 机2主机 + L4 24GB + L20 48GB",
                128000,
                215000,
                169000,
                181000,
                234000,
                "第一年默认",
            ),
            (
                "宽裕48GB",
                "机1 + 机2主机 + L4 + L40S 48GB",
                158489,
                247918,
                200000,
                212000,
                265000,
                "约贵3.1万，同显存档算力更强",
            ),
            (
                "对照32GB",
                "机1 + 机2主机 + L4 + RTX 5090 32GB",
                138000,
                217000,
                174500,
                186500,
                239500,
                "总价接近推荐，但生成显存只有32GB",
            ),
            (
                "旗舰80/96GB",
                "机1 + 机2主机 + L4 + H20 96GB",
                198000,
                350000,
                271000,
                283000,
                336000,
                "约贵10万+；110人第一年不作为默认",
            ),
        ],
    )
    for r in range(4, end8 + 1):
        for c in range(3, 8):
            ws8.cell(row=r, column=c).number_format = "#,##0"
        ws8.row_dimensions[r].height = 36
    autosize(ws8, [14, 46, 18, 18, 16, 16, 22, 36])

    ws8["A9"] = "3年持有成本（推荐方案，电费另计）"
    ws8["A9"].font = BOLD
    write_table(
        ws8,
        10,
        ["项目", "口径", "金额(元)", "说明"],
        [
            ("一期硬件参考", "推荐方案中位，不含选配", 169000, "机1 6.5万 + 机2主机 4.5万 + L4 2.1万 + L20 3.8万"),
            ("选配网络电源", "简易交换机+UPS", 12000, "已有机房可记0"),
            ("年电费", "1.0元/kWh，全年开机，GPU平均负载约30%", 5000, "机1约0.2kW常开；机2主机约0.15kW + L4/L20按30%计。各地电价按实际替换"),
            ("3年电费", "年电费 x 3", 15000, "相对硬件可忽略，但机房电容量要按峰值预留"),
            ("3年TCO中位", "一期硬件 + 选配 + 3年电费", 196000, "不含二期OCR、不含人力、不含品牌维保上浮"),
            ("品牌整机上浮", "联想/戴尔等含3年保", 0, "在白牌中位上常见上浮约15%-30%，询价时单列"),
            ("二期OCR", "机3 + L4", 53000, "扫描件入库堆积后再买；不进一期TCO"),
        ],
    )
    for r in range(11, 18):
        ws8.cell(row=r, column=3).number_format = "#,##0"
        ws8.row_dimensions[r].height = 28

    ws8["A19"] = "峰值功耗(选型用电容量，不是电费)"
    ws8["A19"].font = BOLD
    write_table(
        ws8,
        20,
        ["机器", "峰值约(W)", "说明"],
        [
            ("机1", 300, "CPU服务器峰值，常载更低"),
            ("机2 推荐", 750, "主机 + L4约72W + L20约275-350W，留电源余量到1600W"),
            ("机2 若改4090+L40S", 1200, "4090约450W + L40S约350W + 主机"),
            ("机3", 450, "主机 + 一张24GB卡"),
        ],
    )

    ws8["A26"] = (
        "加总核对: 推荐一期 65000+45000+21000+38000=169000；"
        "区间 45000+35000+18000+30000=128000 至 85000+60000+25000+45000=215000；"
        "含选配中位 169000+12000=181000；含二期 181000+32000+21000=234000。"
        "宽裕方案用L40S参考价69000替换L20的38000，差额+31000。"
        "旗舰用H20参考价140000替换，差额+102000。"
    )
    ws8["A26"].alignment = WRAP
    ws8.merge_cells("A26:H26")
    ws8.row_dimensions[26].height = 48

    ws9 = wb.create_sheet("推荐模型", 1)
    ws9["A1"] = "推荐本地模型组合（L20 48GB 生成 + L4 24GB 检索）"
    ws9["A1"].font = TITLE_FONT
    ws9.merge_cells("A1:H1")
    ws9["A2"] = "最稳妥组合: 主生成 Qwen2.5-32B-Instruct（或 Qwen3-32B）；检索沿用 bge-m3 + bge-reranker-v2-m3。显存为选型估算，上线前以本机压测为准。"
    ws9.merge_cells("A2:H2")
    end9 = write_table(
        ws9,
        4,
        ["角色", "推荐模型", "精度", "预估显存", "部署卡", "部署方式", "状态", "说明"],
        [
            (
                "主生成",
                "Qwen2.5-32B-Instruct（优先）或 Qwen3-32B",
                "INT4优先，余量够可FP8",
                "INT4约16.4GB；FP8约33GB",
                "L20 48GB",
                "vLLM，单卡 tensor-parallel-size=1",
                "采用",
                "中文知识问答成熟，与现有RAG架构同形态。INT4后L20余量给KV和多Agent并发。",
            ),
            (
                "主生成备选",
                "GLM-4-32B-0414",
                "按官方量化",
                "同32B档，需单独压测",
                "L20 48GB",
                "vLLM（以当时支持为准）",
                "备选",
                "MIT可商用。指令遵循和函数调用较强；多Agent工具调用多时再对比，不作为第一天上线模型。",
            ),
            (
                "Embedding",
                "bge-m3",
                "FP16",
                "约2GB",
                "L4 24GB",
                "TEI优先；Xinference可过渡",
                "沿用",
                "与enterprise-kb一致，1024维，无需重嵌入。不更换。",
            ),
            (
                "Reranker",
                "bge-reranker-v2-m3",
                "FP16",
                "约2GB；与bge-m3合计约4.7GB",
                "L4 24GB",
                "TEI优先",
                "沿用",
                "官方显存约2GB。Xinference加载后可能每次调用累积显存，若最新版仍未修复则改TEI并定期重启作为过渡。",
            ),
        ],
    )
    for r in range(5, end9 + 1):
        ws9.row_dimensions[r].height = 48
    autosize(ws9, [14, 42, 24, 28, 14, 36, 10, 55])

    ws9["A10"] = "vLLM 起步参数（Qwen2.5-32B-Instruct INT4/AWQ）"
    ws9["A10"].font = BOLD
    write_table(
        ws9,
        11,
        ["参数", "建议值", "说明"],
        [
            ("模型", "Qwen/Qwen2.5-32B-Instruct-AWQ", "优先INT4/AWQ；质量不够再试FP8"),
            ("gpu-memory-utilization", "0.90", "48GB跑INT4可偏激进，仍留约10%安全余量"),
            ("max-model-len", "8192", "企业问答默认上下文；父块很长再按压测上调"),
            ("max-num-seqs", "8", "同时在飞的序列上限，用来挡住多Agent把KV打满"),
            ("tensor-parallel-size", "1", "单张L20，不做张量并行"),
            ("显存不足时", "先降 max-num-seqs 或 max-model-len", "不要先降 gpu-memory-utilization"),
        ],
    )

    ws9["A19"] = "多Agent与压测"
    ws9["A19"].font = BOLD
    write_table(
        ws9,
        20,
        ["项", "约定"],
        [
            ("调用放大", "每问约2-4次本地LLM，KV压力约为单轮对话的2-4倍"),
            ("排队", "vLLM侧限制同时in-flight；前端显示前方等待。高峰接受约3-10秒排队"),
            ("路由/改写", "与主回答共用同一个32B，不另驻留7B，避免再占L20显存"),
            ("首轮压测", "用vLLM benchmark看约5人并发的吞吐和延迟，确认仍落在可排队范围内"),
            ("检索服务", "L4上TEI常驻bge-m3与bge-reranker-v2-m3，经API给检索调用"),
        ],
    )
    for r in range(21, 26):
        ws9.row_dimensions[r].height = 28
    ws9.column_dimensions["B"].width = 72

    OUT.parent.mkdir(parents=True, exist_ok=True)
    targets = [
        OUT.with_name("intranet-multi-agent-hardware-bom-priced.xlsx"),
        OUT,
    ]
    saved: list[Path] = []
    for path in targets:
        try:
            wb.save(path)
            saved.append(path)
        except PermissionError:
            continue
    if not saved:
        alt = OUT.with_name("intranet-multi-agent-hardware-bom-models.xlsx")
        wb.save(alt)
        print(alt)
    else:
        for path in saved:
            print(path)

    write_summary()


def write_summary() -> None:
    """Single-sheet rollup of hardware, price, and model choices."""
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws["A1"] = "内网多Agent平台 硬件与模型总表"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:I1")
    ws["A2"] = "约110人全员入口 | 高峰可排队3-10秒 | 价格为2026-09国内含税参考价，非正式报价 | 一期推荐合计见表末"
    ws.merge_cells("A2:I2")
    headers = ["阶段", "类别", "项目", "推荐配置", "数量", "参考单价(元)", "小计(元)", "单价区间(元)", "说明"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(headers))

    rows = [
        ("一期", "主机", "机1 业务与数据", "32核 / 128GB / NVMe 2-4TB", 1, 65000, 65000, "4.5万-8.5万", "API、Worker、Postgres、Redis、Milvus、MinIO"),
        ("一期", "主机", "机2 推理主机(不含卡)", "16-32核 / 128GB / NVMe 1-2TB / 双电", 1, 45000, 45000, "3.5万-6.0万", "4U GPU机箱"),
        ("一期", "GPU", "卡A 检索", "NVIDIA L4 24GB", 1, 21000, 21000, "1.8万-2.5万", "常驻检索，不跑生成"),
        ("一期", "模型", "Embedding + Reranker", "bge-m3 + bge-reranker-v2-m3，FP16", 1, None, None, "—", "沿用现有链路，合计约4.7GB，部署在L4。优先TEI"),
        ("一期", "GPU", "卡B 生成", "NVIDIA L20 48GB", 1, 38000, 38000, "3.0万-4.5万", "单卡，不与检索/OCR混部"),
        ("一期", "模型", "主生成", "Qwen2.5-32B-Instruct（优先）或 Qwen3-32B", 1, None, None, "—", "INT4约16.4GB，FP8约33GB。vLLM：utilization 0.90，max-model-len 8192，max-num-seqs 8"),
        ("一期", "模型", "主生成备选", "GLM-4-32B-0414", 1, None, None, "—", "MIT可商用。工具调用多时再对比，不作为第一天模型"),
        ("选配", "配套", "交换机 + UPS", "千兆/万兆小交换机 + 在线式UPS", 1, 12000, 12000, "0.8万-2.0万", "已有机房则可为0"),
        ("二期", "主机", "机3 OCR", "16核 / 64GB / 1TB", 1, 32000, 32000, "2.5万-4.0万", "扫描件入库堆积后再买"),
        ("二期", "GPU+模型", "卡C 解析", "L4 24GB + MinerU pipeline", 1, 21000, 21000, "1.8万-2.5万", "第一年不上VLM后端，避免和问答抢卡"),
        ("约定", "体验", "高峰排队", "接受约3-10秒", None, None, None, "—", "每问约2-4次本地LLM；前端显示前方等待。不上第二张生成卡"),
        ("约定", "规模", "使用场景", "约110人，全员日常入口", None, None, None, "—", "同时在线按25-45人估计"),
    ]
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    for i, row in enumerate(rows, 5):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = THIN
            cell.alignment = WRAP
            if c in (6, 7) and isinstance(v, int):
                cell.number_format = "#,##0"
        ws.row_dimensions[i].height = 36

    # Totals sit on the same sheet, formulas only over priced rows.
    totals = [
        (17, "合计", "一期硬件", "机1+机2主机+L4+L20", "=G5+G6+G7+G9", "12.8万-21.5万", "不含选配、不含二期"),
        (18, "合计", "一期含选配", "上一行 + 交换机/UPS", "=G17+G12", "—", "已有机房可减去选配"),
        (19, "合计", "一期+二期", "再加机3与OCR卡", "=G18+G13+G14", "—", "OCR按需，不进第一笔采购"),
        (20, "合计", "推荐方案3年TCO", "一期含选配 + 3年电费约1.5万", "=G18+15000", "—", "电价按1元/度、GPU平均负载约30%。不含人力与品牌维保上浮"),
    ]
    for r, phase, cat, item, total, band, note in totals:
        values = [phase, cat, item, "", None, None, total, band, note]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = THIN
            cell.alignment = WRAP
            cell.fill = total_fill
            cell.font = BOLD
            if c == 7:
                cell.number_format = '#,##0'
        ws.row_dimensions[r].height = 32

    autosize(ws, [10, 12, 22, 48, 10, 16, 14, 16, 62])
    ws.auto_filter.ref = "A4:I16"
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "内网多Agent平台 硬件与模型总表"
    ws.print_title_rows = "1:4"

    out = OUT.with_name("intranet-multi-agent-总表.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        out = OUT.with_name("intranet-multi-agent-总表-new.xlsx")
        wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
